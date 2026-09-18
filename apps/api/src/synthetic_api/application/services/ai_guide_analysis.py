# -*- coding: utf-8 -*-
# =============================================================================
# 파일명: ai_guide_analysis.py
# 경로: apps/api/src/synthetic_api/application/services/ai_guide_analysis.py
# 목적: 공공 데이터 구조 분석 및 AI 친화도 가이드 모델 생성을 수행함
# 작성자: 개발팀
# 작성일: 2026-09-16
# 수정일: 2026-09-16
# =============================================================================
"""Deterministic, guide-only structural analysis. No remote context resolution."""
from __future__ import annotations
import base64
import csv
import hashlib
import io
import json
import re
from typing import Any
from lxml import etree

MAX_BYTES = 32 * 1024 * 1024
MAX_NODES = 10000000

# These patterns describe generic metadata concepts only.  They intentionally
# do not name a domain-specific field (for example, a particular sensor or
# business measure); the analyzer records candidates and leaves the rule to
# the institution when the source does not prove it.
_DERIVATION_HINT = re.compile(
    r'(?:date|time|timestamp|datetime|year|month|day|hour|minute|second|'
    r'total|sum|avg|average|mean|min|max|count|rate|ratio|percent|'
    r'날짜|일자|시간|시각|연도|월|일|시|분|초|합계|평균|최소|최대|건수|비율|비중)',
    re.IGNORECASE,
)
_QUALITY_FLAG_HINT = re.compile(
    r'(?:flag|quality|status|state|valid|validity|missing|null|imput|'
    r'code|class|category|indicator|품질|플래그|상태|결측|누락|보정|검증|유효|코드|구분|여부)',
    re.IGNORECASE,
)
_QUALITY_FLAG_EXPLICIT_HINT = re.compile(
    r'(?:flag|quality|status|state|valid|validity|missing|null|imput|indicator|code|class|category)',
    re.IGNORECASE,
)
_BINARY_FLAG_VALUES = {
    '0', '1', '2', 'y', 'n', 'yes', 'no', 'true', 'false',
    'valid', 'invalid', 'good', 'bad', 'ok', 'ng', 'unknown',
    '예', '아니오', '정상', '오류', '유효', '무효', '미확인',
}


def _sample_values(field: dict) -> list:
    values = field.get('examples')
    if not isinstance(values, list):
        values = field.get('sample_values')
    return values if isinstance(values, list) else []


def _processing_candidates(fields: list[dict], data_category: str) -> dict:
    """Build conservative, source-grounded processing observations.

    This function never invents a formula or applies a correction.  It only
    records observed missingness, generic structural hints, and candidate flag
    values.  Any rule that would change a value remains null and therefore is
    surfaced as REVIEW_REQUIRED by the canonical finalizer.
    """
    missing_rules = []
    derived_fields = []
    quality_flags = []

    for field in fields:
        path = str(field.get('path') or '')
        name = str(field.get('name') or path)
        field['value_origin'] = 'observed'
        field['derived'] = False

        null_count = int(field.get('null_count') or 0)
        empty_count = int(field.get('empty_count') or 0)
        structural_missing = int(field.get('missing_count') or 0)
        missing_count = null_count + empty_count + structural_missing
        if missing_count:
            occurrences = int(field.get('occurrences') or 0)
            denominator = occurrences + structural_missing
            missing_rules.append({
                'target_field_ids': [path],
                'detected_missing_count': missing_count,
                'detected_missing_ratio': (missing_count / denominator) if denominator else None,
                'method': None,
                'method_description': '결측치 처리 방식 기관 확인 필요',
                'steps': [],
                'grouping_keys': [],
                'temporal_window': None,
                'parameters': {},
                'fallback_method': None,
                'clipping_rule': None,
                'affected_record_count': None,
                'quality_flag_field': None,
                'original_value_preserved': True,
                'limitations': '원천 데이터에서 대체·삭제·보간 규칙을 확인할 수 없음',
            })

        # A temporal/aggregate-shaped name is only a candidate for a possible
        # derivation.  The source field remains observed and no formula is
        # supplied without an explicit contract.
        if _DERIVATION_HINT.search(name) and not set(field.get('types') or ()) <= {'object', 'array'}:
            derived_fields.append({
                'field_id': path,
                'field_name': name,
                'derivation_type': 'observed',
                'source_fields': [path],
                'description': '원천에서 직접 관측된 필드이며 파생·계산 여부는 기관 확인 필요',
                'method': None,
                'formula_or_rule': None,
                'aggregation': None,
                'window': None,
                'unit': None,
                'parameters': {},
                'fallback_rule': None,
                'reproducible': None,
                'confidence': None,
                'notes': '구조적 명칭만으로 파생 규칙을 확정하지 않음',
            })

        samples = _sample_values(field)
        normalized = {str(value).strip().lower() for value in samples if value is not None}
        name_hint = bool(_QUALITY_FLAG_HINT.search(name))
        non_numeric_binary = any(not isinstance(value, (int, float, bool)) for value in samples)
        binary_hint = len(normalized) >= 2 and normalized <= _BINARY_FLAG_VALUES
        explicit_flag_name = bool(_QUALITY_FLAG_EXPLICIT_HINT.search(name))
        # Low-cardinality values alone are insufficient evidence: a numeric
        # measure can also contain 0/1/2.  Require a generic status/code name
        # before surfacing a quality-flag candidate.
        if name_hint and ((binary_hint and (non_numeric_binary or explicit_flag_name)) or not normalized):
            values = []
            seen = set()
            for value in samples:
                key = repr(value)
                if key in seen:
                    continue
                seen.add(key)
                values.append({
                    'code': str(value) if value is not None else None,
                    'name': None,
                    'meaning': None,
                    'value_origin': 'observed',
                })
            quality_flags.append({
                'flag_field': path,
                'description': '품질·상태 플래그 후보이며 코드 의미는 기관 확인 필요',
                'target_fields': [],
                'values': values,
                'meaning': None,
            })

    return {
        'integration': {
            'is_integrated': False,
            'description': '단일 입력 원천을 분석했으며 외부 데이터 결합은 확인되지 않음',
            'method': None,
            'join_type': None,
            'join_keys': [],
            'temporal_alignment': None,
            'spatial_alignment': None,
            'source_dataset_ids': [],
            'external_sources': [],
            'output_description': None,
            'limitations': '외부 원천·결합키·결합 규칙은 입력에서 확인되지 않음',
        },
        'derived_fields': derived_fields,
        'transformations': [],
        'missing_value_processing': missing_rules,
        'outlier_processing': [],
        'quality_flags': quality_flags,
    }


# 데이터를 JSON 문자열로 직렬화함
def dumps(value):
    return json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False, default=str)


# 데이터 값의 세부 데이터 타입을 판별하여 반환함
def kind(value):
    if value is None: return 'null'
    if isinstance(value, bool): return 'boolean'
    if isinstance(value, int): return 'integer'
    if isinstance(value, float): return 'number'
    if isinstance(value, dict): return 'object'
    if isinstance(value, list): return 'array'
    return 'string'


# 텍스트 또는 바이너리 데이터를 분석하여 정형화된 모델 사전을 생성함
def analyze(text: str, fmt: str, binary: str | None = None) -> dict:
    fmt = fmt.lower().strip().lstrip('.')
    if fmt not in {'csv', 'tsv', 'xlsx', 'json', 'jsonld', 'json-ld', 'xml'}:
        raise ValueError('지원 형식: CSV, TSV, XLSX, JSON, JSON-LD, XML')
    raw = base64.b64decode(binary, validate=True) if binary else text.encode('utf-8')
    if len(raw) > MAX_BYTES: raise ValueError('분석 한도 32 MiB 초과: 유효한 파일 단위로 분할하세요.')
    warnings=[]; tables=[]; namespaces={}; traits=set(); fields={}; visited=0
    namespace_bindings=set()

    # 계층 구조 데이터를 재귀 탐색하여 필드 메타데이터를 수집함
    def walk(value, path='', depth=0):
        nonlocal visited
        visited += 1
        if depth > 64 or visited > MAX_NODES: raise ValueError('구조 분석 한도(깊이 64 / 노드 10,000,000) 초과')
        field=fields.setdefault(path or '$', {'path':path or '$', 'types':set(), 'occurrences':0, 'scalar_occurrences':0, 'null_count':0, 'empty_count':0, 'examples':[]})
        field['types'].add(kind(value)); field['occurrences']+=1
        field['scalar_occurrences']+=not isinstance(value,(dict,list))
        field['null_count']+=value is None; field['empty_count']+=value==''
        field['zero_count']=field.get('zero_count',0)+(isinstance(value,(int,float)) and not isinstance(value,bool) and value==0)
        field['false_count']=field.get('false_count',0)+(value is False)
        field['empty_array_count']=field.get('empty_array_count',0)+(isinstance(value,list) and not value)
        if not isinstance(value,(dict,list)) and len(field['examples'])<3 and not any(type(value) is type(example) and value == example for example in field['examples']): field['examples'].append(value)
        if isinstance(value,dict):
            for k,v in value.items():
                if k.startswith('@'): traits.add(k)
                if k == '@context': traits.add('context-' + kind(v))
                if k in {'@graph','@list','@set','@value','@reverse','@included'}: traits.add('jsonld-' + k[1:])
                walk(v,path+'/'+k.replace('~','~0').replace('/','~1').replace('*','~2'),depth+1)
        elif isinstance(value,list):
            for v in value: walk(v,path+'/*',depth+1)

    # 표 데이터를 정형 레코드 사전 목록으로 변환함
    def table(name, headers, records):
        headers=[str(h) if h is not None else '' for h in headers]
        if not headers or any(not h.strip() for h in headers) or len(headers)!=len(set(headers)):
            raise ValueError(f'{name}: 빈 헤더 또는 중복 헤더를 수정하세요.')
        result=[]
        for row in records:
            if len(row)!=len(headers): raise ValueError(f'{name}: 헤더와 레코드의 열 수 불일치')
            result.append(dict(zip(headers,row)))
        tables.append({'name':name,'row_count':len(result),'columns':headers})
        return result

    if fmt in {'csv','tsv'}:
        rows=list(csv.reader(io.StringIO(text.lstrip('\ufeff'), newline=''),delimiter='\t' if fmt=='tsv' else ',',strict=True))
        if not rows: raise ValueError('빈 데이터입니다.')
        data=table('data',rows[0],rows[1:]); category='file'
    elif fmt=='xlsx':
        if not binary: raise ValueError('XLSX는 바이너리 파일 업로드가 필요합니다.')
        from openpyxl import load_workbook
        from zipfile import ZipFile
        with ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist())>512*1024*1024: raise ValueError('XLSX 압축 해제 크기 초과')
        workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
        data={}
        try:
            for sheet in workbook:
                iterator=sheet.iter_rows(values_only=True); headers=next(iterator,None)
                if headers is None: continue
                data[sheet.title]=table(sheet.title,headers,iterator)
        finally: workbook.close()
        warnings.append('수식은 계산하지 않고 원문으로 보존합니다. 시트 첫 행을 헤더로 사용합니다.')
        category='file'
    elif fmt in {'json','jsonld','json-ld'}:
        # JSON 중복 키 존재 여부를 검증하며 사전을 구성함
        def unique(pairs):
            result={}
            for k,v in pairs:
                if k in result: raise ValueError(f'중복 JSON 키: {k}')
                result[k]=v
            return result
        # 유효하지 않은 JSON 수치에 대해 예외를 발생시킴
        def invalid(v): raise ValueError(f'비표준 JSON 숫자: {v}')
        data=json.loads(text.lstrip('\ufeff'),object_pairs_hook=unique,parse_constant=invalid); category='api'
    else:
        if '<!DOCTYPE' in text.upper() or '<!ENTITY' in text.upper(): raise ValueError('DTD/외부 엔티티 XML은 지원하지 않습니다.')
        root=etree.fromstring(text.encode('utf-8'),etree.XMLParser(resolve_entities=False,no_network=True,load_dtd=False))
        for node in root.iter():
            if not isinstance(node.tag,str): continue
            for prefix,uri in node.nsmap.items():
                namespace_bindings.add((str(prefix or '(default)'), uri))
                namespaces[str(prefix or '(default)')]=uri

        # XML 노드를 재귀적으로 파싱하여 파이썬 딕셔너리로 변환함
        def xml_node(node,depth=0):
            if depth>64: raise ValueError('XML 깊이 한도 초과')
            result={'@attributes':dict(node.attrib)} if node.attrib else {}
            content=[]; groups={}
            if node.text is not None: content.append({'text':node.text})
            for child in node:
                if not isinstance(child.tag,str): continue
                value=xml_node(child,depth+1)
                groups.setdefault(child.tag,[]).append(value)
                content.append({'element':child.tag,'value':value})
                if child.tail is not None: content.append({'text':child.tail})
            for k,v in groups.items():
                result[k]=v if len(v)>1 else v[0]
                if len(v)>1: traits.add('repeated-elements')
            if groups:
                if any(c.get('text','').strip() for c in content): result['#content']=content; traits.add('mixed-content')
            else: result['#text']=None if node.get('{http://www.w3.org/2001/XMLSchema-instance}nil') in {'true','1'} else node.text if node.text is not None else ''
            if node.get('{http://www.w3.org/2001/XMLSchema-instance}nil') in {'true','1'}: traits.add('xsi:nil')
            return result
        data={root.tag:xml_node(root)}; category='api'; traits.add('xml')
        warnings.append('XML 문자값은 XSD 없이 숫자/날짜로 강제 변환하지 않습니다. QName은 {URI}local 형식입니다.')
    walk(data)
    for f in fields.values():
        f['types']=sorted(f['types'])
        parent_path=f['path'].rsplit('/',1)[0] or '$'
        parent=fields.get(parent_path)
        f['missing_count']=max(0,parent['occurrences']-f['occurrences']) if f['path'] != '$' and parent and parent['types']==['object'] else None
    if any(k in traits for k in {'@context','@id','@type','@graph','@value','@list'}) or fmt in {'jsonld','json-ld'}:
        traits.add('json-ld')
        warnings.append('JSON-LD 구문 유형화만 수행합니다. 원격 @context 조회·RDF 확장·SHACL 검증은 미수행이며 REVIEW_REQUIRED입니다.')
    candidates=[]

    # 구조 내 배열 및 레코드셋 후보 경로를 탐색함
    def find_arrays(v,path='',depth=0):
        if isinstance(v,list):
            candidates.append({'path':path or '/', 'count':len(v), 'item_types':sorted({kind(x) for x in v})})
        elif isinstance(v,dict):
            for k,x in v.items(): find_arrays(x,path+'/'+k.replace('~','~0').replace('/','~1').replace('*','~2'),depth+1)
    find_arrays(data)
    leaf=[f for f in fields.values() if not set(f['types']) <= {'object','array'}]
    cells=sum(f['scalar_occurrences'] for f in leaf); missing=sum(f['null_count']+f['empty_count'] for f in leaf)
    quality=[{'category':'COMPLETENESS','status':'MEASURED','scope':'observed_leaf_values','missing':missing,'observed':cells,'score':round(100*(cells-missing)/cells,2) if cells else None}]
    quality += [{'category':k,'status':'REVIEW_REQUIRED','score':None} for k in ['VALIDITY','CONSISTENCY','UNIQUENESS','TIMELINESS','ACCURACY']]
    field_list = list(fields.values())
    processing = _processing_candidates(field_list, category)
    return {
        'schema_version': '2.0',
        'format': fmt,
        'data_category': category,
        'sha256': hashlib.sha256(raw).hexdigest(),
        'byte_size': len(raw),
        'root_type': kind(data),
        'traits': sorted(traits),
        'namespaces': namespaces,
        'namespace_bindings': [{'prefix': p, 'uri': u} for p, u in sorted(namespace_bindings)],
        'fields': field_list,
        'record_sets': candidates,
        'tables': tables,
        'quality_metrics': quality,
        'processing': processing,
        'warnings': warnings,
        'review_required': ['소관기관', '라이선스', '갱신주기', '개인정보 및 편향 검토']
        + (['endpoint', 'HTTP method', '인증', '요청 파라미터', '오류 코드', '페이지네이션']
           if category == 'api' else ['단위', '코드 사전', '결합키', '보간 및 가공 이력']),
    }


# 분석 모델과 제목을 기반으로 AI 친화 가이드 마크다운을 렌더링함
def render_markdown(model, title):
    category='API' if model['data_category']=='api' else '파일데이터'
    lines=[f'# AI 친화 가이드 — {title}',f'\n## AI 친화 {category} 가이드',f"형식: {model['format']} / 루트 유형: {model['root_type']} / SHA-256: {model['sha256']}",'\n## 데이터셋 구조', '원본 구조의 경로와 관측 타입을 보존합니다. 배열 /*는 관측 항목의 합집합이며 필수 필드를 의미하지 않습니다.', '```json',dumps({'traits':model['traits'],'namespaces':model['namespaces'],'record_sets':model['record_sets'],'tables':model['tables']}),'```','\n## 데이터 사전','| 경로 | 관측 타입 | 출현 | null | 빈 문자열 | 누락 |','|---|---|---:|---:|---:|---:|']
    for f in model['fields']:
        path=f['path'].replace('|','\\|').replace('\n','\\n')
        lines.append(f"| {path} | {', '.join(f['types'])} | {f['occurrences']} | {f['null_count']} | {f['empty_count']} | {f['missing_count'] if f['missing_count'] is not None else '해당 없음'} |")
    lines+=['\n## 품질 진단','입력값 채움률은 업로드 데이터에서 null·빈 문자열이 아닌 관측값의 비율입니다. 모집단 대표성, 정확성, 필수 필드 충족 여부를 보증하지 않습니다.','```json',dumps(model['quality_metrics']),'```','\n## 기관 확인 필요 (REVIEW_REQUIRED)']+['- '+x for x in model['review_required']]
    lines+=['\n## 파싱 및 활용 지침','- null, 빈 문자열, 누락, 0, false를 구분하고 원천 값을 임의 보정하지 않습니다.','- 구조 분석만으로 OpenAPI/XSD 적합성 또는 AI 학습 적합성을 확정하지 않습니다.','- 학습/검증 분리, 편향, 저작권, 개인정보, 원천 및 가공 이력을 검토합니다.']+['- '+x for x in model['warnings']]
    if category=='API':
        lines += ['\n## API 서비스 및 호출 계약',
                  '| 항목 | 값 | 상태 |', '|---|---|---|',
                  '| 서비스 URL / HTTP method | 제공되지 않음 | REVIEW_REQUIRED |',
                  '| 인증 방식 / 요청 파라미터 | 제공되지 않음 | REVIEW_REQUIRED |',
                  '| 페이지네이션 / 호출 제한 | 제공되지 않음 | REVIEW_REQUIRED |',
                  '| 오류 코드 / 재시도 정책 | 제공되지 않음 | REVIEW_REQUIRED |',
                  '\n## 응답 파싱 절차',
                  '1. 응답 형식과 문자 인코딩을 확인한 뒤 전체 문서 구문을 검증합니다.',
                  '2. 위 구조 사전의 레코드셋 후보를 확인하고 실제 업무 레코드 경로를 선택합니다.',
                  '3. XML namespace와 속성, JSON-LD context와 그래프 관계를 유지합니다.',
                  '4. 선택 필드 누락·null·빈 배열·오류 응답을 별도로 처리합니다.',
                  '- 응답 샘플은 API 계약이 아닙니다. 요청 예제는 제공기관 계약 확인 후 작성합니다.']
    else:
        lines += ['\n## 파일 적재 및 가공 이력',
                  '- XLSX 시트별 헤더·단위·수식·결합키를 확인하고 CSV 인코딩 및 구분자를 문서화합니다.',
                  '- 결합 방법, 보간 여부, 단위 변환, 원천과 파생 필드 구분: REVIEW_REQUIRED.',
                  '- 파일 해시와 버전을 보관하고 원천 데이터 및 보정 여부를 함께 제공합니다.',
                  '\n## AI 활용 시나리오 검토',
                  '- 예측·분류·이상 탐지의 목적과 목표 변수를 먼저 정의합니다.',
                  '- 시계열·기관·개체별 분리 기준을 확인해 학습/검증 간 정보 누출을 방지합니다.']
    return '\n'.join(lines)


# 분석 모델을 DCAT 표준 준수 JSON-LD 포맷으로 직렬화함
def render_jsonld(model,title):
    return dumps({'@context':{'dcat':'http://www.w3.org/ns/dcat#','dct':'http://purl.org/dc/terms/'},'@id':'urn:sha256:'+model['sha256'],'@type':'dcat:Dataset','dct:title':title,'dcat:distribution':{'@type':'dcat:Distribution','dct:format':model['format']}})


# 분석 모델을 정형화된 내부 XML 포맷으로 변환함
def render_xml(model):
    root=etree.Element('aiReadyDataset',schemaVersion='2.0')
    # 내부 요소 노드를 재귀적으로 생성하여 XML 트리에 추가함
    def add(parent,key,value):
        node=etree.SubElement(parent,'entry',keyJson=dumps(str(key)),type=kind(value))
        if isinstance(value,dict):
            for k,v in value.items(): add(node,k,v)
        elif isinstance(value,list):
            for i,v in enumerate(value): add(node,i,v)
        elif value is not None: node.text=dumps(value)
    for key,value in model.items(): add(root,key,value)
    return etree.tostring(root,encoding='unicode',pretty_print=True)
