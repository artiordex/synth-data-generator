# -*- coding: utf-8 -*-
# =============================================================================
# 파일명: profile.py
# 경로: apps/api/src/synthetic_api/application/services/ai_guide_document/profile.py
# 목적: 원본 데이터셋의 한계 내 프로파일링 및 통계 분석을 수행함
# 작성자: 개발팀
# 작성일: 2026-09-16
# 수정일: 2026-09-16
# =============================================================================
"""Bounded, deterministic source profiling for every guide output."""
from __future__ import annotations

import hashlib
import io
import json
import math
import re
from collections import Counter
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook

from ..ai_guide_analysis import analyze, kind


# 날짜/시간 객체 또는 문자열을 ISO 포맷 타임스탬프로 변환함
def timestamp(value):
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str) and re.match(r'^\d{4}[-/]\d\d[-/]\d\d', value):
        try:
            return datetime.fromisoformat(value.replace('/', '-')).isoformat()
        except ValueError:
            pass
    return None


class Observations:
    # 관측 통계 객체를 초기화함
    def __init__(self):
        self.types = Counter()
        self.count = self.null = self.empty = self.zero = self.false = 0
        self.numeric = 0
        self.mean = 0.0
        self.minimum = self.maximum = None
        self.formula = self.uncached = 0
        self.times = Counter()

    # 단일 관측값 또는 수식을 통계에 반영함
    def add(self, value, formula=False, cached=None):
        self.count += 1
        self.types['formula' if formula else kind(value)] += 1
        if formula:
            self.formula += 1
            self.uncached += cached is None
            return
        self.null += value is None
        self.empty += isinstance(value, str) and value == ''
        self.false += value is False
        self.zero += not isinstance(value, bool) and isinstance(value, (int, float)) and value == 0
        if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
            self.numeric += 1
            self.mean += (value - self.mean) / self.numeric
            self.minimum = value if self.minimum is None else min(value, self.minimum)
            self.maximum = value if self.maximum is None else max(value, self.maximum)
        stamp = timestamp(value)
        if stamp:
            self.times[stamp] += 1

    # 수집된 관측치들을 집계하여 최종 통계 사전을 반환함
    def finish(self):
        temporal = None
        if self.times:
            stamps = sorted(self.times)
            deltas = Counter((datetime.fromisoformat(b) - datetime.fromisoformat(a)).total_seconds()
                             for a, b in zip(stamps, stamps[1:]))
            interval = deltas.most_common(1)[0][0] if deltas else None
            gaps = [{'start': a, 'end': b, 'seconds': (datetime.fromisoformat(b)-datetime.fromisoformat(a)).total_seconds()}
                    for a, b in zip(stamps, stamps[1:])
                    if interval and (datetime.fromisoformat(b)-datetime.fromisoformat(a)).total_seconds() > interval]
            temporal = dict(start=stamps[0], end=stamps[-1], valid=sum(self.times.values()),
                            unique=len(stamps), duplicate_timestamps=sum(self.times.values())-len(stamps),
                            modal_interval_seconds=interval, gaps=gaps,
                            rule='정렬한 고유 시각의 최빈 간격 초과 구간. 업무상 누락 여부·개체별 주기는 기관 확인 필요.')
        return dict(occurrences=self.count, types=sorted(self.types), null_count=self.null,
                    empty_count=self.empty, zero_count=self.zero, false_count=self.false,
                    numeric_count=self.numeric, numeric_missing_count=self.null+self.empty+self.uncached,
                    min_value=self.minimum, max_value=self.maximum,
                    mean=self.mean if self.numeric else None, formula_count=self.formula,
                    uncached_formula_count=self.uncached, temporal=temporal)


# 바이트 데이터를 파싱하여 통계 및 품질 메타데이터를 프로파일링함
def profile_bytes(raw: bytes, fmt: str, name: str) -> dict:
    fmt = fmt.lower().lstrip('.')
    if fmt != 'xlsx':
        try:
            text = raw.decode('utf-8-sig')
            encoding = 'UTF-8'
        except UnicodeDecodeError:
            if fmt not in {'csv', 'tsv'}:
                raise ValueError('UTF-8 입력이 필요합니다.')
            text = raw.decode('cp949')
            encoding = 'CP949'
        model = analyze(text, fmt)
        model.update(name=name, byte_size=len(raw), sha256=hashlib.sha256(raw).hexdigest(),
                     scope='FULL', encoding=encoding)
        for field in model['fields']:
            field.pop('examples', None)  # No source credentials or signed URLs in exports.
        model['declared_counts'] = []
        if fmt in {'json', 'jsonld', 'json-ld'}:
            value = json.loads(text)
            # JSON 노드에서 선언된 총 건수나 페이지 정보를 추출함
            def declared(node, path=''):
                if isinstance(node, dict):
                    for key, val in node.items():
                        if key.lower() in {'totalcount', 'numofrows', 'pageno'}:
                            model['declared_counts'].append({'path': path+'/'+key, 'value': val})
                        declared(val, path+'/'+key)
                elif isinstance(node, list):
                    return
            declared(value)
        elif fmt == 'xml':
            from lxml import etree
            root = etree.fromstring(raw, etree.XMLParser(resolve_entities=False, no_network=True))
            for node in root.iter():
                if isinstance(node.tag, str) and etree.QName(node).localname.lower() in {'totalcount', 'numofrows', 'pageno'}:
                    model['declared_counts'].append({'path': root.getroottree().getpath(node), 'value': node.text})
        model['observed_records'] = (sum(t['row_count'] for t in model['tables']) if model['data_category']=='file'
                                     else model['record_sets'][0]['count'] if len(model['record_sets']) == 1 else None)
        return model
    from zipfile import ZipFile
    with ZipFile(io.BytesIO(raw)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 512*1024*1024:
            raise ValueError('XLSX 압축 해제 크기 한도 512 MiB 초과')
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    cached_book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    fields, tables = [], []
    try:
        for sheet in workbook:
            rows = sheet.iter_rows()
            header_cells = next(rows, None)
            if not header_cells:
                tables.append(dict(name=sheet.title, row_count=0, columns=[]))
                continue
            headers = [str(c.value) if c.value is not None else '' for c in header_cells]
            if any(not h.strip() for h in headers) or len(set(headers)) != len(headers):
                raise ValueError(f'{sheet.title}: 빈 헤더 또는 중복 헤더')
            cached_rows = cached_book[sheet.title].iter_rows()
            next(cached_rows, None)
            stats = [Observations() for _ in headers]
            date_cols=[i for i,h in enumerate(headers) if re.search(r'일자|날짜|date',h,re.I)]
            hour_cols=[i for i,h in enumerate(headers) if re.search(r'^(작업)?시간$|(^|[ _])hour$',h,re.I)]
            date_index=date_cols[0] if len(date_cols)==1 else None
            hour_index=hour_cols[0] if len(hour_cols)==1 else None
            entity_indexes=[i for i,h in enumerate(headers) if re.search(r'(^id$|[ _]id$|이름$)',h,re.I)]
            time_stats=Observations(); entity_times={}; invalid_time=0
            count = blank = 0
            for row, cached in zip(rows, cached_rows):
                if all(c.value is None and c.data_type != 'inlineStr' for c in row):
                    blank += 1
                    continue
                count += 1
                for stat, cell, cache in zip(stats, row, cached):
                    value = '' if cell.value is None and cell.data_type == 'inlineStr' else cell.value
                    stat.add(value, cell.data_type == 'f', cache.value)
                if date_index is not None:
                    raw_date=row[date_index].value
                    try:
                        stamp=datetime.strptime(str(raw_date),'%Y%m%d') if re.fullmatch(r'\d{8}',str(raw_date)) else datetime.fromisoformat(str(raw_date))
                        if hour_index is not None:
                            hour=row[hour_index].value
                            if not isinstance(hour,int) or isinstance(hour,bool) or not 0<=hour<=23: raise ValueError('시간 범위')
                            stamp+=timedelta(hours=hour)
                        time_stats.add(stamp)
                        group=tuple(str(row[i].value) for i in entity_indexes)
                        entity_times.setdefault(group,Counter())[stamp.isoformat()]+=1
                    except (ValueError,TypeError): invalid_time+=1
            temporal=time_stats.finish()['temporal']
            if temporal:
                group_summaries=[]
                for key,counts in entity_times.items():
                    stat=Observations();stat.times=counts
                    result=stat.finish()['temporal']
                    group_summaries.append(dict(key=list(key),**result))
                temporal.update(date_field=headers[date_index],hour_field=headers[hour_index] if hour_index is not None else None,
                                invalid_count=invalid_time,entity_fields=[headers[i] for i in entity_indexes],groups=group_summaries,
                                interpretation_status='AUTO_INFERRED',parsing_rule='날짜 이름 후보 + YYYYMMDD/ISO 유효성 검사; 시간 이름 후보의 정수 0..23을 시각으로 결합. 개체 후보별 분리. 업무 시간대와 키는 미확정.')
            tables.append(dict(name=sheet.title, row_count=count, columns=headers, blank_rows=blank,temporal=temporal))
            for header, stat in zip(headers, stats):
                path = '/'+sheet.title.replace('~','~0').replace('/','~1')+'/*/'+header.replace('~','~0').replace('/','~1')
                fields.append(dict(path=path, name=header, sheet=sheet.title, missing_count=0, **stat.finish()))
    finally:
        workbook.close()
        cached_book.close()
    cells = sum(f['occurrences'] for f in fields)
    missing = sum(f['null_count']+f['empty_count'] for f in fields)
    return dict(name=name, format=fmt, data_category='file', byte_size=len(raw),
                sha256=hashlib.sha256(raw).hexdigest(), root_type='workbook', scope='FULL',
                encoding=None, tables=tables, fields=fields, record_sets=[], declared_counts=[],
                namespaces={}, traits=['tabular'], observed_records=sum(t['row_count'] for t in tables),
                quality_metrics=[dict(category='COMPLETENESS', score=100*(cells-missing)/cells if cells else None,
                                      missing=missing, observed=cells)],
                warnings=['업로드 파일 전수 조사이며 모집단 전체를 뜻하지 않습니다.',
                          '수식 원문과 캐시 유무를 구분하며 수식은 수치 통계에서 제외합니다.',
                          '날짜별 중복은 다중 개체일 수 있어 중복 레코드로 판정하지 않습니다.'])


# 동일 데이터셋의 JSON 및 XML 샘플 간 일치성을 비교함
def compare_json_xml(json_raw: bytes, xml_raw: bytes) -> dict:
    """Explicit sample correspondence; no namespace stripping or record union."""
    from lxml import etree
    left = json.loads(json_raw.decode('utf-8-sig'))
    root = etree.fromstring(xml_raw, etree.XMLParser(resolve_entities=False, no_network=True))
    # XML 노드를 JSON과 비교 가능한 딕셔너리로 변환함
    def convert(node):
        if node.attrib:
            return {'@attributes':dict(node.attrib), '#text':node.text}
        if not len(node):
            return node.text if node.text is not None else ''
        grouped = {}
        for child in node:
            grouped.setdefault(child.tag, []).append(convert(child))
        return {k:v if len(v)>1 else v[0] for k,v in grouped.items()}
    right = {root.tag:convert(root)}
    differences, mappings = [], []
    # 두 데이터 트리 간의 값, 타입, 구조적 차이를 재귀 비교함
    def compare(a,b,path='$'):
        if isinstance(a,list) and isinstance(b,dict) and len(b)==1:
            key, values = next(iter(b.items()))
            mappings.append(dict(path=path, rule=f'JSON 배열 ↔ XML {key} 반복 요소'))
            b=values if isinstance(values,list) else [values]
        if isinstance(a,dict) and isinstance(b,dict):
            for key in sorted(a.keys() | b.keys()):
                if key not in a or key not in b:
                    differences.append(dict(path=path+'/'+key, kind='missing_key'))
                else: compare(a[key],b[key],path+'/'+key)
        elif isinstance(a,list) and isinstance(b,list):
            if len(a)!=len(b): differences.append(dict(path=path,kind='record_count',json=len(a),xml=len(b)))
            for i,(x,y) in enumerate(zip(a,b)): compare(x,y,path+f'/{i}')
        else:
            expected = str(a).lower() if isinstance(a,bool) else str(a) if isinstance(a,(int,float)) else a
            if expected != b:
                differences.append(dict(path=path,kind='value_or_type'))
    compare(left,right)
    return dict(equivalent=not differences, differences=differences, mappings=mappings,
                rule='QName 유지; 단일 자식 반복 컨테이너를 배열에 대응; 레코드 순서·키·모든 값 비교. JSON 숫자/boolean만 XML 정규 문자열로 대응. null과 빈 문자열은 다름.',
                scope='제공된 두 파일의 값과 envelope 전수 비교. 운영 API 또는 전체 모집단 동일성은 미확정.')
