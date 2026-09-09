import zipfile

from hwpx.document import HwpxDocument

from synthetic_api.routes.v1 import converter


def _make_converter_pdf(path):
    import pymupdf as fitz

    with fitz.open() as doc:
        page = doc.new_page(width=420, height=320)
        page.draw_rect(fitz.Rect(30, 30, 390, 85), color=(0.1, 0.2, 0.4), fill=(0.1, 0.2, 0.4))
        page.insert_text((45, 62), "Document Title", fontsize=15, color=(1, 1, 1))
        page.insert_text((45, 125), "Structured paragraph with value 123", fontsize=11)
        page.draw_rect(fitz.Rect(45, 155, 375, 230), color=(0.3, 0.3, 0.3))
        page.insert_text((58, 178), "Name", fontsize=10)
        page.insert_text((190, 178), "Amount", fontsize=10)
        page.insert_text((58, 208), "Alpha", fontsize=10)
        page.insert_text((190, 208), "100", fontsize=10)
        doc.save(path)


def _make_ordered_docx(path):
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '<Override PartName="/word/styles.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
        '</Types>'
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/>'
        '</Relationships>'
    )
    doc_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId5" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
        'Target="https://example.com" TargetMode="External"/>'
        '</Relationships>'
    )
    styles = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:style w:type="paragraph" w:styleId="Normal"><w:name w:val="Normal"/></w:style>'
        '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/></w:style>'
        '</w:styles>'
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<w:body>'
        '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Title</w:t></w:r></w:p>'
        '<w:p><w:r><w:t xml:space="preserve">Before table </w:t></w:r>'
        '<w:hyperlink r:id="rId5"><w:r><w:t>link</w:t></w:r></w:hyperlink></w:p>'
        '<w:tbl>'
        '<w:tr><w:tc><w:p><w:r><w:t>Name</w:t></w:r></w:p></w:tc>'
        '<w:tc><w:p><w:r><w:t>Value</w:t></w:r></w:p></w:tc></w:tr>'
        '<w:tr><w:tc><w:p><w:r><w:t>Alpha</w:t></w:r></w:p></w:tc>'
        '<w:tc><w:p><w:r><w:t>100</w:t></w:r></w:p></w:tc></w:tr>'
        '</w:tbl>'
        '<w:p><w:r><w:t>After table</w:t></w:r></w:p>'
        '<w:sectPr/>'
        '</w:body></w:document>'
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/_rels/document.xml.rels", doc_rels)
        archive.writestr("word/styles.xml", styles)
        archive.writestr("word/document.xml", document)


def test_pure_hwp_to_hwpx_writes_valid_package(tmp_path, monkeypatch):
    source = tmp_path / "source.hwp"
    source.write_bytes(b"placeholder")
    target = tmp_path / "target.hwpx"
    monkeypatch.setattr(converter, "_extract_hwp_paragraphs_pure", lambda _: ["첫 문단", "둘째 문단"])

    converter._convert_hwp_to_hwpx_pure(source, target)

    doc = HwpxDocument.open(target)
    assert doc.validate().ok
    text = doc.text.plain()
    assert "첫 문단" in text
    assert "둘째 문단" in text


def test_hwpx_validation_rejects_broken_zip(tmp_path):
    target = tmp_path / "broken.hwpx"
    target.write_text("not a zip", encoding="utf-8")

    import pytest

    with pytest.raises(RuntimeError, match="ZIP 패키지"):
        converter._validate_hwpx_output(target)


def test_convert_hwp_to_docx_with_tables(tmp_path):
    import docx
    from synthetic_engine.exporters.hwp_high_fidelity_docx_converter import convert_any_hwp_to_docx
    from pathlib import Path

    sample = Path("storage/uploads/0f7776311d2c4361a4eb2b1dfeb1b32f_참고.hwp")
    if sample.exists():
        target = tmp_path / "참고.docx"
        out_path = convert_any_hwp_to_docx(sample, target)
        assert out_path.exists()
        assert out_path.stat().st_size > 1000
        doc = docx.Document(out_path)
        assert len(doc.tables) >= 1
        t = doc.tables[0]
        assert len(t.rows) >= 5
        assert len(t.columns) == 7


def test_convert_hwp_to_high_fidelity_hwpx_with_tables(tmp_path):
    import zipfile
    from pathlib import Path
    from synthetic_engine.exporters.hwp_high_fidelity_hwpx_converter import convert_any_hwp_to_hwpx
    from hwpx.document import HwpxDocument

    sample = Path("storage/uploads/0f7776311d2c4361a4eb2b1dfeb1b32f_참고.hwp")
    if sample.exists():
        target = tmp_path / "참고.hwpx"
        out_path = convert_any_hwp_to_hwpx(sample, target)
        assert out_path.exists()
        assert out_path.stat().st_size > 1000
        assert zipfile.is_zipfile(out_path)

        doc = HwpxDocument.open(out_path)
        assert doc.validate().ok
        plain = doc.text.plain()
        assert "기본" in plain or "식품" in plain or "번호" in plain or "1" in plain


def test_convert_api_hwp_to_hwpx_endpoint():
    from fastapi.testclient import TestClient
    from synthetic_api.main import app
    from pathlib import Path

    client = TestClient(app)
    sample = Path("storage/uploads/0f7776311d2c4361a4eb2b1dfeb1b32f_참고.hwp")
    if sample.exists():
        res = client.post(
            "/api/v1/converter/convert",
            files={"file": ("참고.hwp", sample.read_bytes(), "application/octet-stream")},
            data={"target_format": "hwpx"}
        )
        assert res.status_code == 200
        data = res.json()
        assert data["status"] == "success"
        assert data["target_format"] == "HWPX"
        assert data["download_url"] is not None
        assert data.get("html_preview") is not None
        assert data.get("markdown_preview") is not None


def test_structured_pdf_parse_returns_preview_and_metadata(tmp_path):
    source = tmp_path / "layout.pdf"
    _make_converter_pdf(source)

    parsed = converter._build_structured_document_parse(source, ".pdf", "layout.pdf")

    assert "Document Title" in parsed["markdown"]
    assert parsed["html_preview"]
    structure = parsed["structure"]
    assert structure["format"] == "PDF"
    assert structure["pages_count"] == 1
    assert structure["block_count"] >= 1
    assert "PyMuPDF visual layout" in structure["parser_engines"]


def test_structured_hwpx_parse_returns_blocks_and_markdown(tmp_path):
    source = tmp_path / "sample.hwpx"
    doc = HwpxDocument.new()
    doc.add_paragraph("한글 표준 문서 제목")
    doc.add_paragraph("본문 구조화 파싱 테스트")
    doc.save_to_path(source)

    parsed = converter._build_structured_document_parse(source, ".hwpx", "sample.hwpx")

    assert "본문 구조화 파싱 테스트" in parsed["markdown"]
    assert parsed["html_preview"]
    assert parsed["structure"]["format"] == "HWPX"
    assert parsed["structure"]["block_count"] >= 1


def test_convert_pdf_to_markdown_endpoint_includes_document_structure(tmp_path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from synthetic_api.core.config import settings

    for attr in ("UPLOAD_DIR", "OUTPUT_DIR"):
        folder = tmp_path / attr
        folder.mkdir()
        monkeypatch.setattr(settings, attr, folder)

    app = FastAPI()
    app.include_router(converter.router, prefix="/api/v1")
    client = TestClient(app)

    source = tmp_path / "layout.pdf"
    _make_converter_pdf(source)
    res = client.post(
        "/api/v1/converter/convert",
        files={"file": ("layout.pdf", source.read_bytes(), "application/pdf")},
        data={"target_format": "md"},
    )

    assert res.status_code == 200, res.text
    data = res.json()
    assert data["target_format"] == "MD"
    assert "Document Title" in data["markdown_preview"]
    assert data["document_structure"]["format"] == "PDF"
    assert data["document_structure"]["block_count"] >= 1


def test_convert_docx_to_markdown_preserves_block_order_and_links(tmp_path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from synthetic_api.core.config import settings

    for attr in ("UPLOAD_DIR", "OUTPUT_DIR"):
        folder = tmp_path / attr
        folder.mkdir()
        monkeypatch.setattr(settings, attr, folder)

    app = FastAPI()
    app.include_router(converter.router, prefix="/api/v1")
    client = TestClient(app)

    source = tmp_path / "ordered.docx"
    _make_ordered_docx(source)
    res = client.post(
        "/api/v1/converter/convert",
        files={"file": ("ordered.docx", source.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"target_format": "md"},
    )

    assert res.status_code == 200, res.text
    data = res.json()
    md = data["markdown_preview"]
    assert data["document_structure"]["format"] == "DOCX"
    assert "[link](https://example.com)" in md
    assert md.index("Before table") < md.index("| Name | Value |") < md.index("After table")


def test_convert_docx_to_html_and_hwpx_outputs_structured_documents(tmp_path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from synthetic_api.core.config import settings

    for attr in ("UPLOAD_DIR", "OUTPUT_DIR"):
        folder = tmp_path / attr
        folder.mkdir()
        monkeypatch.setattr(settings, attr, folder)

    app = FastAPI()
    app.include_router(converter.router, prefix="/api/v1")
    client = TestClient(app)

    source = tmp_path / "ordered.docx"
    _make_ordered_docx(source)

    html_res = client.post(
        "/api/v1/converter/convert",
        files={"file": ("ordered.docx", source.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"target_format": "html"},
    )
    assert html_res.status_code == 200, html_res.text
    html = html_res.json()["html_preview"]
    assert '<a href="https://example.com">link</a>' in html
    assert html.index("Before table") < html.index("<table>") < html.index("After table")

    hwpx_res = client.post(
        "/api/v1/converter/convert",
        files={"file": ("ordered.docx", source.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"target_format": "hwpx"},
    )
    assert hwpx_res.status_code == 200, hwpx_res.text
    out_path = settings.OUTPUT_DIR / "converted" / hwpx_res.json()["file_name"]
    assert zipfile.is_zipfile(out_path)


def test_convert_pdf_to_hwpx_endpoint_writes_valid_package(tmp_path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from synthetic_api.core.config import settings

    for attr in ("UPLOAD_DIR", "OUTPUT_DIR"):
        folder = tmp_path / attr
        folder.mkdir()
        monkeypatch.setattr(settings, attr, folder)

    app = FastAPI()
    app.include_router(converter.router, prefix="/api/v1")
    client = TestClient(app)

    source = tmp_path / "layout.pdf"
    _make_converter_pdf(source)
    res = client.post(
        "/api/v1/converter/convert",
        files={"file": ("layout.pdf", source.read_bytes(), "application/pdf")},
        data={"target_format": "hwpx"},
    )

    assert res.status_code == 200, res.text
    data = res.json()
    out_path = settings.OUTPUT_DIR / "converted" / data["file_name"]
    assert data["target_format"] == "HWPX"
    assert zipfile.is_zipfile(out_path)
    assert data["document_structure"]["format"] == "PDF"


def test_convert_pdf_to_docx_endpoint_writes_word_file(tmp_path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from synthetic_api.core.config import settings

    for attr in ("UPLOAD_DIR", "OUTPUT_DIR"):
        folder = tmp_path / attr
        folder.mkdir()
        monkeypatch.setattr(settings, attr, folder)

    app = FastAPI()
    app.include_router(converter.router, prefix="/api/v1")
    client = TestClient(app)

    source = tmp_path / "layout.pdf"
    _make_converter_pdf(source)
    res = client.post(
        "/api/v1/converter/convert",
        files={"file": ("layout.pdf", source.read_bytes(), "application/pdf")},
        data={"target_format": "docx"},
    )

    assert res.status_code == 200, res.text
    data = res.json()
    out_path = settings.OUTPUT_DIR / "converted" / data["file_name"]
    assert data["target_format"] == "DOCX"
    assert out_path.suffix == ".docx"
    assert zipfile.is_zipfile(out_path)
    assert data["document_structure"]["format"] == "PDF"


def test_convert_pdf_to_hwp_endpoint_routes_to_hwp_exporter(tmp_path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from synthetic_api.core.config import settings
    import synthetic_engine.exporters.pdf_high_fidelity_converter as pdf_exporter

    for attr in ("UPLOAD_DIR", "OUTPUT_DIR"):
        folder = tmp_path / attr
        folder.mkdir()
        monkeypatch.setattr(settings, attr, folder)

    def fake_pdf_to_hwp(source, output):
        output.write_bytes(b"hwp-binary")

    monkeypatch.setattr(pdf_exporter, "convert_pdf_to_high_fidelity_hwp", fake_pdf_to_hwp)

    app = FastAPI()
    app.include_router(converter.router, prefix="/api/v1")
    client = TestClient(app)

    source = tmp_path / "layout.pdf"
    _make_converter_pdf(source)
    res = client.post(
        "/api/v1/converter/convert",
        files={"file": ("layout.pdf", source.read_bytes(), "application/pdf")},
        data={"target_format": "hwp"},
    )

    assert res.status_code == 200, res.text
    data = res.json()
    out_path = settings.OUTPUT_DIR / "converted" / data["file_name"]
    assert data["target_format"] == "HWP"
    assert out_path.suffix == ".hwp"
    assert out_path.read_bytes() == b"hwp-binary"
    assert data["document_structure"]["format"] == "PDF"


def test_convert_hwpx_to_markdown_endpoint_includes_document_structure(tmp_path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from synthetic_api.core.config import settings

    for attr in ("UPLOAD_DIR", "OUTPUT_DIR"):
        folder = tmp_path / attr
        folder.mkdir()
        monkeypatch.setattr(settings, attr, folder)

    app = FastAPI()
    app.include_router(converter.router, prefix="/api/v1")
    client = TestClient(app)

    source = tmp_path / "sample.hwpx"
    doc = HwpxDocument.new()
    doc.add_paragraph("HWPX 구조화 문서")
    doc.add_paragraph("마크다운 변환 확인")
    doc.save_to_path(source)

    res = client.post(
        "/api/v1/converter/convert",
        files={"file": ("sample.hwpx", source.read_bytes(), "application/octet-stream")},
        data={"target_format": "md"},
    )

    assert res.status_code == 200, res.text
    data = res.json()
    assert data["target_format"] == "MD"
    assert "마크다운 변환 확인" in data["markdown_preview"]
    assert data["document_structure"]["format"] == "HWPX"


def test_structured_preview_is_not_truncated_for_large_document(tmp_path):
    source = tmp_path / "large.md"
    filler = "\n\n".join(f"문서 본문 {idx:04d} " + ("가" * 120) for idx in range(2300))
    source.write_text(f"# 전체 미리보기\n\n{filler}\n\n끝 페이지 확인", encoding="utf-8")

    parsed = converter._build_structured_document_parse(source, ".md", "large.md")

    assert len(parsed["html_preview"]) > 250_000
    assert "끝 페이지 확인" in parsed["html_preview"]
