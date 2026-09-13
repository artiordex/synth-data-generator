# -*- coding: utf-8 -*-
# =============================================================================
# 파일명: test_pseudonym_formats.py
# 경로: apps/api/tests/test_pseudonym_formats.py
# 목적: 다양한 서식의 개인정보 가명화 결과 무결성을 검증함
# 작성자: 개발팀
# 작성일: 2026-09-09
# 수정일: 2026-09-13
# =============================================================================
from pathlib import Path

import pandas as pd
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from synthetic_api.core.config import settings
from synthetic_api.routes.v1.datasets import router
from synthetic_engine.profiling.pseudonym_input import read_pseudonym_input


# client 작업을 수행함
@pytest.fixture
def client(tmp_path, monkeypatch):
    for attr in ('UPLOAD_DIR', 'OUTPUT_DIR'):
        folder = tmp_path / attr
        folder.mkdir()
        monkeypatch.setattr(settings, attr, folder)
    app = FastAPI()
    app.include_router(router)
    return TestClient(app)


# tabular upload and real output 기능의 정상 동작 및 제약조건을 테스트함
@pytest.mark.parametrize('fmt', ['csv', 'xlsx', 'tsv', 'json', 'parquet'])
def test_tabular_upload_and_real_output(client, fmt):
    frame = pd.DataFrame({'email': ['sample@example.org'], 'category': ['A']})
    source = settings.UPLOAD_DIR / f'input.{fmt}'
    if fmt == 'xlsx': frame.to_excel(source, index=False)
    elif fmt == 'json': frame.to_json(source, orient='records')
    elif fmt == 'parquet': frame.to_parquet(source, index=False)
    else: frame.to_csv(source, sep='\t' if fmt == 'tsv' else ',', index=False)
    assert client.get('/datasets/profile', params={'file_name': source.name, 'pseudonym': True}).status_code == 200
    response = client.post('/datasets/pseudonymize', json={'file_name': source.name,
        'pii_actions': {'email': 'mask'}, 'export_format': fmt})
    assert response.status_code == 200, response.text
    output = settings.OUTPUT_DIR / 'pseudonymized' / response.json()['file_name']
    result = read_pseudonym_input(output)
    assert result.iloc[0]['email'] != 'sample@example.org'
    assert result.iloc[0]['category'] == 'A'


# document input and export 기능의 정상 동작 및 제약조건을 테스트함
@pytest.mark.parametrize('fmt', ['md', 'docx', 'hwpx', 'pdf'])
def test_document_input_and_export(client, fmt):
    source = settings.UPLOAD_DIR / f'input.{fmt}'
    if fmt == 'md': source.write_text('Before sample@example.org\n\n| label | value |\n| --- | --- |\n| item | example |\n\nAfter', encoding='utf-8')
    elif fmt == 'docx':
        from docx import Document
        doc = Document()
        doc.add_paragraph('Before sample@example.org')
        doc.add_table(rows=2, cols=1).cell(1, 0).text = 'Table content'
        doc.add_paragraph('After')
        doc.save(source)
    elif fmt == 'hwpx':
        from hwpx.document import HwpxDocument
        doc = HwpxDocument.new()
        doc.add_paragraph('Before sample@example.org')
        doc.add_paragraph('After')
        doc.save_to_path(source)
    else:
        from synthetic_engine.exporters.document_exporter import export_pseudonymized_document
        export_pseudonymized_document(pd.DataFrame({'문서_내용': ['Before sample@example.org', 'After']}), 'pdf', source)
    profile = client.get('/datasets/profile', params={'file_name': source.name, 'pseudonym': True})
    assert profile.status_code == 200, profile.text
    assert [c['name'] for c in profile.json()['columns']] == ['문단번호', '문서_내용']
    text = ' '.join(row['문서_내용'] for row in profile.json()['preview'])
    assert 'Before' in text and 'After' in text
    response = client.post('/datasets/pseudonymize', json={'file_name': source.name,
        'pii_actions': {'문서_내용': 'mask'}, 'export_format': fmt})
    if fmt in {'pdf', 'hwpx'}:
        assert response.status_code == 422
        assert '원본 서식 유지' in response.json()['detail']
        return
    assert response.status_code == 200, response.text
    output = settings.OUTPUT_DIR / 'pseudonymized' / response.json()['file_name']
    text = ' '.join(read_pseudonym_input(output)['문서_내용'].astype(str))
    assert 'sample@example.org' not in text
    assert 'After' in text


# export error does not return disguised csv 기능의 정상 동작 및 제약조건을 테스트함
def test_export_error_does_not_return_disguised_csv(client, monkeypatch):
    (settings.UPLOAD_DIR / 'input.csv').write_text('email\nsample@example.org', encoding='utf-8')
    # failed 작업을 수행함
    def failed(*args, **kwargs): raise RuntimeError('export unavailable')
    monkeypatch.setattr('synthetic_api.routes.v1.datasets.export_pseudonymized_document', failed)
    response = client.post('/datasets/pseudonymize', json={'file_name': 'input.csv',
        'pii_actions': {'email': 'mask'}, 'export_format': 'pdf'})
    assert response.status_code == 422
    assert not list((settings.OUTPUT_DIR / 'pseudonymized').glob('*.pdf'))


# unsupported format is not silent csv 기능의 정상 동작 및 제약조건을 테스트함
def test_unsupported_format_is_not_silent_csv(client):
    (settings.UPLOAD_DIR / 'input.csv').write_text('email\nsample@example.org', encoding='utf-8')
    response = client.post('/datasets/pseudonymize', json={'file_name': 'input.csv', 'export_format': 'hwp'})
    assert response.status_code == 422


# 한글(HWP) input requires native edit 기능의 정상 동작 및 제약조건을 테스트함
def test_hwp_input_requires_native_edit(client):
    template = Path(__file__).resolve().parents[3] / 'storage/templates/원본데이터 명세서.hwp'
    response = client.post('/datasets/upload', files={'file': ('review.hwp', template.read_bytes())})
    assert response.status_code == 200
    name = response.json()['filename']
    profile = client.get('/datasets/profile', params={'file_name': name, 'pseudonym': True})
    assert profile.status_code == 200, profile.text
    assert profile.json()['row_count'] > 0
    response = client.post('/datasets/pseudonymize', json={'file_name': name,
        'pii_actions': {'문서_내용': 'mask'}, 'export_format': 'hwpx'})
    assert response.status_code == 422
    assert '원본 서식 유지' in response.json()['detail']
