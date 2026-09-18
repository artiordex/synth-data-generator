# -*- coding: utf-8 -*-
# =============================================================================
# 파일명: test_public_data_conversion.py
# 경로: apps/api/tests/test_public_data_conversion.py
# 목적: 공공데이터 응답 구조, 영문명 추천·확인 및 원천 값 보존을 검증함
# 작성자: 개발팀
# 작성일: 2026-09-17
# 수정일: 2026-09-18
# =============================================================================
import io
import csv
import json
import xml.etree.ElementTree as ET
from datetime import datetime
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook

from synthetic_api.application.services import public_data_conversion as public
from synthetic_api.core.config import settings
from synthetic_api.routes.v1.converter import router


# 원천 문자열의 코드·빈 값·특수문자를 두 포맷에서 보존하는지 검증함
def test_csv_public_envelope_matches_example_structure():
    source=public.read_source('코드,설명\n001,"A & B < C"\n002,\n'.encode(),'table.csv')
    model=public.public_payload(source,{'코드':'code','설명':'description'})
    body=model['response']['body']
    assert body==dict(totalCount=2,pageNo=0,numOfRows=2,items=[
        {'code':'001','description':'A & B < C'},{'code':'002','description':''}])
    root=ET.fromstring(public.public_xml(model))
    assert [c.tag for c in root]==['header','body']
    assert [c.tag for c in root.find('header')]==['resultMsg','resultCode']
    assert [c.tag for c in root.find('body')]==['totalCount','pageNo','numOfRows','items']
    assert root.find('./body/items/item/code').text=='001'
    assert root.find('./body/items/item/description').text=='A & B < C'


# 엑셀 시트 선택·수식 원문·날짜·불리언·null과 0을 보존하는지 검증함
def test_excel_selected_sheet_native_values_and_xml_nil():
    book=Workbook();first=book.active;first.title='First';first.append(['unused']);first.append([1])
    sheet=book.create_sheet('Second');sheet.append(['코드','수치','상태','날짜','빈값','수식'])
    sheet.append(['001',0,False,datetime(2026,1,2),None,'=1+2'])
    stream=io.BytesIO();book.save(stream)
    source=public.read_source(stream.getvalue(),'table.xlsx',sheet_name='Second')
    assert source['sheets']==['First','Second'] and source['sheet_name']=='Second'
    names=dict(zip(source['headers'],['code','value','flag','date','empty','formula']))
    model=public.public_payload(source,names)
    assert model['response']['body']['items'][0]==dict(code='001',value=0,flag=False,
        date='2026-01-02T00:00:00',empty=None,formula='=1+2')
    row=ET.fromstring(public.public_xml(model)).find('./body/items/item')
    assert row.find('flag').text=='false' and row.find('value').text=='0'
    assert row.find('empty').attrib['{'+public.XSI+'}nil']=='true'
    with pytest.raises(ValueError,match='시트'):
        public.read_source(stream.getvalue(),'table.xlsx',sheet_name='Missing')


@pytest.mark.parametrize('names',[
    {'코드':'code'}, {'코드':'same','설명':'Same'}, {'코드':'한글','설명':'description'},
    {'코드':'1code','설명':'description'}, {'코드':'xmlField','설명':'description'},
    {'코드':'code','설명':'description','없는컬럼':'extra'}, [],
])
# 불완전하거나 부적합한 컬럼명 매핑을 거부하는지 검증함
def test_invalid_or_incomplete_mapping_rejected(names):
    with pytest.raises(ValueError): public.validate_names(['코드','설명'],names)


@pytest.mark.parametrize('raw',[b'a,a\n1,2',b'a,\n1,2',b'a,b\n1,2,3'])
# 중복·빈 헤더와 비정형 CSV 행을 거부하는지 검증함
def test_duplicate_blank_headers_and_nonrectangular_csv_rejected(raw):
    with pytest.raises(ValueError): public.read_source(raw,'table.csv')


# AI 미사용·실패 상태를 AUTO_INFERRED로 위장하지 않도록 검증함
def test_missing_key_returns_manual_draft_without_false_ai_claim(monkeypatch):
    monkeypatch.setattr(settings,'OPENAI_API_KEY','')
    monkeypatch.delenv('OPENAI_API_KEY',raising=False)
    result=public.recommend_names(['이름','field1','field1_1'])
    assert not result['ai_powered'] and result['model'] is None
    assert result['fields'][0]['status']=='REVIEW_REQUIRED'
    assert len({f['english_name'] for f in result['fields']})==3
    assert result['warning']


# 외부 모델에는 컬럼명만 보내며 결과를 엄격히 검증하는지 확인함
@pytest.mark.parametrize('bad',[False,True])
def test_gpt_mini_structured_suggestions_and_invalid_ai_fallback(monkeypatch,bad):
    monkeypatch.setattr(settings,'OPENAI_API_KEY','test-key')
    requests=[]
    # 외부 추천 API 요청을 검증하는 가짜 응답을 생성함
    def fake_open(request,timeout):
        body=json.loads(request.data);requests.append(body)
        assert timeout==20 and body['model']=='gpt-4o-mini'
        assert body['response_format']['json_schema']['strict'] is True
        suggestions={'fields':[{'index':0,'english_name':'한글' if bad else 'regionName','reason':'지역 이름 번역'}]}
        return io.BytesIO(json.dumps({'choices':[{'message':{'content':json.dumps(suggestions)}}]}).encode())
    monkeypatch.setattr(public.urllib.request,'urlopen',fake_open)
    result=public.recommend_names(['지역명'])
    user=json.loads(requests[0]['messages'][1]['content'])
    assert user==[{'index':0,'column_name':'지역명'}]
    assert result['ai_powered'] is (not bad)
    assert result['fields'][0]['status']==('REVIEW_REQUIRED' if bad else 'AUTO_INFERRED')
    assert result['fields'][0]['confidence'] is None


@pytest.fixture
# 격리된 저장소를 사용하는 테스트용 FastAPI 클라이언트를 생성함
def client(tmp_path,monkeypatch):
    for attr in ('UPLOAD_DIR','OUTPUT_DIR'):
        folder=tmp_path/attr;folder.mkdir();monkeypatch.setattr(settings,attr,folder)
    monkeypatch.setattr(settings,'OPENAI_API_KEY','')
    monkeypatch.delenv('OPENAI_API_KEY',raising=False)
    app=FastAPI();app.include_router(router,prefix='/api/v1')
    with TestClient(app) as connection: yield connection


# 사람이 수정·확인한 동일 매핑으로 JSON/XML을 생성하고 이력을 남기는지 검증함
@pytest.mark.parametrize('target',['json','xml'])
def test_api_human_edited_names_and_mapping_download(client,target):
    raw='코드,값\n001,0\n002,\n'.encode()
    recommendation=client.post('/api/v1/converter/field-names',files={'file':('data.csv',raw,'text/csv')})
    assert recommendation.status_code==200 and not recommendation.json()['ai_powered']
    names={'코드':'customCode','값':'customValue'}
    response=client.post('/api/v1/converter/convert',files={'file':('data.csv',raw,'text/csv')},data={
        'target_format':target,'field_names':json.dumps(names),'field_names_confirmed':'true'})
    assert response.status_code==200,response.text
    result=response.json()
    assert result['columns']==['customCode','customValue']
    assert result['preview']==[{'customCode':'001','customValue':'0'},{'customCode':'002','customValue':''}]
    assert result['dataset_profile']=='public_data'
    assert all(f['status']=='USER_CONFIRMED' for f in result['field_mappings'])
    output=settings.OUTPUT_DIR/'converted'/result['file_name']
    mapping=json.loads(output.with_name(output.name+'.field-mappings.json').read_text(encoding='utf-8'))
    assert [f['english_name'] for f in mapping['field_mappings']]==list(names.values())
    if target=='json':
        assert json.loads(output.read_text(encoding='utf-8'))['response']['body']['items']==result['preview']
    else:
        assert ET.parse(output).find('./body/items/item/customCode').text=='001'
    back=client.post('/api/v1/converter/convert',files={'file':('result.'+target,output.read_bytes(),'application/octet-stream')},data={'target_format':'csv'})
    assert back.status_code==200,back.text
    csv_text=(settings.OUTPUT_DIR/'converted'/back.json()['file_name']).read_text(encoding='utf-8-sig')
    assert 'customCode,customValue' in csv_text and '001,0' in csv_text


# 변환 확인 없이 출력하거나 가짜 한글 영문명을 생성하지 않는지 검증함
def test_api_requires_confirmation_and_does_not_export_fake_korean_names(client):
    raw='지역명\n청주\n'.encode()
    for data in ({'target_format':'json'},{'target_format':'json','field_names':json.dumps({'지역명':'regionName'})}):
        response=client.post('/api/v1/converter/convert',files={'file':('data.csv',raw,'text/csv')},data=data)
        assert response.status_code==400 and '확인' in response.json()['detail']
    malformed=client.post('/api/v1/converter/convert',files={'file':('data.csv',raw,'text/csv')},data={'target_format':'json','field_names':'{bad'})
    assert malformed.status_code==400


# 기본 공공데이터 프로파일과 records 호환 프로파일을 검증함
def test_api_ascii_headers_default_public_and_explicit_records_compatibility(client):
    for profile in ('auto','records'):
        response=client.post('/api/v1/converter/convert',files={'file':('data.csv',b'code\n001','text/csv')},
            data={'target_format':'json','dataset_profile':profile})
        assert response.status_code==200,response.text
        output=settings.OUTPUT_DIR/'converted'/response.json()['file_name']
        value=json.loads(output.read_text(encoding='utf-8'))
        assert isinstance(value,dict) if profile=='auto' else isinstance(value,list)


# XML 변환 시 허용되지 않는 제어문자를 거부하는지 검증함
def test_xml_rejects_invalid_control_characters_without_silent_data_loss():
    source={'headers':['field'],'rows':[['a\x01b']]}
    with pytest.raises(ValueError,match='제어문자'):
        public.public_xml(public.public_payload(source,{'field':'field'}))


# XLS 고유 셀 타입을 pandas 타입 추론 없이 변환하는지 검증함
def test_xls_native_boolean_number_and_empty_cells(monkeypatch):
    import xlrd
    released=[]
    rows=[
        [SimpleNamespace(ctype=xlrd.XL_CELL_TEXT,value=name) for name in ['코드','수치','상태','빈값']],
        [SimpleNamespace(ctype=xlrd.XL_CELL_TEXT,value='001'),
         SimpleNamespace(ctype=xlrd.XL_CELL_NUMBER,value=0.0),
         SimpleNamespace(ctype=xlrd.XL_CELL_BOOLEAN,value=0),
         SimpleNamespace(ctype=xlrd.XL_CELL_EMPTY,value='')],
    ]
    sheet=SimpleNamespace(nrows=2,row=lambda i: rows[i])
    book=SimpleNamespace(datemode=0,sheet_names=lambda:['Data'],
        sheet_by_name=lambda _:sheet,release_resources=lambda:released.append(True))
    monkeypatch.setattr(xlrd,'open_workbook',lambda **_:book)
    source=public.read_source(b'fixture','table.xls')
    assert source['rows']==[['001',0,False,None]]
    assert type(source['rows'][0][1]) is int and type(source['rows'][0][2]) is bool
    assert released==[True]


# 선택한 엑셀 시트를 API 추천·변환 단계에서 유지하는지 검증함
def test_excel_api_selected_sheet_is_used_for_recommendation_and_export(client):
    book=Workbook();book.active.title='First';book.active.append(['unused']);book.active.append([999])
    sheet=book.create_sheet('Second');sheet.append(['코드','상태']);sheet.append(['001',False])
    stream=io.BytesIO();book.save(stream);raw=stream.getvalue()
    recommendation=client.post('/api/v1/converter/field-names',files={'file':('book.xlsx',raw)},data={'sheet_name':'Second'})
    assert recommendation.status_code==200,recommendation.text
    assert recommendation.json()['sheets']==['First','Second']
    assert [f['original_name'] for f in recommendation.json()['fields']]==['코드','상태']
    response=client.post('/api/v1/converter/convert',files={'file':('book.xlsx',raw)},data={
        'target_format':'xml','sheet_name':'Second','field_names':json.dumps({'코드':'code','상태':'flag'}),
        'field_names_confirmed':'true'})
    assert response.status_code==200,response.text
    assert response.json()['sheet_name']=='Second'
    assert response.json()['preview']==[{'code':'001','flag':False}]


# 손상된 엑셀 입력을 클라이언트 오류로 반환하는지 검증함
def test_inspection_invalid_workbook_is_client_error(client):
    response=client.post('/api/v1/converter/field-names',files={'file':('broken.xlsx',b'not a zip')})
    assert response.status_code==400


# CP949 입력에서 한글과 코드 값을 보존하는지 검증함
def test_cp949_decoding_preserves_codes_and_na_tokens():
    raw='지역명,코드\n청주,001\nNA,NULL\n'.encode('cp949')
    source=public.read_source(raw,'table.csv',encoding='cp949')
    assert source['rows']==[['청주','001'],['NA','NULL']]


# 삭제·배포되지 않는 외부 샘플 대신 테스트 계약을 표현하는 대표 응답을 생성함
def _structured_public_model():
    items = [
        {'recordId':'PUB-001', 'name':'시청 앞', 'latitude':'36.6357', 'longitude':'127.4917'},
        {'recordId':'PUB-002', 'name':'터미널 앞', 'latitude':'36.6283', 'longitude':'127.4325'},
        {'recordId':'PUB-003', 'name':'공원 입구', 'latitude':'36.6421', 'longitude':'127.4890'},
    ]
    return {'response':{'header':{'resultCode':'00','resultMsg':'NORMAL_CODE(정상)'},
        'body':{'totalCount':len(items),'pageNo':0,'numOfRows':len(items),'items':items}}}


# 대표 공공데이터 응답의 JSON·XML 동등성과 왕복 변환을 검증함
def test_structured_public_model_is_equivalent_and_xml_json_roundtrip():
    original = _structured_public_model()
    source = public.read_structured_source(json.dumps(original, ensure_ascii=False).encode(), 'public-data.json')
    assert public.public_payload(source) == original
    body = original['response']['body']
    assert len(body['items']) == 3 and body['totalCount'] == 3
    assert body['pageNo'] == 0 and body['numOfRows'] == 3
    xml_source = public.read_structured_source(public.public_xml(original), 'public-data.xml')
    assert public.public_payload(xml_source) == original
    assert public.public_payload(public.read_structured_source(public.public_xml(original), 'back.xml')) == original


# 단일 행·누락 키·중첩 값·형식 부적합 키·기본 타입을 구분하여 보존함
@pytest.mark.parametrize('rows', [[], [{}], [
    {'code':'001', 'active':False, 'value':0, 'real':1.5, 'empty':'', 'null':None,
     'nested':{'x':' A & B ', 'none':None}, 'array':[1, '01', False, None, {'key':'value'}],
     'emptyObject':{}, 'emptyArray':[], '한글 필드':'내용', '1first':'number',
     '@id':'attribute-like', '#text':'text-like', '{urn:test}key':'namespaced-key'},
    {'code':'002'},
]])
# JSON·XML 왕복 변환에서 다양한 값과 누락 키를 보존하는지 검증함
def test_structured_json_xml_preserves_all_values_and_missing_keys(rows):
    source = public.read_structured_source(json.dumps(rows, ensure_ascii=False).encode(), 'data.json')
    model = public.public_payload(source)
    restored = public.public_payload(public.read_structured_source(public.public_xml(model), 'result.xml'))
    assert restored == model
    assert restored['response']['body']['items'] == rows


# 일반 XML의 반복 레코드·반복 필드·속성·네임스페이스를 손실 없이 읽음
def test_arbitrary_xml_repeated_fields_attributes_and_namespaces():
    raw = b'<dataset xmlns:n="urn:geo"><rows><row id="001"><code>001</code><tag>A</tag><tag>B</tag><n:point><n:x>12</n:x></n:point></row><row id="002"><code>002</code></row></rows></dataset>'
    source = public.read_structured_source(raw, 'data.xml')
    assert source['record_path'] == '/dataset/rows/row'
    rows = source['items']
    assert rows == [
        {'@id':'001','code':'001','tag':['A','B'],'{urn:geo}point':{'{urn:geo}x':'12'}},
        {'@id':'002','code':'002'},
    ]
    assert public.public_payload(public.read_structured_source(public.public_xml(public.public_payload(source)), 'back.xml'))['response']['body']['items'] == rows


# 복잡한 단일 XML 레코드의 중첩 업무 필드를 보존하는지 검증함
def test_single_xml_record_keeps_nested_business_field():
    source = public.read_structured_source(b'<record><detail><code>001</code></detail></record>', 'data.xml')
    assert source['items'] == [{'detail': {'code': '001'}}]


# 단일 JSON 객체 내부의 배열을 별도 레코드로 오인하지 않는지 검증함
def test_single_json_object_does_not_extract_internal_object_array_as_records():
    row = {'code':'001','points':[{'x':1},{'x':2}]}
    source = public.read_structured_source(json.dumps(row).encode(), 'data.json')
    assert source['items'] == [row]
    assert source['record_path'] == ''


# 명시적 루트 선택 시 복합 단일 레코드를 보존하는지 검증함
def test_explicit_root_selection_preserves_a_complex_single_record():
    row = {'data':[{'x':1},{'x':2}]}
    source = public.read_structured_source(json.dumps(row).encode(), 'data.json', '$')
    assert source['items'] == [row] and source['source_metadata'] is None


@pytest.mark.parametrize('suffix,raw,header', [
    ('json', b'{"header":{"resultCode":"00"},"body":{"totalCount":20,"items":[{"code":"001"}]}}', {'resultCode':'00'}),
    ('json', b'{"response":{"body":{"totalCount":20,"items":[{"code":"001"}]}}}', None),
    ('xml', b'<response><header/><body><totalCount>20</totalCount><items><item><code>001</code></item></items></body></response>', {}),
    ('xml', b'<root><header><resultCode>00</resultCode></header><body><totalCount>20</totalCount><items><item><code>001</code></item></items></body></root>', {'resultCode':'00'}),
])
# 직접 응답 및 헤더 누락 응답의 페이지 메타데이터를 보존하는지 검증함
def test_direct_or_missing_header_envelopes_keep_pagination_without_invented_api_status(suffix,raw,header):
    model = public.public_payload(public.read_structured_source(raw, 'data.'+suffix))
    assert model['response']['header'] == header
    assert model['response']['body'] == {'totalCount':20,'items':[{'code':'001'}]}
    assert public.public_payload(public.read_structured_source(public.public_xml(model), 'back.xml')) == model


@pytest.mark.parametrize('suffix,raw,path', [
    ('json', b'{"data":{"rows":[{"code":"001"},{"code":"002"}]}}', '/data/rows'),
    ('json', b'{"response":{"header":{"resultCode":"00"},"body":{"items":{"item":{"code":"001"}}}}}', '/response/body/items'),
    ('xml', b'<response><header><resultCode>00</resultCode></header><body><items><item><code>001</code></item></items></body></response>', '/response/body/items'),
    ('xml', b'<records><record><field name="code">001</field></record></records>', '/records/record'),
])
# JSON·XML 래퍼와 단일 항목 변형의 레코드 경로를 해석하는지 검증함
def test_wrapper_and_single_item_variants(suffix,raw,path):
    source = public.read_structured_source(raw, 'data.' + suffix)
    assert source['record_path'] == path
    assert source['items'][0] == {'code':'001'}


@pytest.mark.parametrize('suffix,raw,path', [
    ('json', b'{"a":[{"code":"001"}],"b":[{"code":"002"}]}', '/b'),
    ('xml', b'<root><a><row><code>001</code></row><row><code>002</code></row></a><b><row><code>003</code></row><row><code>004</code></row></b></root>', '/root/b/row'),
])
# 복수 목록의 모호성을 감지하고 명시적 경로를 적용하는지 검증함
def test_ambiguous_lists_require_explicit_selection(suffix,raw,path):
    with pytest.raises(ValueError, match='record_path'):
        public.read_structured_source(raw, 'data.'+suffix)
    source = public.read_structured_source(raw, 'data.'+suffix, path)
    assert source['record_path'] == path
    assert source['items'][0]['code'] in {'002','003'}


@pytest.mark.parametrize('suffix,raw', [
    ('json', b'[{"code":"001","code":"002"}]'),
    ('json', b'[{"value":NaN}]'),
    ('json', b'[{"value":1e999}]'),
    ('json', b'[1,2,3]'),
    ('xml', b'<!DOCTYPE rows [<!ENTITY x "value">]><rows><row>&x;</row></rows>'),
    ('xml', '<!DOCTYPE rows [<!ENTITY x "value">]><rows><row>&x;</row></rows>'.encode('utf-16')),
    ('xml', b'<row>before<code>001</code>after</row>'),
    ('xml', b'<row>before<code>001</code></row>'),
])
# 유효하지 않거나 안전하지 않은 구조화 입력을 조용히 버리지 않는지 검증함
def test_structured_invalid_or_unsafe_data_is_not_silently_discarded(suffix,raw):
    with pytest.raises(ValueError):
        public.read_structured_source(raw, 'data.'+suffix)


# API 기본 동작이 양쪽 응답 구조를 사용하고 CSV에는 items만 저장하는지 검증함
@pytest.mark.parametrize('input_format,target', [('xml','json'),('json','xml'),('json','csv'),('xml','csv')])
def test_api_structured_bidirectional_outputs_and_original_pagination(client,input_format,target):
    model = {'response':{'header':{'resultCode':'00','resultMsg':'NORMAL_CODE(정상)'},
        'body':{'totalCount':2,'pageNo':3,'numOfRows':2,'items':[
            {'code':'001','value':'0','nested':{'key':'a,b'},'list':['a','b']},
            {'code':'002','value':''},
        ]}}}
    raw = public.public_xml(model) if input_format == 'xml' else json.dumps(model).encode()
    response = client.post('/api/v1/converter/convert', files={'file':('data.'+input_format,raw)}, data={'target_format':target})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['dataset_profile'] == 'public_data' and result['rows_count'] == 2
    output = settings.OUTPUT_DIR / 'converted' / result['file_name']
    mapping = json.loads(output.with_name(output.name+'.field-mappings.json').read_text(encoding='utf-8'))
    assert mapping['response_metadata']['body'] == {'totalCount':2,'pageNo':3,'numOfRows':2}
    if target == 'csv':
        with output.open(encoding='utf-8-sig',newline='') as stream:
            rows = list(csv.DictReader(stream))
        assert rows[0] == {'code':'001','value':'0','nested':'{"key":"a,b"}','list':'["a","b"]'}
        assert rows[1] == {'code':'002','value':'','nested':'','list':''}
        assert result['warnings']
    else:
        restored = public.public_payload(public.read_structured_source(output.read_bytes(),output.name))
        assert restored == model


# 일반 XML·JSON 배열을 공공데이터 응답으로 변환하는지 검증함
def test_api_generic_xml_and_json_arrays_use_public_envelope(client):
    for filename, raw in [('data.xml',b'<rows><row><code>001</code></row><row><code>002</code></row></rows>'),
                          ('data.json',b'[{"code":"001"},{"code":"002"}]')]:
        response = client.post('/api/v1/converter/convert',files={'file':(filename,raw)},data={'target_format':'json'})
        assert response.status_code == 200,response.text
        output = settings.OUTPUT_DIR / 'converted' / response.json()['file_name']
        body = json.loads(output.read_text(encoding='utf-8'))['response']['body']
        assert body == {'totalCount':2,'pageNo':0,'numOfRows':2,'items':[{'code':'001'},{'code':'002'}]}


# 모호한 XML의 record_path 지정 변환을 검증함
def test_api_record_path_can_resolve_ambiguous_xml(client):
    raw = b'<root><a><row><code>001</code></row><row><code>002</code></row></a><b><row><code>003</code></row><row><code>004</code></row></b></root>'
    bad = client.post('/api/v1/converter/convert',files={'file':('data.xml',raw)},data={'target_format':'json'})
    assert bad.status_code == 400 and 'record_path' in bad.json()['detail']
    good = client.post('/api/v1/converter/convert',files={'file':('data.xml',raw)},data={'target_format':'json','record_path':'/root/b/row'})
    assert good.status_code == 200,good.text
    assert good.json()['preview'] == [{'code':'003'},{'code':'004'}]


@pytest.mark.parametrize('separator', [',',';','\t','|'])
# CSV 구분자 자동 감지와 빈 데이터 행 보존을 검증함
def test_csv_separator_detection_and_empty_data_rows(separator):
    raw = f'code{separator}value\n001{separator}A\n{separator}\n002{separator}B\n'.encode()
    source = public.read_source(raw, 'data.csv')
    assert source['headers'] == ['code','value']
    assert source['rows'] == [['001','A'],['',''],['002','B']]


@pytest.mark.parametrize('raw', [b'<rows/>', b'<items/>', b'<records/>'])
# 빈 XML 컨테이너에서 가짜 행을 생성하지 않는지 검증함
def test_empty_xml_containers_have_no_fabricated_rows(raw):
    assert public.read_structured_source(raw, 'data.xml')['items'] == []


# 단일 XML item 컨테이너의 레코드를 읽는지 검증함
def test_single_nested_xml_item_container():
    source = public.read_structured_source(b'<root><items><item><code>001</code></item></items></root>', 'data.xml')
    assert source['items'] == [{'code':'001'}]


@pytest.mark.parametrize('input_format,target', [('xml','json'),('json','xml'),('json','csv')])
@pytest.mark.parametrize('profile', ['auto', 'public_data', 'records'])
# 대표 공공데이터의 컬럼명과 무관하게 구조 변환하는지 검증함
def test_api_structured_conversion_is_not_specific_to_sample_columns(client,input_format,target,profile):
    original = _structured_public_model()
    raw = json.dumps(original, ensure_ascii=False).encode() if input_format == 'json' else public.public_xml(original)
    response = client.post(
        '/api/v1/converter/convert',
        files={'file':('public-data.'+input_format, raw)},
        data={'target_format':target,'dataset_profile':profile},
    )
    assert response.status_code == 200,response.text
    result = response.json()
    assert result['rows_count'] == len(original['response']['body']['items'])
    assert result['columns_count'] == len(original['response']['body']['items'][0])
    output = settings.OUTPUT_DIR/'converted'/result['file_name']
    if target == 'csv':
        with output.open(encoding='utf-8-sig',newline='') as stream:
            assert list(csv.DictReader(stream)) == original['response']['body']['items']
    else:
        assert public.public_payload(public.read_structured_source(output.read_bytes(),output.name)) == original
        if target == 'xml':
            root = ET.parse(output).getroot()
            assert [c.tag for c in root] == ['header','body']
            assert [c.tag for c in root.find('body')] == ['totalCount','pageNo','numOfRows','items']
            # 들여쓰기·XML 선언을 제외한 모든 요소·속성·순서·업무 값이 원본과 같은지 검증함
            # XML 요소 비교용 서명을 생성함
            def signature(element):
                return (element.tag,element.attrib,element.text if not list(element) else (element.text or '').strip(),
                        [signature(child) for child in element])
            assert signature(root) == signature(ET.fromstring(public.public_xml(original)))
            assert not any('.' in element.tag or '[' in element.tag for element in root.iter())
            assert not root.findall('.//field')
        preview = public.public_payload(public.read_structured_source(result['structured_preview'].encode(), 'preview.'+target))
        assert preview['response']['body']['items'] == original['response']['body']['items'][:100]
        assert preview['response']['body']['totalCount'] == original['response']['body']['totalCount']


# 일반 래퍼의 원본 메타데이터를 매핑 파일에 보존하는지 검증함
def test_generic_wrapper_metadata_is_kept_in_mapping_file(client):
    raw = b'{"source":"agency","data":{"updated":"2026-01-01","rows":[{"code":"001"}]}}'
    response = client.post('/api/v1/converter/convert',files={'file':('data.json',raw)},data={'target_format':'csv'})
    assert response.status_code == 200,response.text
    output = settings.OUTPUT_DIR/'converted'/response.json()['file_name']
    mapping = json.loads(output.with_name(output.name+'.field-mappings.json').read_text(encoding='utf-8'))
    assert mapping['source_metadata'] == {'source':'agency','data':{'updated':'2026-01-01'}}


# XML 제어문자 입력을 API 클라이언트 오류로 반환하는지 검증함
def test_api_xml_control_character_is_input_error(client):
    response = client.post('/api/v1/converter/convert',files={'file':('data.json',b'[{"value":"bad\\u0001text"}]')},data={'target_format':'xml'})
    assert response.status_code == 400 and '제어문자' in response.json()['detail']
