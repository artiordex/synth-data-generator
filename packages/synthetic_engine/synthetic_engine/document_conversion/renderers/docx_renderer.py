# -*- coding: utf-8 -*-
# =============================================================================
# 파일명: docx_renderer.py
# 경로: packages/synthetic_engine/synthetic_engine/document_conversion/renderers/docx_renderer.py
# 목적: IR 트리를 DOCX 워드 문서로 렌더링 변환함.
# 작성자: AI Agent
# 작성일: 2026-09-13
# 수정일: 2026-09-13
# =============================================================================
"""Map document IR to native WordprocessingML without fetching resources."""
from io import BytesIO

from ..core.ir import (
    ConversionWarning, FieldIR, HyperlinkIR, ImageIR, LineBreakIR, ParagraphIR, TableIR, TabIR, TextRunIR,
)
from ..core.capabilities import TargetCapabilities
from ..geometry.table_geometry import build_virtual_grid


# property 작업을 수행함
def _property(parent, name, **attributes):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    node = parent.find(qn('w:' + name))
    if node is None:
        node = OxmlElement('w:' + name)
        parent.append(node)
    for key, value in attributes.items():
        node.set(qn('w:' + key), str(value))
    return node


# format run 작업을 수행함
def _format_run(run, source):
    from docx.shared import Pt, RGBColor
    run.bold, run.italic = source.bold, source.italic
    run.font.size = Pt(source.size_pt)
    run.font.superscript = source.superscript
    if source.subscript:
        run.font.subscript = True
    run.font.strike = source.strike.value == 'single'
    run.font.double_strike = source.strike.value == 'double'
    props = run._r.get_or_add_rPr()
    _property(props, 'u', val={'dashed': 'dash', 'wavy': 'wave'}.get(
        source.underline.value, source.underline.value))
    _property(props, 'spacing', val=round(source.letter_spacing_pt * 20))
    _property(props, 'w', val=round(source.scale_percent))
    fonts = {}
    if source.font_family_en:
        fonts.update(ascii=source.font_family_en, hAnsi=source.font_family_en)
    elif source.font_family_ko:
        fonts.update(ascii=source.font_family_ko, hAnsi=source.font_family_ko)
    if source.font_family_ko:
        fonts['eastAsia'] = source.font_family_ko
    if fonts:
        _property(props, 'rFonts', **fonts)
    if source.language:
        _property(props, 'lang', val=source.language, eastAsia=source.language)
    if source.color_hex:
        run.font.color.rgb = RGBColor.from_string(source.color_hex.lstrip('#'))
    if source.bg_color_hex:
        _property(props, 'shd', val='clear', fill=source.bg_color_hex.lstrip('#'))


# unmapped 작업을 수행함
def _unmapped(warnings, source, feature):
    warnings.append(ConversionWarning(
        'DOCX_UNMAPPED_CONTENT', f'DOCX renderer cannot fully map {feature}.',
        source_ref=getattr(source, 'source_ref', None), feature=feature,
    ))


# inlines 작업을 수행함
def _inlines(paragraph, inlines, warnings):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.text.run import Run
    from docx.shared import Pt
    stack = [(item, paragraph._p) for item in reversed(inlines)]
    while stack:
        item, parent = stack.pop()
        if isinstance(item, HyperlinkIR):
            link = OxmlElement('w:hyperlink')
            if item.target.startswith('#'):
                link.set(qn('w:anchor'), item.target[1:])
            else:
                link.set(qn('r:id'), paragraph.part.relate_to(item.target, RT.HYPERLINK, is_external=True))
            if item.title:
                link.set(qn('w:tooltip'), item.title)
            parent.append(link)
            stack.extend((child, link) for child in reversed(item.inlines))
            continue
        if isinstance(item, FieldIR):
            # Emit only page fields as executable fields, retaining cached text
            # for other instructions without executing them.
            command = {'PAGE_NUMBER': 'PAGE', 'PAGE': 'PAGE', 'NUMPAGES': 'NUMPAGES',
                       'PAGE_COUNT': 'NUMPAGES', 'SECTIONPAGES': 'SECTIONPAGES'}.get(item.field_type.upper())
            if command:
                field = OxmlElement('w:fldSimple')
                field.set(qn('w:instr'), command)
                field.set(qn('w:dirty'), 'true')
                parent.append(field)
                parent = field
            else:
                _unmapped(warnings, item, 'field:' + item.field_type)
            element = OxmlElement('w:r')
            parent.append(element)
            Run(element, paragraph).text = item.cached_text or ''
            continue
        element = OxmlElement('w:r')
        parent.append(element)
        run = Run(element, paragraph)
        if isinstance(item, TextRunIR):
            run.text = item.text
            _format_run(run, item)
        elif isinstance(item, TabIR):
            if item.position_pt is not None:
                paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(item.position_pt))
            run.add_tab()
        elif isinstance(item, LineBreakIR):
            run.add_break()
        else:
            _unmapped(warnings, item, type(item).__name__)


# 문단 작업을 수행함
def _paragraph(container, source, warnings):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    paragraph = container.add_paragraph()
    if source.list_type is not None or source.list_level is not None:
        _unmapped(warnings, source, 'list-numbering')
    if source.heading_level is not None and 1 <= source.heading_level <= 9:
        paragraph.style = 'Heading ' + str(source.heading_level)
    paragraph.alignment = {
        'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
        'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
        'distribute': WD_ALIGN_PARAGRAPH.DISTRIBUTE,
    }[source.align.value]
    layout = paragraph.paragraph_format
    layout.line_spacing = source.line_spacing_percent / 100
    layout.space_before, layout.space_after = Pt(source.space_before_pt), Pt(source.space_after_pt)
    layout.left_indent, layout.first_line_indent = Pt(source.indent_pt), Pt(-source.hanging_pt)
    layout.page_break_before = source.page_break_before
    layout.keep_with_next = source.keep_with_next
    layout.keep_together = source.keep_lines_together
    _inlines(paragraph, source.inlines, warnings)
    return paragraph


# 셀 스타일 서식 작업을 수행함
def _cell_style(target, source):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.shared import Pt
    target.vertical_alignment = {
        'top': WD_CELL_VERTICAL_ALIGNMENT.TOP, 'center': WD_CELL_VERTICAL_ALIGNMENT.CENTER,
        'bottom': WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
    }[source.vertical_align.value]
    if source.width_pt is not None:
        target.width = Pt(source.width_pt)
    props = target._tc.get_or_add_tcPr()
    if source.bg_color_hex:
        _property(props, 'shd', val='clear', fill=source.bg_color_hex.lstrip('#'))
    margins = _property(props, 'tcMar')
    for side, value in zip(('top', 'right', 'bottom', 'left'), source.padding_pt):
        _property(margins, side, w=round(value * 20), type='dxa')
    borders = _property(props, 'tcBorders')
    names = {'slash': 'tr2bl', 'backslash': 'tl2br', 'diagonal_up': 'tr2bl',
             'diagonal_down': 'tl2br', 'tr2bl': 'tr2bl', 'tl2br': 'tl2br',
             'top': 'top', 'bottom': 'bottom', 'left': 'left', 'right': 'right',
             'start': 'start', 'end': 'end', 'insideH': 'insideH', 'insideV': 'insideV'}
    styles = {'none': 'nil', 'solid': 'single', 'double': 'double', 'dotted': 'dotted',
              'dashed': 'dashed', 'dash_dot': 'dotDash'}
    for side, border in source.borders.items():
        if side in names:
            _property(borders, names[side], val=styles[border.style.value],
                      sz=round(border.width_pt * 8), color=border.color_hex.lstrip('#'))


# 블록 목록 작업을 수행함
def _blocks(elements, container, available_width, warnings):
    from docx.shared import Pt
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
    from docx.oxml.ns import qn
    stack = [('block', block, container, available_width) for block in reversed(elements)]
    while stack:
        kind, block, target, width = stack.pop()
        if kind == 'finish_cell':
            if not len(target._tc) or target._tc[-1].tag != qn('w:p'):
                target.add_paragraph()
            continue
        if isinstance(block, ParagraphIR):
            _paragraph(target, block, warnings)
        elif isinstance(block, TableIR):
            grid = build_virtual_grid(block)
            if not grid or not grid[0]:
                continue
            if block.caption is not None:
                _paragraph(target, block.caption, warnings)
            # Header/footer add_table requires width; body/cell add_table does not.
            if hasattr(target, 'is_linked_to_previous'):
                table = target.add_table(rows=len(grid), cols=len(grid[0]), width=Pt(width))
            else:
                table = target.add_table(rows=len(grid), cols=len(grid[0]))
            table.alignment = {'left': WD_TABLE_ALIGNMENT.LEFT, 'center': WD_TABLE_ALIGNMENT.CENTER,
                               'right': WD_TABLE_ALIGNMENT.RIGHT}[block.alignment.value]
            table.autofit = False
            column_widths = block.column_widths_pt
            if len(column_widths) == len(grid[0]):
                for index, value in enumerate(column_widths):
                    table.columns[index].width = Pt(value)
                    for cell in table.columns[index].cells:
                        cell.width = Pt(value)
            total = block.total_width_pt or sum(column_widths)
            if total:
                _property(table._tbl.tblPr, 'tblW', w=round(total * 20), type='dxa')
            for index, row in enumerate(table.rows):
                props = row._tr.get_or_add_trPr()
                if index < block.repeat_header_rows:
                    _property(props, 'tblHeader', val='true')
                if block.cant_split:
                    _property(props, 'cantSplit', val='true')
            children = []
            for row in block.rows:
                for cell in row:
                    destination = table.cell(cell.row_index, cell.col_index)
                    if cell.row_span > 1 or cell.col_span > 1:
                        destination = destination.merge(table.cell(
                            cell.row_index + cell.row_span - 1, cell.col_index + cell.col_span - 1))
                    _cell_style(destination, cell)
                    if cell.height_pt is not None and cell.row_span == 1:
                        target_row = table.rows[cell.row_index]
                        target_row.height = Pt(max(cell.height_pt, (target_row.height.pt if target_row.height else 0)))
                        target_row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
                    for paragraph in list(destination.paragraphs):
                        destination._tc.remove(paragraph._p)
                    child_width = cell.width_pt or sum(column_widths[cell.col_index:cell.col_index + cell.col_span]) or width
                    child_width = max(1, child_width - cell.padding_pt[1] - cell.padding_pt[3])
                    children.extend(('block', child, destination, child_width) for child in cell.content)
                    children.append(('finish_cell', None, destination, child_width))
            stack.extend(reversed(children))
        elif isinstance(block, ImageIR):
            if block.opacity != 1 or block.rotation_deg != 0:
                _unmapped(warnings, block, 'image-opacity-or-rotation')
            target.add_paragraph().add_run().add_picture(
                BytesIO(block.image_bytes), width=Pt(block.width_pt), height=Pt(block.height_pt))
            if block.caption is not None:
                _paragraph(target, block.caption, warnings)
        else:
            _unmapped(warnings, block, type(block).__name__)
            target.add_paragraph('Unsupported source object')


class DocxRenderer:
    capabilities = TargetCapabilities(
        supports_pages=True, supports_headers=True, supports_footers=True,
        supports_nested_tables=True, supports_rowspan=True, supports_colspan=True,
        supports_font_scaling=True, supports_letter_spacing=True, supports_diagonal_borders=True,
    )

    # render 작업을 수행함
    def render(self, document, output_path):
        """
            @description 문서 중간 표현을 대상 포맷으로 렌더링함
            @param {document} - 메서드 입력값임
            @param {output_path} - 메서드 입력값임
        """
        from docx import Document
        from docx.shared import Pt
        from docx.enum.section import WD_SECTION_START, WD_ORIENT
        output = Document()
        for index, section in enumerate(document.sections):
            target = output.sections[0] if index == 0 else output.add_section(WD_SECTION_START.NEW_PAGE)
            target.orientation = WD_ORIENT.LANDSCAPE if section.orientation.value == 'landscape' else WD_ORIENT.PORTRAIT
            if section.page_width_pt is not None:
                target.page_width = Pt(section.page_width_pt)
            if section.page_height_pt is not None:
                target.page_height = Pt(section.page_height_pt)
            for name in ('top', 'bottom', 'left', 'right'):
                setattr(target, name + '_margin', Pt(getattr(section, 'margin_' + name + '_pt')))
            target.header_distance, target.footer_distance = Pt(section.header_distance_pt), Pt(section.footer_distance_pt)
            _property(target._sectPr, 'cols', num=section.column_count, space=round(section.column_gap_pt * 20))
            width = target.page_width.pt - target.left_margin.pt - target.right_margin.pt
            for name in ('header', 'footer'):
                source = getattr(section, name)
                if index == 0 and source is None:
                    continue
                destination = getattr(target, name)
                destination.is_linked_to_previous = False
                if source is not None:
                    for paragraph in list(destination.paragraphs):
                        destination._element.remove(paragraph._p)
                    _blocks(source.elements, destination, width, document.warnings)
                    if not destination.paragraphs:
                        destination.add_paragraph()
            _blocks(section.elements, output, width, document.warnings)
        output.save(output_path)
        return output_path
