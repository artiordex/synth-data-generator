# -*- coding: utf-8 -*-
# =============================================================================
# 파일명: xlsx_renderer.py
# 경로: packages/synthetic_engine/synthetic_engine/document_conversion/renderers/xlsx_renderer.py
# 목적: IR 트리를 XLSX 엑셀 스프레드시트로 렌더링함.
# 작성자: AI Agent
# 작성일: 2026-09-13
# 수정일: 2026-09-13
# =============================================================================
"""Logical tables become sheets; nested tables receive independent sheets."""
from io import BytesIO
from ..core.ir import ParagraphIR, TableIR, ImageIR, TextRunIR, ConversionWarning
from ..core.capabilities import TargetCapabilities
from .text import inline_text


class XlsxRenderer:
    capabilities = TargetCapabilities(supports_rowspan=True, supports_colspan=True)

    # render 작업을 수행함
    def render(self, document, output_path):
        """
            @description 문서 중간 표현을 대상 포맷으로 렌더링함
            @param {document} - 메서드 입력값임
            @param {output_path} - 메서드 입력값임
        """
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.drawing.image import Image
        workbook = Workbook()
        main = workbook.active
        main.title = 'Document'
        for section in document.sections:
            for block in section.elements:
                if isinstance(block, ParagraphIR):
                    cell = main.cell(main.max_row + 1, 1, inline_text(block.inlines))
                    cell.data_type = 's'
        for index, table in enumerate((b for b in document.iter_blocks() if isinstance(b, TableIR)), 1):
            sheet = workbook.create_sheet(f'Table {index}')
            for column, width in enumerate(table.column_widths_pt, 1):
                sheet.column_dimensions[get_column_letter(column)].width = max(0.1, (width / 0.75 - 5) / 7)
            for row in table.rows:
                for cell in row:
                    text = '\n'.join(inline_text(p.inlines) for p in cell.content if isinstance(p, ParagraphIR))
                    value = cell.formula if cell.formula is not None else cell.cached_value if cell.cached_value is not None else text
                    target = sheet.cell(cell.row_index+1, cell.col_index+1, value)
                    if isinstance(value, str) and cell.formula is None:
                        target.data_type = 's'
                    paragraphs = [p for p in cell.content if isinstance(p, ParagraphIR)]
                    run = next((r for p in paragraphs for r in p.inlines if isinstance(r, TextRunIR)), None)
                    target.alignment = Alignment(wrap_text=True, vertical=cell.vertical_align.value,
                        horizontal=('distributed' if paragraphs[0].align.value == 'distribute' else paragraphs[0].align.value)
                        if paragraphs else 'left')
                    if run is not None:
                        target.font = Font(name=run.font_family_ko or run.font_family_en, size=run.size_pt,
                            bold=run.bold, italic=run.italic, color=run.color_hex or '000000',
                            strike=run.strike.value != 'none',
                            underline=run.underline.value if run.underline.value in {'single', 'double'} else None,
                            vertAlign='superscript' if run.superscript else 'subscript' if run.subscript else None)
                    if cell.bg_color_hex:
                        target.fill = PatternFill('solid', fgColor=cell.bg_color_hex)
                    styles = {'none': None, 'solid': 'thin', 'double': 'double',
                              'dotted': 'dotted', 'dashed': 'dashed', 'dash_dot': 'dashDot'}
                    sides = {key: Side(style=styles[border.style.value], color=border.color_hex)
                             for key, border in cell.borders.items() if key in {'top', 'bottom', 'left', 'right'}}
                    target.border = Border(**sides)
                    if cell.height_pt is not None:
                        sheet.row_dimensions[cell.row_index + 1].height = max(
                            sheet.row_dimensions[cell.row_index + 1].height or 0, cell.height_pt / cell.row_span)
                    if cell.number_format:
                        target.number_format = cell.number_format
                    if cell.row_span > 1 or cell.col_span > 1:
                        sheet.merge_cells(start_row=cell.row_index+1, start_column=cell.col_index+1,
                                          end_row=cell.row_index+cell.row_span, end_column=cell.col_index+cell.col_span)
                    for block in cell.content:
                        if isinstance(block, ImageIR):
                            picture = Image(BytesIO(block.image_bytes))
                            picture.width, picture.height = block.width_pt / 0.75, block.height_pt / 0.75
                            sheet.add_image(picture, target.coordinate)
        document.warnings.append(ConversionWarning('XLSX_GRID_DEGRADATION',
            'Paragraph layout is flattened. Nested tables use separate sheets; mixed run styles, '
            'floating images, exact font-dependent widths and formula cached values are not preserved.', feature='grid_layout'))
        workbook.save(output_path)
        workbook.close()
        return output_path
