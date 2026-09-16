# -*- coding: utf-8 -*-
# =============================================================================
# 파일명: gov_doc_data_parser.py
# 경로: packages/synthetic_engine/synthetic_engine/document_conversion/parsers/gov_doc_data_parser.py
# 목적: 공공 기술검토 문서(DOCX, HWPX) 및 API 명세 문서에서 텍스트로 내포된
#       XML·JSON 페이로드 및 요청/응답 파라미터, 에러코드를 자동 감지하여
#       구조화된 표(데이터 그리드 표, 헤더 메타 표, 파라미터 규격 표)로 파싱하고,
#       'AI친화_고가치_데이터셋_파일데이터용_템플릿.docx' 표준 디자인을 준수하여 보고서를 생성하는 통합 엔진
# 작성자: 개발팀
# 작성일: 2026-09-16
# =============================================================================
"""Unified parser and report renderer for extracting tables, API specifications, and embedded XML/JSON data grids."""
from __future__ import annotations

from dataclasses import asdict, dataclass, field
import io
import json
from pathlib import Path
import re
from typing import Any, Dict, List, Optional, Tuple, Union
import xml.etree.ElementTree as ET
import zipfile

import docx
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml as parse_oxml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

# -----------------------------------------------------------------------------
# 공문서 표준 색상 및 스타일 상수 ('AI친화_고가치_데이터셋_파일데이터용_템플릿.docx' 디자인 기준)
# -----------------------------------------------------------------------------
COLOR_PRIMARY = RGBColor(30, 58, 138)     # Deep Navy (#1E3A8A)
COLOR_TEXT = RGBColor(31, 41, 55)         # Dark Slate (#1F2937)
COLOR_MUTED = RGBColor(107, 114, 128)     # Gray (#6B7280)
COLOR_HEADER_BG = "EBF2FA"                # Soft Blue-Gray Header
COLOR_ALT_ROW = "F9FAFB"                  # Soft Zebra Gray
COLOR_FIRST_COL = "F8FAFC"                # First Column Soft Gray
COLOR_BORDER = "CBD5E1"                   # Slate Border (#CBD5E1)


def clean_text(text: Optional[str]) -> str:
    """공백 및 줄바꿈 정리."""
    if not text:
        return ""
    return re.sub(r"[ \t\r\f\v]+", " ", text).strip()


# -----------------------------------------------------------------------------
# DOCX 공문서 서식 헬퍼 함수들
# -----------------------------------------------------------------------------
def set_cell_background(cell, fill_hex: str):
    """셀 배경색 설정."""
    shading = parse_oxml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_margins(cell, top=120, bottom=120, left=140, right=140):
    """셀 내부 여백 설정 (dxa 단위)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_oxml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """테이블 괘선(테두리) 설정."""
    tblPr = table._tbl.tblPr
    borders = parse_oxml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def format_table_standard(table, col_widths: List[float], header_bg: str = COLOR_HEADER_BG):
    """'AI친화_고가치_데이터셋_파일데이터용_템플릿.docx' 표준 표 디자인 일괄 적용."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, COLOR_BORDER)

    for row_idx, row in enumerate(table.rows):
        if row_idx == 0:
            # 헤더 행 반복 속성
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_oxml(f'<w:tblHeader {nsdecls("w")}/>'))

        for col_idx, cell in enumerate(row.cells):
            if col_idx < len(col_widths):
                cell.width = Inches(col_widths[col_idx])

            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)

            # 배경색 적용
            if row_idx == 0:
                set_cell_background(cell, header_bg)
            elif col_idx == 0 and len(row.cells) > 2:
                set_cell_background(cell, COLOR_FIRST_COL)
            elif row_idx % 2 == 1:
                set_cell_background(cell, COLOR_ALT_ROW)

            # 문단 및 글꼴 스타일 적용
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                if row_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "맑은 고딕"
                    run._r.get_or_add_rPr().append(parse_oxml(f'<w:rFonts {nsdecls("w")} w:eastAsia="맑은 고딕"/>'))
                    run.font.size = Pt(9)
                    if row_idx == 0:
                        run.font.bold = True
                        run.font.color.rgb = COLOR_TEXT
                    else:
                        run.font.color.rgb = COLOR_TEXT


# -----------------------------------------------------------------------------
# 표 및 파싱 데이터 모델
# -----------------------------------------------------------------------------
@dataclass
class SimpleTable:
    """표 데이터 모델."""
    title: str
    headers: List[str]
    rows: List[List[str]]
    category: str = "general"  # general, parameter, error_code, overview, operation, data_grid, header_meta
    raw_source: Optional[str] = None

    # 표 데이터를 마크다운 표 형식 문자열로 직렬화함
    def to_markdown(self) -> str:
        lines = []
        if self.title:
            lines.append(f"### {self.title}\n")
        if not self.headers and not self.rows:
            return lines[0] if lines else ""

        headers = self.headers or [f"열{i+1}" for i in range(len(self.rows[0]) if self.rows else 0)]
        lines.append("| " + " | ".join(h.replace("|", "\\|") for h in headers) + " |")
        lines.append("| " + " | ".join("---" for _ in headers) + " |")
        for r in self.rows:
            padded = [(r[i] if i < len(r) else "") for i in range(len(headers))]
            lines.append("| " + " | ".join(val.replace("|", "\\|").replace("\n", "<br>") for val in padded) + " |")
        lines.append("")
        return "\n".join(lines)


@dataclass
class ParsedXmlJsonData:
    """XML 또는 JSON에서 파싱된 구조화 데이터."""
    source_type: str  # 'xml' or 'json'
    root_tag: str
    header_meta_table: Optional[SimpleTable] = None
    data_grid_table: Optional[SimpleTable] = None
    schema_table: Optional[SimpleTable] = None
    raw_snippet: str = ""


@dataclass
class DocumentParseResult:
    """문서 전체 파싱 결과 모델."""
    filename: str
    format: str
    title: str
    paragraph_count: int = 0
    total_tables_count: int = 0
    overview_tables: List[SimpleTable] = field(default_factory=list)
    operation_tables: List[SimpleTable] = field(default_factory=list)
    parameter_tables: List[SimpleTable] = field(default_factory=list)
    payload_data_tables: List[SimpleTable] = field(default_factory=list)
    error_code_tables: List[SimpleTable] = field(default_factory=list)
    quality_tables: List[SimpleTable] = field(default_factory=list)
    other_tables: List[SimpleTable] = field(default_factory=list)
    embedded_payloads: List[ParsedXmlJsonData] = field(default_factory=list)

    # 문서 내 모든 추출 표 목록을 단일 리스트로 병합 반환함
    def all_tables(self) -> List[SimpleTable]:
        all_t = []
        all_t.extend(self.overview_tables)
        all_t.extend(self.operation_tables)
        all_t.extend(self.parameter_tables)
        all_t.extend(self.payload_data_tables)
        all_t.extend(self.error_code_tables)
        all_t.extend(self.quality_tables)
        all_t.extend(self.other_tables)
        return all_t

    # 문서 전체 파싱 결과를 마크다운 보고서로 변환함
    def to_markdown(self) -> str:
        md = [
            f"# [문서 파싱 결과 보고서] {self.title}",
            f"- **파일명**: `{self.filename}` (형식: {self.format.upper()})",
            f"- **추출된 총 표 수**: {self.total_tables_count}개\n",
            "---\n"
        ]

        if self.payload_data_tables:
            md.append("## 1. [핵심] XML·JSON 응답 파싱 데이터 그리드 표 (Data Grid)")
            md.append("> 문서 내 텍스트로 내포되어 있던 XML/JSON 응답 페이로드를 행·열 2차원 표로 정밀 파싱한 결과입니다.\n")
            for t in self.payload_data_tables:
                md.append(t.to_markdown())

        if self.parameter_tables:
            md.append("## 2. API 요청 및 응답 파라미터 명세 표")
            for t in self.parameter_tables:
                md.append(t.to_markdown())

        if self.overview_tables:
            md.append("## 3. 서비스 및 데이터셋 개요 표")
            for t in self.overview_tables:
                md.append(t.to_markdown())

        if self.operation_tables:
            md.append("## 4. 상세기능(오퍼레이션) 목록 표")
            for t in self.operation_tables:
                md.append(t.to_markdown())

        if self.error_code_tables:
            md.append("## 5. 오류 및 에러 코드 정의 표")
            for t in self.error_code_tables:
                md.append(t.to_markdown())

        if self.quality_tables:
            md.append("## 6. 품질 진단 및 검증 표")
            for t in self.quality_tables:
                md.append(t.to_markdown())

        if self.other_tables:
            md.append("## 7. 기타 추출 표 목록")
            for t in self.other_tables:
                md.append(t.to_markdown())

        return "\n".join(md)

    # 파싱 결과를 딕셔너리 구조로 변환함
    def to_dict(self) -> Dict[str, Any]:
        return {
            "filename": self.filename,
            "format": self.format,
            "title": self.title,
            "paragraph_count": self.paragraph_count,
            "total_tables_count": self.total_tables_count,
            "overview_tables": [asdict(t) for t in self.overview_tables],
            "operation_tables": [asdict(t) for t in self.operation_tables],
            "parameter_tables": [asdict(t) for t in self.parameter_tables],
            "payload_data_tables": [asdict(t) for t in self.payload_data_tables],
            "error_code_tables": [asdict(t) for t in self.error_code_tables],
            "quality_tables": [asdict(t) for t in self.quality_tables],
            "other_tables": [asdict(t) for t in self.other_tables],
        }


# -----------------------------------------------------------------------------
# 파서 메인 클래스
# -----------------------------------------------------------------------------
class GovDocDataParser:
    """공공 기술검토 문서 및 XML/JSON 표 자동 파싱 엔진."""

    @staticmethod
    def parse_xml_snippet(xml_text: str, label_prefix: str = "API 응답") -> Optional[ParsedXmlJsonData]:
        """XML 문자열에서 레코드 배열, 헤더 메타정보, 스키마 표를 파싱함."""
        clean_xml = xml_text.strip()
        match = re.search(r"(<(?:\?xml|[a-zA-Z0-9_\-]+:?[a-zA-Z0-9_\-]*).*?>.*</[a-zA-Z0-9_\-]+>)", clean_xml, re.DOTALL)
        if match:
            clean_xml = match.group(1)
        elif not (clean_xml.startswith("<") and clean_xml.endswith(">")):
            return None

        try:
            root = ET.fromstring(clean_xml)
        except Exception:
            try:
                parser = etree.XMLParser(recover=True, encoding="utf-8")
                lxml_root = etree.fromstring(clean_xml.encode("utf-8"), parser=parser)
                clean_xml = etree.tostring(lxml_root, encoding="utf-8").decode("utf-8")
                root = ET.fromstring(clean_xml)
            except Exception:
                return None

        # 1) 헤더 메타정보 파싱
        header_meta: Dict[str, str] = {}
        for h_tag in ["header", "cmmMsgHeader"]:
            h_node = root.find(f".//{h_tag}")
            if h_node is not None:
                for ch in h_node:
                    tag_name = ch.tag.split("}")[-1]
                    header_meta[tag_name] = clean_text(ch.text)

        body_node = root.find(".//body")
        if body_node is not None:
            for ch in body_node:
                tag_name = ch.tag.split("}")[-1]
                if tag_name not in ("items", "item", "rows", "row", "records", "record"):
                    header_meta[tag_name] = clean_text(ch.text)

        header_table = None
        if header_meta:
            header_table = SimpleTable(
                title=f"{label_prefix} 헤더 및 상태 메타정보 표",
                headers=["항목명(속성)", "설정값"],
                rows=[[k, v] for k, v in header_meta.items()],
                category="header_meta"
            )

        # 2) 데이터 레코드 표 (Data Grid) 파싱
        items = root.findall(".//item")
        if not items:
            items = root.findall(".//row")
        if not items:
            items = root.findall(".//record")
        if not items:
            items = root.findall(".//entry")

        records: List[Dict[str, str]] = []
        for it in items:
            row_dict: Dict[str, str] = {}
            for ch in it:
                tag_name = ch.tag.split("}")[-1]
                row_dict[tag_name] = clean_text(ch.text)
            if row_dict:
                records.append(row_dict)

        data_grid_table = None
        schema_table = None

        if records:
            columns: List[str] = []
            for r in records:
                for k in r.keys():
                    if k not in columns:
                        columns.append(k)

            grid_rows = [[r.get(c, "") for c in columns] for r in records]
            data_grid_table = SimpleTable(
                title=f"{label_prefix} 데이터 레코드 표 ({len(records)}건 관측)",
                headers=columns,
                rows=grid_rows,
                category="data_grid"
            )

            schema_rows = []
            for col in columns:
                samples = [r[col] for r in records if r.get(col)]
                sample_val = samples[0] if samples else ""
                val_type = "string"
                if sample_val.isdigit():
                    val_type = "integer"
                elif re.match(r"^\d+\.\d+$", sample_val):
                    val_type = "number"
                elif re.match(r"^\d{4}-\d{2}-\d{2}", sample_val):
                    val_type = "date"
                schema_rows.append([col, f"response/body/items/item/{col}", val_type, "1" if all(col in r for r in records) else "0", sample_val])

            schema_table = SimpleTable(
                title=f"{label_prefix} 데이터 필드 스키마 명세 표",
                headers=["항목명(영문)", "계층 구조 경로", "데이터 타입", "필수여부", "샘플 데이터"],
                rows=schema_rows,
                category="parameter"
            )

        return ParsedXmlJsonData(
            source_type="xml",
            root_tag=root.tag.split("}")[-1],
            header_meta_table=header_table,
            data_grid_table=data_grid_table,
            schema_table=schema_table,
            raw_snippet=clean_xml[:500]
        )

    @staticmethod
    def parse_json_snippet(json_text: str, label_prefix: str = "API 응답") -> Optional[ParsedXmlJsonData]:
        """JSON 문자열에서 레코드 배열, 헤더 메타정보, 스키마 표를 파싱함."""
        clean_json = json_text.strip()
        match = re.search(r"(\[.*\]|\{.*\})", clean_json, re.DOTALL)
        if match:
            clean_json = match.group(1)
        else:
            return None

        try:
            data = json.loads(clean_json)
        except Exception:
            return None

        header_meta: Dict[str, str] = {}
        records: List[Dict[str, Any]] = []

        if isinstance(data, list):
            for item in data:
                if isinstance(item, dict):
                    records.append(item)
        elif isinstance(data, dict):
            # JSON 객체 내 데이터 레코드 및 헤더 메타데이터를 재귀 탐색함
            def search_container(obj: Any):
                nonlocal records, header_meta
                if isinstance(obj, dict):
                    for k, v in obj.items():
                        if k in ("items", "item", "rows", "row", "records", "record", "data") and isinstance(v, list):
                            for it in v:
                                if isinstance(it, dict):
                                    records.append(it)
                        elif isinstance(v, (str, int, float, bool)):
                            if k in ("resultCode", "resultMsg", "totalCount", "pageNo", "numOfRows", "status", "code", "message"):
                                header_meta[k] = str(v)
                        elif isinstance(v, dict):
                            search_container(v)

            search_container(data)
            if not records:
                flat = {k: str(v) for k, v in data.items() if not isinstance(v, (dict, list))}
                if flat:
                    records.append(flat)

        header_table = None
        if header_meta:
            header_table = SimpleTable(
                title=f"{label_prefix} 헤더 및 상태 메타정보 표",
                headers=["항목명(속성)", "설정값"],
                rows=[[k, v] for k, v in header_meta.items()],
                category="header_meta"
            )

        data_grid_table = None
        schema_table = None

        if records:
            columns: List[str] = []
            for r in records:
                for k in r.keys():
                    if k not in columns:
                        columns.append(k)

            grid_rows = [[clean_text(str(r.get(c, ""))) for c in columns] for r in records]
            data_grid_table = SimpleTable(
                title=f"{label_prefix} 데이터 레코드 표 ({len(records)}건 관측)",
                headers=columns,
                rows=grid_rows,
                category="data_grid"
            )

            schema_rows = []
            for col in columns:
                samples = [str(r[col]) for r in records if r.get(col) is not None]
                sample_val = samples[0] if samples else ""
                schema_rows.append([col, f"response/{col}", "string", "1" if all(col in r for r in records) else "0", sample_val])

            schema_table = SimpleTable(
                title=f"{label_prefix} 데이터 필드 스키마 명세 표",
                headers=["항목명(영문)", "계층 구조 경로", "데이터 타입", "필수여부", "샘플 데이터"],
                rows=schema_rows,
                category="parameter"
            )

        return ParsedXmlJsonData(
            source_type="json",
            root_tag="object" if isinstance(data, dict) else "array",
            header_meta_table=header_table,
            data_grid_table=data_grid_table,
            schema_table=schema_table,
            raw_snippet=clean_json[:500]
        )

    @classmethod
    def parse_docx(cls, file_source: Union[str, Path, bytes, io.BytesIO], filename: str = "document.docx") -> DocumentParseResult:
        """DOCX 문서를 분석하여 표와 내포된 XML/JSON 데이터를 정밀 파싱함."""
        if isinstance(file_source, (str, Path)):
            doc = docx.Document(str(file_source))
            p_name = Path(file_source).name
        else:
            bio = io.BytesIO(file_source) if isinstance(file_source, bytes) else file_source
            doc = docx.Document(bio)
            p_name = filename

        result = DocumentParseResult(
            filename=p_name,
            format="docx",
            title=p_name.replace(".docx", ""),
            paragraph_count=len(doc.paragraphs),
            total_tables_count=len(doc.tables)
        )

        for p in doc.paragraphs[:5]:
            p_txt = clean_text(p.text)
            if len(p_txt) > 5 and not p_txt.startswith("※"):
                result.title = p_txt
                break

        # 전체 단락에서 독자적으로 존재하는 XML/JSON 메시지 탐색
        full_doc_paragraphs_text = "\n".join(p.text for p in doc.paragraphs)
        xml_matches = re.findall(r"(<(?:OpenAPI_ServiceResponse|response|cmmMsgHeader)>.*?</(?:OpenAPI_ServiceResponse|response|cmmMsgHeader)>)", full_doc_paragraphs_text, re.DOTALL)
        for idx, xml_str in enumerate(xml_matches):
            parsed_pj = cls.parse_xml_snippet(xml_str, label_prefix=f"본문 XML 메시지 #{idx+1}")
            if parsed_pj:
                result.embedded_payloads.append(parsed_pj)
                if parsed_pj.header_meta_table:
                    result.payload_data_tables.append(parsed_pj.header_meta_table)
                if parsed_pj.data_grid_table:
                    result.payload_data_tables.append(parsed_pj.data_grid_table)
                if parsed_pj.schema_table:
                    result.parameter_tables.append(parsed_pj.schema_table)

        # 표(Table) 순회 및 분석
        for t_idx, table in enumerate(doc.tables):
            rows = table.rows
            if not rows:
                continue

            extracted_matrix: List[List[str]] = []
            for r in rows:
                row_vals = []
                seen_cells = set()
                for c in r.cells:
                    if c._tc not in seen_cells:
                        seen_cells.add(c._tc)
                        row_vals.append(clean_text(c.text))
                extracted_matrix.append(row_vals)

            if not extracted_matrix or not extracted_matrix[0]:
                continue

            headers = extracted_matrix[0]
            data_rows = extracted_matrix[1:]
            header_str = " ".join(headers)
            full_table_text = " ".join(" ".join(r) for r in extracted_matrix)

            # XML 또는 JSON 셀 탐색
            cell_payload_found = False
            for r in rows:
                for c in r.cells:
                    c_raw = c.text.strip()
                    if "<response>" in c_raw or "<item>" in c_raw or "<cmmMsgHeader>" in c_raw or (c_raw.startswith("<") and "</" in c_raw):
                        parsed_xml = cls.parse_xml_snippet(c_raw, label_prefix=f"표 #{t_idx+1} [XML 응답]")
                        if parsed_xml:
                            cell_payload_found = True
                            result.embedded_payloads.append(parsed_xml)
                            if parsed_xml.header_meta_table:
                                result.payload_data_tables.append(parsed_xml.header_meta_table)
                            if parsed_xml.data_grid_table:
                                result.payload_data_tables.append(parsed_xml.data_grid_table)
                            if parsed_xml.schema_table:
                                result.parameter_tables.append(parsed_xml.schema_table)
                    elif (c_raw.startswith("{") and "}" in c_raw) or (c_raw.startswith("[") and "]" in c_raw):
                        parsed_json = cls.parse_json_snippet(c_raw, label_prefix=f"표 #{t_idx+1} [JSON 응답]")
                        if parsed_json:
                            cell_payload_found = True
                            result.embedded_payloads.append(parsed_json)
                            if parsed_json.header_meta_table:
                                result.payload_data_tables.append(parsed_json.header_meta_table)
                            if parsed_json.data_grid_table:
                                result.payload_data_tables.append(parsed_json.data_grid_table)
                            if parsed_json.schema_table:
                                result.parameter_tables.append(parsed_json.schema_table)

            # 표 성격 분류
            t_title = f"표 #{t_idx+1}"
            if "항목명(영문)" in header_str and "항목구분" in header_str:
                is_req = any("serviceKey" in "".join(r) or "요청" in "".join(r) for r in data_rows)
                t_title += " [요청 파라미터 명세]" if is_req else " [응답 파라미터 명세]"
                result.parameter_tables.append(SimpleTable(title=t_title, headers=headers, rows=data_rows, category="parameter"))
            elif "에러코드" in header_str or "오류코드" in header_str or ("코드" in header_str and "에러메시지" in header_str):
                t_title += " [오류/에러 코드 명세]"
                result.error_code_tables.append(SimpleTable(title=t_title, headers=headers, rows=data_rows, category="error_code"))
            elif "상세기능명" in header_str or "상세기능 번호" in header_str:
                t_title += " [상세기능 명세]"
                result.operation_tables.append(SimpleTable(title=t_title, headers=headers, rows=data_rows, category="operation"))
            elif "API 서비스 정보" in full_table_text or "API명(국문)" in full_table_text or "데이터셋 개요" in full_table_text:
                t_title += " [서비스/데이터셋 개요]"
                result.overview_tables.append(SimpleTable(title=t_title, headers=headers, rows=data_rows, category="overview"))
            elif not cell_payload_found:
                result.other_tables.append(SimpleTable(title=t_title, headers=headers, rows=data_rows, category="general"))

        return result

    @classmethod
    def parse_hwpx(cls, file_source: Union[str, Path, bytes, io.BytesIO], filename: str = "document.hwpx") -> DocumentParseResult:
        """HWPX 문서를 분석하여 표와 내포된 XML/JSON 데이터를 정밀 파싱함."""
        bio = io.BytesIO(file_source) if isinstance(file_source, bytes) else file_source
        p_name = Path(file_source).name if isinstance(file_source, (str, Path)) else filename

        result = DocumentParseResult(
            filename=p_name,
            format="hwpx",
            title=p_name.replace(".hwpx", ""),
        )

        with zipfile.ZipFile(file_source if isinstance(file_source, (str, Path)) else bio) as z:
            section_names = sorted([n for n in z.namelist() if n.startswith("Contents/section") and n.endswith(".xml")])
            HP = "http://www.hancom.co.kr/hwpml/2011/paragraph"

            total_tables = 0
            total_paragraphs = 0

            for sec_name in section_names:
                root = etree.fromstring(z.read(sec_name))
                paragraphs = root.findall(f".//{{{HP}}}p")
                total_paragraphs += len(paragraphs)
                tables = root.findall(f".//{{{HP}}}tbl")
                total_tables += len(tables)

                for t_idx, t in enumerate(tables):
                    rows = t.findall(f"{{{HP}}}tr")
                    if not rows:
                        continue

                    extracted_matrix: List[List[str]] = []
                    for r in rows:
                        cells = r.findall(f"{{{HP}}}tc")
                        row_vals = ["".join(c.itertext()).strip().replace("\n", " ") for c in cells]
                        extracted_matrix.append(row_vals)

                    if not extracted_matrix or not extracted_matrix[0]:
                        continue

                    headers = extracted_matrix[0]
                    data_rows = extracted_matrix[1:]
                    header_str = " ".join(headers)
                    full_table_text = " ".join(" ".join(r) for r in extracted_matrix)

                    cell_payload_found = False
                    for r in extracted_matrix:
                        for val in r:
                            if "<response>" in val or "<item>" in val or (val.startswith("<") and "</" in val):
                                parsed_xml = cls.parse_xml_snippet(val, label_prefix=f"HWPX 표 #{t_idx+1} [XML 응답]")
                                if parsed_xml:
                                    cell_payload_found = True
                                    result.embedded_payloads.append(parsed_xml)
                                    if parsed_xml.header_meta_table:
                                        result.payload_data_tables.append(parsed_xml.header_meta_table)
                                    if parsed_xml.data_grid_table:
                                        result.payload_data_tables.append(parsed_xml.data_grid_table)
                                    if parsed_xml.schema_table:
                                        result.parameter_tables.append(parsed_xml.schema_table)
                            elif (val.startswith("{") and "}" in val) or (val.startswith("[") and "]" in val):
                                parsed_json = cls.parse_json_snippet(val, label_prefix=f"HWPX 표 #{t_idx+1} [JSON 응답]")
                                if parsed_json:
                                    cell_payload_found = True
                                    result.embedded_payloads.append(parsed_json)
                                    if parsed_json.header_meta_table:
                                        result.payload_data_tables.append(parsed_json.header_meta_table)
                                    if parsed_json.data_grid_table:
                                        result.payload_data_tables.append(parsed_json.data_grid_table)
                                    if parsed_json.schema_table:
                                        result.parameter_tables.append(parsed_json.schema_table)

                    t_title = f"HWPX 표 #{t_idx+1}"
                    if "원천 구조 경로" in header_str or "필드명" in header_str or "데이터사전" in full_table_text:
                        t_title += " [데이터 사전(Data Dictionary)]"
                        result.parameter_tables.append(SimpleTable(title=t_title, headers=headers, rows=data_rows, category="parameter"))
                    elif "품질지표" in header_str or "COMPLETENESS" in full_table_text:
                        t_title += " [6대 품질진단 결과]"
                        result.quality_tables.append(SimpleTable(title=t_title, headers=headers, rows=data_rows, category="quality"))
                    elif "구분" in header_str and "내용" in header_str and len(headers) == 4:
                        t_title += " [데이터셋 개요]"
                        result.overview_tables.append(SimpleTable(title=t_title, headers=headers, rows=data_rows, category="overview"))
                    elif "확인 대상" in header_str or "처리 기준" in header_str:
                        t_title += " [기관 확인 필요 목록]"
                        result.other_tables.append(SimpleTable(title=t_title, headers=headers, rows=data_rows, category="review_required"))
                    elif not cell_payload_found:
                        result.other_tables.append(SimpleTable(title=t_title, headers=headers, rows=data_rows, category="general"))

            result.paragraph_count = total_paragraphs
            result.total_tables_count = total_tables

        return result

    @classmethod
    def parse_file(cls, file_path: Union[str, Path]) -> DocumentParseResult:
        """파일 경로로부터 확장자를 자동 판별하여 문서를 표로 파싱함."""
        p = Path(file_path)
        if not p.exists():
            raise FileNotFoundError(f"파일을 찾을 수 없습니다: {file_path}")

        ext = p.suffix.lower()
        if ext == ".docx":
            return cls.parse_docx(p)
        elif ext == ".hwpx":
            return cls.parse_hwpx(p)
        elif ext == ".xml":
            xml_text = p.read_text(encoding="utf-8", errors="replace")
            res = DocumentParseResult(filename=p.name, format="xml", title=p.stem)
            pj = cls.parse_xml_snippet(xml_text, label_prefix=f"{p.stem} XML 데이터")
            if pj:
                res.embedded_payloads.append(pj)
                if pj.header_meta_table:
                    res.payload_data_tables.append(pj.header_meta_table)
                if pj.data_grid_table:
                    res.payload_data_tables.append(pj.data_grid_table)
                if pj.schema_table:
                    res.parameter_tables.append(pj.schema_table)
                res.total_tables_count = len(res.all_tables())
            return res
        elif ext in (".json", ".jsonld"):
            json_text = p.read_text(encoding="utf-8", errors="replace")
            res = DocumentParseResult(filename=p.name, format="json", title=p.stem)
            pj = cls.parse_json_snippet(json_text, label_prefix=f"{p.stem} JSON 데이터")
            if pj:
                res.embedded_payloads.append(pj)
                if pj.header_meta_table:
                    res.payload_data_tables.append(pj.header_meta_table)
                if pj.data_grid_table:
                    res.payload_data_tables.append(pj.data_grid_table)
                if pj.schema_table:
                    res.parameter_tables.append(pj.schema_table)
                res.total_tables_count = len(res.all_tables())
            return res
        else:
            raise ValueError(f"지원하지 않는 문서 형식입니다: {ext} (지원: .docx, .hwpx, .xml, .json)")

    @classmethod
    def parse_text(cls, text: str, format: str = "xml", filename: str = "document") -> DocumentParseResult:
        """원문 텍스트(XML, JSON)를 직접 파싱하여 DocumentParseResult 생성."""
        fmt = format.lower().lstrip(".")
        res = DocumentParseResult(filename=filename, format=fmt, title=filename.replace(f".{fmt}", ""))
        if fmt == "xml":
            pj = cls.parse_xml_snippet(text, label_prefix=f"{res.title} XML 데이터")
        else:
            pj = cls.parse_json_snippet(text, label_prefix=f"{res.title} JSON 데이터")
        if pj:
            res.embedded_payloads.append(pj)
            if pj.header_meta_table:
                res.payload_data_tables.append(pj.header_meta_table)
            if pj.data_grid_table:
                res.payload_data_tables.append(pj.data_grid_table)
            if pj.schema_table:
                res.parameter_tables.append(pj.schema_table)
            res.total_tables_count = len(res.all_tables())
        return res

    @classmethod
    def parse_bytes(cls, data: bytes, format: str, filename: str = "document") -> DocumentParseResult:
        """바이트 스트림으로부터 문서(DOCX, HWPX) 또는 텍스트(XML, JSON) 파싱."""
        fmt = format.lower().lstrip(".")
        if fmt == "docx":
            return cls.parse_docx(io.BytesIO(data), filename=filename)
        elif fmt == "hwpx":
            return cls.parse_hwpx(io.BytesIO(data), filename=filename)
        elif fmt in ("xml", "json", "jsonld"):
            text = data.decode("utf-8", errors="replace")
            return cls.parse_text(text, format=fmt, filename=filename)
        else:
            raise ValueError(f"지원하지 않는 포맷입니다: {fmt}")



# -----------------------------------------------------------------------------
# 'AI친화_고가치_데이터셋_파일데이터용_템플릿.docx' 디자인을 100% 준수하는 DOCX 보고서 생성기
# -----------------------------------------------------------------------------
def render_parsed_report_docx(result: DocumentParseResult, output_path: Union[str, Path]) -> Path:
    """파싱된 결과(XML/JSON 데이터 그리드 및 표)를 '파일데이터용 템플릿.docx' 디자인 양식으로 DOCX 파일 생성."""
    doc = Document()
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    # 1. 페이지 여백 설정 (20mm / 0.8인치 표준)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 2. 표지 / 헤더 타이틀 ('AI친화_고가치_데이터셋_파일데이터용_템플릿.docx'와 100% 동일한 디자인)
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(14)
    run_t = title_p.add_run("AI 친화·고가치 공공데이터셋 설명서")
    run_t.font.name = "맑은 고딕"
    run_t._r.get_or_add_rPr().append(parse_oxml(f'<w:rFonts {nsdecls("w")} w:eastAsia="맑은 고딕"/>'))
    run_t.font.size = Pt(20)
    run_t.font.bold = True
    run_t.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(22)
    run_sub = sub_p.add_run(f"[{result.title}] 오픈API 및 데이터셋 표(Table) 파싱 보고서")
    run_sub.font.name = "맑은 고딕"
    run_sub._r.get_or_add_rPr().append(parse_oxml(f'<w:rFonts {nsdecls("w")} w:eastAsia="맑은 고딕"/>'))
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = COLOR_MUTED

    # 공통 섹션 제목 추가 함수
    def add_section_heading(text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "맑은 고딕"
        run._r.get_or_add_rPr().append(parse_oxml(f'<w:rFonts {nsdecls("w")} w:eastAsia="맑은 고딕"/>'))
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    # DOCX 문서에 소제목 단락을 추가함
    def add_sub_heading(text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "맑은 고딕"
        run._r.get_or_add_rPr().append(parse_oxml(f'<w:rFonts {nsdecls("w")} w:eastAsia="맑은 고딕"/>'))
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_TEXT
        return p

    # DOCX 문서에 안내 주석 단락을 추가함
    def add_notice(text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "맑은 고딕"
        run._r.get_or_add_rPr().append(parse_oxml(f'<w:rFonts {nsdecls("w")} w:eastAsia="맑은 고딕"/>'))
        run.font.size = Pt(8.5)
        run.font.color.rgb = COLOR_MUTED
        return p

    # SimpleTable 데이터를 DOCX 워드 표 객체로 렌더링함
    def render_simple_table_to_doc(stable: SimpleTable):
        add_sub_heading(f"□ {stable.title}")
        headers = stable.headers or [f"열{i+1}" for i in range(len(stable.rows[0]) if stable.rows else 1)]
        col_cnt = len(headers)
        row_cnt = len(stable.rows) + 1

        t = doc.add_table(rows=row_cnt, cols=col_cnt)
        # 열 너비 자동 분배 (총 너비 6.7인치)
        total_width = 6.7
        base_w = round(total_width / col_cnt, 2)
        col_widths = [base_w] * col_cnt

        for c_idx, h in enumerate(headers):
            t.rows[0].cells[c_idx].text = h

        for r_idx, r in enumerate(stable.rows, start=1):
            for c_idx in range(col_cnt):
                val = r[c_idx] if c_idx < len(r) else ""
                t.rows[r_idx].cells[c_idx].text = val

        format_table_standard(t, col_widths)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------------------
    # 섹션 1: [핵심] XML·JSON 데이터 그리드 표 (Data Grid)
    # -------------------------------------------------------------------------
    if result.payload_data_tables:
        add_section_heading("Ⅰ. XML·JSON 응답 파싱 데이터 그리드 표 (Data Grid)")
        add_notice("※ 원문 문서에 텍스트로 수록되어 있던 XML/JSON 응답 페이로드를 행·열 2차원 표로 정밀 전개한 결과입니다.")
        for st in result.payload_data_tables:
            render_simple_table_to_doc(st)

    # -------------------------------------------------------------------------
    # 섹션 2: API 요청 및 응답 파라미터 명세 표
    # -------------------------------------------------------------------------
    if result.parameter_tables:
        add_section_heading("Ⅱ. 전체 필드 데이터 사전 및 파라미터 명세 (Data Dictionary)")
        add_notice("※ 항목구분 : 필수(1), 옵션(0) / 데이터셋에 포함된 모든 관측 컬럼에 대한 물리명, 논리명, 데이터타입 및 제약조건")
        for st in result.parameter_tables:
            render_simple_table_to_doc(st)

    # -------------------------------------------------------------------------
    # 섹션 3: 서비스 및 데이터셋 개요 표
    # -------------------------------------------------------------------------
    if result.overview_tables:
        add_section_heading("Ⅲ. 데이터셋 일반 개요")
        for st in result.overview_tables:
            render_simple_table_to_doc(st)

    # -------------------------------------------------------------------------
    # 섹션 4: 상세기능(오퍼레이션) 목록 표
    # -------------------------------------------------------------------------
    if result.operation_tables:
        add_section_heading("Ⅳ. 상세기능(오퍼레이션) 목록")
        for st in result.operation_tables:
            render_simple_table_to_doc(st)

    # -------------------------------------------------------------------------
    # 섹션 5: 오류 및 에러 코드 명세 표
    # -------------------------------------------------------------------------
    if result.error_code_tables:
        add_section_heading("Ⅴ. 오픈API 오류 및 에러 코드 명세")
        add_notice("※ 연계 게이트웨이 표준 에러코드 및 기관 자체 에러코드 정의")
        for st in result.error_code_tables:
            render_simple_table_to_doc(st)

    # -------------------------------------------------------------------------
    # 섹션 6: 품질 진단 및 기타 표
    # -------------------------------------------------------------------------
    if result.quality_tables:
        add_section_heading("Ⅵ. AI 친화도 및 6대 품질 진단 결과 ([REF-02])")
        for st in result.quality_tables:
            render_simple_table_to_doc(st)

    if result.other_tables:
        add_section_heading("Ⅶ. 기타 추출 표 목록")
        for st in result.other_tables:
            render_simple_table_to_doc(st)

    doc.save(str(out))
    print(f"[성공] 파싱 결과 DOCX 보고서 생성 완료: {out.resolve()} ({out.stat().st_size:,} bytes)")
    return out
