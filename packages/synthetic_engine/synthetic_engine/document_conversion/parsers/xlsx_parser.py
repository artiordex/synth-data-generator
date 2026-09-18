# -*- coding: utf-8 -*-
# =============================================================================
# 파일명: xlsx_parser.py
# 경로: packages/synthetic_engine/synthetic_engine/document_conversion/parsers/xlsx_parser.py
# 목적: XLSX 엑셀 스프레드시트를 분석하여 IR 트리로 파싱함.
# 작성자: AI Agent
# 작성일: 2026-09-13
# 수정일: 2026-09-13
# =============================================================================
"""Spreadsheet worksheets become independent logical IR tables."""
from pathlib import Path
from .package import DocumentPackage
from ..core.ir import DocumentIR, SectionIR, TableIR, TableCellIR, ParagraphIR, TextRunIR, ConversionWarning
from ..core.source_ref import SourceRef


class XlsxParser:
    # parse 작업을 수행함
    def parse(self, path: Path):
        """
            @description 입력 문서를 파싱하여 중간 표현을 생성함
            @param {path} - 메서드 입력값임
        """
        from openpyxl import load_workbook
        from openpyxl.cell.cell import MergedCell
        with DocumentPackage(path) as package:
            if 'xl/workbook.xml' not in package.names:
                raise ValueError('Not an XLSX workbook')
        workbook = load_workbook(path, data_only=False, keep_links=False)
        cached = load_workbook(path, data_only=True, keep_links=False)
        document = DocumentIR(source_format='xlsx', source_path=str(path))
        try:
            for index, sheet in enumerate(workbook, 1):
                if sheet.max_row * sheet.max_column > 1000000:
                    raise ValueError('Worksheet exceeds one million grid cells')
                section = SectionIR(source_ref=SourceRef('xlsx', section_no=index, object_id=sheet.title))
                document.sections.append(section)
                rows = [[] for _ in range(sheet.max_row)]
                spans = {(r.min_row, r.min_col): (r.max_row-r.min_row+1, r.max_col-r.min_col+1) for r in sheet.merged_cells.ranges}
                for row in sheet:
                    for cell in row:
                        if isinstance(cell, MergedCell):
                            continue
                        if cell.value is None and not cell.has_style and (cell.row, cell.column) not in spans:
                            continue
                        formula = cell.value if cell.data_type == 'f' else None
                        value = cached[sheet.title].cell(cell.row, cell.column).value if formula else cell.value
                        text = '' if value is None else str(value)
                        row_span, col_span = spans.get((cell.row, cell.column), (1, 1))
                        ref = SourceRef('xlsx', section_no=index, row_index=cell.row-1, col_index=cell.column-1, object_id=cell.coordinate)
                        rows[cell.row-1].append(TableCellIR(cell.row-1, cell.column-1,
                            row_span=row_span, col_span=col_span, formula=formula,
                            cached_value=value if isinstance(value, (str, int, float, bool)) else text,
                            number_format=cell.number_format,
                            content=[ParagraphIR(inlines=[TextRunIR(text, font_family_en=cell.font.name,
                                size_pt=cell.font.sz or 11, bold=bool(cell.font.b), italic=bool(cell.font.i), source_ref=ref)])], source_ref=ref))
                section.elements.append(TableIR(rows=rows, source_ref=section.source_ref))
        finally:
            workbook.close()
            cached.close()
        document.warnings.append(ConversionWarning('XLSX_PARTIAL_LAYOUT',
            'Formula and cached value are separate. Charts, images, conditional formats and display-width metrics are not fully mapped.'))
        return document
