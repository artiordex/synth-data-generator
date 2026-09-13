# -*- coding: utf-8 -*-
# =============================================================================
# 파일명: pdf_high_fidelity_converter.py
# 경로: packages/synthetic_engine/synthetic_engine/exporters/pdf_high_fidelity_converter.py
# 목적: PDF 문서 레이아웃·표·텍스트를 유형화(Typology)하여 고충실도로 변환함
# 작성자: 개발팀
# 작성일: 2026-09-09
# 수정일: 2026-09-12
# =============================================================================
from __future__ import annotations

import os
import logging
import colorsys
import re
import io
import base64
import uuid
import zipfile
import tempfile
import html as html_lib
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple

import pymupdf
import pdfplumber
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


# rgb to hex 작업을 수행함
def rgb_to_hex(rgb: Optional[Tuple[float, ...]]) -> Optional[str]:
    """Convert RGB float or int tuple to hex string #rrggbb."""
    if not rgb:
        return None
    r, g, b = [int(max(0, min(255, c * 255 if isinstance(c, float) and c <= 1.0 else c))) for c in rgb[:3]]
    return f"#{r:02x}{g:02x}{b:02x}"


# int 색상 to hex 작업을 수행함
def int_color_to_hex(color_int: Optional[int]) -> str:
    """Convert integer RGB color from PyMuPDF to hex string."""
    if color_int is None:
        return "#000000"
    r = (color_int >> 16) & 255
    g = (color_int >> 8) & 255
    b = color_int & 255
    return f"#{r:02x}{g:02x}{b:02x}"


# lightness 작업을 수행함
def _lightness(color: str) -> float:
    rgb = tuple(int(color[i:i+2], 16) / 255 for i in (1, 3, 5))
    return colorsys.rgb_to_hls(*rgb)[1]


# 어두운 배경 여부 및 유효성을 판별함
def is_dark(hex_c: Optional[str]) -> bool:
    """Check if color is dark (brightness < 130)."""
    if not hex_c or not hex_c.startswith("#"):
        return False
    h = hex_c.lstrip("#")
    if len(h) != 6:
        return False
    try:
        r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
        return ((r * 299 + g * 587 + b * 114) / 1000) < 130
    except ValueError:
        return False


# 워드(DOCX) set 셀 background 작업을 수행함
def _docx_set_cell_background(cell, hex_color: str):
    """Set background color of a table cell in DOCX."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color.lstrip("#")}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


# 워드(DOCX) set 셀 margins 작업을 수행함
def _docx_set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding of a table cell in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)


# 워드(DOCX) set 셀 테두리 작업을 수행함
def _docx_set_cell_border(cell, **kwargs):
    """Set cell borders in DOCX."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" w:sz="{edge_data.get("sz", 4)}" w:space="0" w:color="{edge_data.get("color", "auto").lstrip("#")}"/>'
            tcBorders.append(parse_xml(tag))
        else:
            tag = f'<w:{edge} {nsdecls("w")} w:val="none"/>'
            tcBorders.append(parse_xml(tag))
    tcPr.append(tcBorders)


# 바운딩 박스 tuple 작업을 수행함
def _bbox_tuple(value: Any) -> Tuple[float, float, float, float]:
    """Return a normalized PyMuPDF/pdfplumber bbox tuple."""
    if hasattr(value, "x0"):
        return (float(value.x0), float(value.y0), float(value.x1), float(value.y1))
    return tuple(float(v) for v in value[:4])  # type: ignore[index]


# 바운딩 박스 intersection area 작업을 수행함
def _bbox_intersection_area(
    a: Tuple[float, float, float, float],
    b: Tuple[float, float, float, float],
) -> float:
    x0 = max(a[0], b[0])
    y0 = max(a[1], b[1])
    x1 = min(a[2], b[2])
    y1 = min(a[3], b[3])
    if x1 <= x0 or y1 <= y0:
        return 0.0
    return (x1 - x0) * (y1 - y0)


# 바운딩 박스 area 작업을 수행함
def _bbox_area(bbox: Tuple[float, float, float, float]) -> float:
    return max(0.0, bbox[2] - bbox[0]) * max(0.0, bbox[3] - bbox[1])


# fitz rect contains 작업을 수행함
def fitz_rect_contains(outer, inner) -> bool:
    return outer[0] <= inner[0] and outer[1] <= inner[1] and outer[2] >= inner[2] and outer[3] >= inner[3]


# overlap ratio 작업을 수행함
def _overlap_ratio(
    inner: Tuple[float, float, float, float],
    outer: Tuple[float, float, float, float],
) -> float:
    area = _bbox_area(inner)
    if area <= 0:
        return 0.0
    return _bbox_intersection_area(inner, outer) / area


# join 셀 텍스트 작업을 수행함
def _join_cell_text(parts: List[str]) -> str:
    return "".join(parts)


# escape HTML 웹 문서 작업을 수행함
def _escape_html(value: Any) -> str:
    return html_lib.escape(str(value or ""), quote=True)


# escape 마크다운 셀 작업을 수행함
def _escape_md_cell(value: Any) -> str:
    return str(value if value is not None else "").replace("|", "\\|").replace("\n", "<br/>")


# 행 셀 목록 작업을 수행함
def _row_cells(row: Dict[str, Any]) -> List[Dict[str, Any]]:
    if "cells" in row:
        return row["cells"]
    count = 4 if row.get("type") == "4col" else 2
    return [{"text": row.get(f"c{i}", "")} for i in range(count)]


# 컬럼 너비 목록 작업을 수행함
def _column_widths(element: Dict[str, Any], total: float) -> List[float]:
    count = max((len(_row_cells(row)) for row in element["rows"]), default=1)
    edges = element.get("col_edges", [])
    widths = [b-a for a, b in zip(edges, edges[1:])]
    if len(widths) != count or any(w <= 0 for w in widths):
        widths = [1.0] * count
    return [total * w / sum(widths) for w in widths]


# inline 이미지 HTML 웹 문서 작업을 수행함
def _inline_image_html(image: Dict[str, Any]) -> str:
    mime = 'image/jpeg' if image.get('format') in ('jpg', 'jpeg') else 'image/png'
    data = base64.b64encode(image['image_bytes']).decode('ascii')
    width = max(1, image['bbox'][2] - image['bbox'][0])
    return f'<img src="data:{mime};base64,{data}" alt="" style="width:{width}pt;max-width:100%;height:auto"/>'


# excel inline 이미지 목록 작업을 수행함
def _excel_inline_images(sheet, row: int, column: int, images: list) -> None:
    from openpyxl.drawing.image import Image
    for image in images:
        drawing = Image(io.BytesIO(image['image_bytes']))
        drawing.width = max(1, image['bbox'][2] - image['bbox'][0]) * 96 / 72
        drawing.height = max(1, image['bbox'][3] - image['bbox'][1]) * 96 / 72
        sheet.add_image(drawing, f'{get_column_letter(column)}{row}')
        sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 15, drawing.height * 72 / 96)


class HighFidelityPdfDoc:
    """
    Parses and reconstructs a PDF document into high-fidelity semantic components:
    - Banners (Title, Subtitle)
    - Alert / Notice callouts
    - Section Heading bars
    - Cards (Consultation records, transaction applications, medical notes)
    - Tabular Key-Value grids with merged columns
    - General body paragraphs
    - Page footers
    """

    # HighFidelityPdfDoc 인스턴스 멤버 변수 및 초기 설정을 구성함
    def __init__(self, pdf_path: Path):
        self.pdf_path = Path(pdf_path).resolve()
        self.doc = pymupdf.open(str(self.pdf_path))
        self.pages: List[Dict[str, Any]] = []
        self._parse()

    # close 작업을 수행함
    def close(self):
        if self.doc and not self.doc.is_closed:
            self.doc.close()

    # parse 작업을 수행함
    def _parse(self):
        plumber_tables_by_page = self._extract_pdfplumber_tables()

        for p_idx, page in enumerate(self.doc):
            p_w, p_h = page.rect.width, page.rect.height

            # 1. Extract major background drawings (width >= p_w * 0.6)
            raw_shapes = []
            for d in page.get_drawings():
                r = d.get("rect")
                if not r or r.width < p_w * 0.6:
                    continue
                # Skip full page canvas
                if r.height >= p_h * 0.7:
                    continue

                fill = rgb_to_hex(d.get("fill"))
                stroke = rgb_to_hex(d.get("color"))

                stype = "box"
                if is_dark(fill) and r.y0 < p_h * 0.35:
                    stype = "banner"
                elif fill and _lightness(fill) > 0.8 and r.height >= 35:
                    stype = "alert"
                elif fill and 0.35 < _lightness(fill) < 0.95 and r.height < 35 and r.y0 < p_h * 0.35 and len(page.get_textbox(r)) < 60:
                    stype = "section_bar"
                elif (stroke or fill == "#ffffff") and r.height >= 35:
                    stype = "card"

                raw_shapes.append({
                    "type": stype,
                    "rect": _bbox_tuple(r),
                    "fill": fill,
                    "stroke": stroke,
                    "height": r.height,
                })

            # Deduplicate shapes overlapping at the same Y coordinate
            shapes = []
            for s in sorted(raw_shapes, key=lambda x: x["rect"][1]):
                if not shapes:
                    shapes.append(s)
                else:
                    last = shapes[-1]
                    if abs(last["rect"][1] - s["rect"][1]) < 5 and abs(last["rect"][3] - s["rect"][3]) < 5:
                        if s["type"] in ("card", "alert", "banner"):
                            shapes[-1] = s
                    else:
                        shapes.append(s)

            # 2. Extract text blocks and spans
            text_blocks = page.get_text("dict")["blocks"]
            structured_blocks = []
            for b in text_blocks:
                if "lines" not in b:
                    continue
                lines_data = []
                for l in b["lines"]:
                    spans = []
                    for s in l["spans"]:
                        t = s["text"]
                        if t:
                            spans.append({
                                "text": t,
                                "size": round(s["size"], 1),
                                "bold": bool(s["flags"] & 16 or "bold" in s["font"].lower()),
                                "color": int_color_to_hex(s.get("color")),
                                "font": s["font"],
                                "bbox": _bbox_tuple(s["bbox"]),
                            })
                    if spans:
                        spans.sort(key=lambda x: (x["bbox"][0], x["bbox"][1]))
                        lines_data.append(spans)
                if lines_data:
                    lines_data.sort(key=lambda l: (
                        round(min(s["bbox"][1] for s in l) / 2) * 2,
                        min(s["bbox"][0] for s in l),
                    ))
                    full_t = "\n".join("".join(s["text"] for s in l) for l in lines_data).strip()
                    structured_blocks.append({
                        "bbox": _bbox_tuple(b["bbox"]),
                        "lines": lines_data,
                        "text": full_t,
                        "source_order": len(structured_blocks),
                    })

            # 3. Match blocks to container shapes
            matched_blocks = set()
            page_elements = []

            for table in plumber_tables_by_page.get(p_idx, []):
                page_elements.append(table)
                table_bbox = table["bbox"]
                table_text = ''.join(c['text'] for row in table['rows'] for c in _row_cells(row))
                table_text = re.sub(r'\s+', '', table_text)
                for b_idx, b in enumerate(structured_blocks):
                    if b_idx in matched_blocks:
                        continue
                    remaining_lines = []
                    for line in b['lines']:
                        remaining = [span for span in line if not (
                            _overlap_ratio(span['bbox'], table_bbox) >= 0.95
                            and re.sub(r'\s+', '', span['text']) in table_text)]
                        if remaining:
                            remaining_lines.append(remaining)
                    if not remaining_lines:
                        matched_blocks.add(b_idx)
                    elif remaining_lines != b['lines']:
                        b['lines'] = remaining_lines
                        b['text'] = '\n'.join(''.join(span['text'] for span in line) for line in remaining_lines)
                        boxes = [span['bbox'] for line in remaining_lines for span in line]
                        b['bbox'] = (min(r[0] for r in boxes), min(r[1] for r in boxes), max(r[2] for r in boxes), max(r[3] for r in boxes))

            for s in shapes:
                s_blocks = []
                sy0, sy1 = s["rect"][1], s["rect"][3]
                for b_idx, b in enumerate(structured_blocks):
                    if b_idx in matched_blocks:
                        continue
                    if _overlap_ratio(b["bbox"], s["rect"]) >= 0.45:
                        s_blocks.append(b)
                        matched_blocks.add(b_idx)

                if s_blocks:
                    s_blocks.sort(key=lambda b: self._reading_order_key(b))
                    # Double check container type
                    b_txt = "\n".join(b["text"] for b in s_blocks)
                    final_type = s["type"]
                    if "안내" in b_txt and ("가명" in b_txt or "테스트" in b_txt or "참고" in b_txt or "※" in b_txt):
                        final_type = "alert"
                    elif final_type == "box":
                        if is_dark(s["fill"]):
                            final_type = "banner"
                        elif s["height"] >= 35:
                            final_type = "card"

                    page_elements.append({
                        "type": final_type,
                        "y0": sy0,
                        "x0": s["rect"][0],
                        "bbox": s["rect"],
                        "rect": s["rect"],
                        "fill": s["fill"],
                        "stroke": s["stroke"],
                        "blocks": s_blocks,
                        "source": "pymupdf_shape",
                    })

            for table_row in self._detect_table_rows_from_text(structured_blocks, matched_blocks, p_w):
                page_elements.append(table_row)
                for b_idx in table_row.get("block_indices", []):
                    matched_blocks.add(b_idx)

            # 4. Handle remaining blocks (Tables, Standalone Headings, Paragraphs, Footers)
            for b_idx, b in enumerate(structured_blocks):
                if b_idx in matched_blocks:
                    continue
                by0, by1 = b["bbox"][1], b["bbox"][3]
                btxt = b["text"].strip()

                if by0 > p_h - 60 and re.match(r"^\s*-\s*\d+\s*-\s*$", btxt):
                    page_elements.append({
                        "type": "footer",
                        "y0": by0,
                        "x0": b["bbox"][0],
                        "bbox": b["bbox"],
                        "text": btxt,
                        "blocks": [b],
                        "source": "pymupdf_text",
                    })
                elif re.match(r"^\s*\d+\.\s+", btxt) or any(s["size"] >= 12 and s["bold"] for l in b["lines"] for s in l):
                    page_elements.append({"type": "paragraph", "y0": by0, "x0": b["bbox"][0],
                                          "bbox": b["bbox"], "blocks": [b], "source": "pymupdf_text"})
                else:
                    # 일반 본문 텍스트 블록은 인위적인 table_row 분할 없이 자연스러운 paragraph로 보존함
                    page_elements.append({
                        "type": "paragraph",
                        "y0": by0,
                        "x0": b["bbox"][0],
                        "bbox": b["bbox"],
                        "blocks": [b],
                        "text": btxt,
                        "source": "pymupdf_text",
                    })

            # 2.5 Extract embedded images from PDF page
            seen_image_xrefs = set()
            for img_info in page.get_images(full=True):
                xref = img_info[0]
                if xref in seen_image_xrefs:
                    continue
                seen_image_xrefs.add(xref)
                try:
                    base_img = self.doc.extract_image(xref)
                    img_bytes = base_img["image"]
                    img_ext = base_img["ext"]
                    iw, ih = base_img["width"], base_img["height"]
                    rects = page.get_image_rects(xref)
                    if rects:
                        r = rects[0]
                        img_bbox = (r.x0, r.y0, r.x1, r.y1)
                    else:
                        img_bbox = (42.5, p_h / 2, p_w - 42.5, p_h / 2 + 100)
                    for r in rects:
                        img_bbox = (r.x0, r.y0, r.x1, r.y1)
                        page_elements.append({
                            "type": "image",
                            "y0": img_bbox[1],
                            "x0": img_bbox[0],
                            "bbox": img_bbox,
                            "image_bytes": img_bytes,
                            "format": img_ext,
                            "width": iw,
                            "height": ih,
                            "source": "pdf_embedded_image",
                        })
                except Exception as exc:
                    logging.warning('PDF image extraction failed at page %s xref %s: %s', p_idx + 1, xref, exc)

            # Sort all elements by vertical reading position
            page_elements.sort(key=self._reading_order_key)

            # Merge consecutive table_row items into a single table element
            condensed_elements = []
            current_table_rows = []
            current_table_meta = []

            # flush current 표(테이블) 행 목록 작업을 수행함
            def _flush_current_table_rows():
                nonlocal current_table_rows, current_table_meta
                if not current_table_rows:
                    return
                # 최소 2행 이상이고 최소 2열 이상인 경우에만 표로 승격함
                max_cols = max((r.get("col_count", 0) for r in current_table_rows), default=0)
                if len(current_table_rows) >= 2 and max_cols >= 2:
                    y0 = min(m.get("y0", 0) for m in current_table_meta)
                    x0 = min(m.get("x0", 0) for m in current_table_meta)
                    x1 = max(m.get("bbox", (0, 0, 0, 0))[2] for m in current_table_meta)
                    y1 = max(m.get("bbox", (0, 0, 0, 0))[3] for m in current_table_meta)
                    condensed_elements.append({
                        "type": "table",
                        "y0": y0,
                        "x0": x0,
                        "bbox": (x0, y0, x1, y1),
                        "rows": current_table_rows,
                        "source": current_table_meta[0].get("source", "pymupdf_text"),
                    })
                else:
                    # 1행이거나 1열인 가짜 표는 원래의 본문 문단(paragraph)으로 안전하게 복원함
                    for m in current_table_meta:
                        blocks = m.get("blocks", [])
                        condensed_elements.append({
                            "type": "paragraph",
                            "y0": m.get("y0", 0),
                            "x0": m.get("x0", 0),
                            "bbox": m.get("bbox", (0, 0, 0, 0)),
                            "blocks": blocks,
                            "text": "\n".join(b.get("text", "") for b in blocks).strip() if blocks else "",
                            "source": "pymupdf_text",
                        })
                current_table_rows = []
                current_table_meta = []

            for elem in page_elements:
                if elem["type"] == "table_row":
                    current_table_rows.append(elem["cells"])
                    current_table_meta.append(elem)
                else:
                    _flush_current_table_rows()
                    condensed_elements.append(elem)

            _flush_current_table_rows()

            cells = [cell for element in condensed_elements if element['type'] == 'table'
                     for row in element['rows'] for cell in _row_cells(row) if cell.get('bbox')]
            standalone = []
            for element in condensed_elements:
                containing = [cell for cell in cells if element['type'] == 'image'
                              and fitz_rect_contains(cell['bbox'], element['bbox'])]
                if containing:
                    owner = min(containing, key=lambda cell: _bbox_area(cell['bbox']))
                    owner.setdefault('inline_images', []).append(element)
                else:
                    standalone.append(element)
            condensed_elements = standalone
            footer_text = next((e.get("text") for e in condensed_elements if e.get("type") == "footer"), None)
            self.pages.append({
                "page_num": p_idx + 1,
                "width": p_w,
                "height": p_h,
                "elements": condensed_elements,
                "footer_text": footer_text,
            })

        # 문서 구조 및 레이아웃 유형화(Typology) 분석 적용함
        try:
            from synthetic_engine.document_conversion.typology import DocumentTypologyPipeline
            from synthetic_engine.document_conversion.typology.models import ComponentType
            typology_pipe = DocumentTypologyPipeline(self.doc, self.pdf_path)
            self.typology_result = typology_pipe.run(self.pages)
            self.tokens = self.typology_result.tokens

            for p_idx, p_data in enumerate(self.pages):
                if p_idx < len(self.typology_result.pages):
                    t_page = self.typology_result.pages[p_idx]
                    p_data["archetype"] = t_page.archetype.value
                    p_data["running_header"] = t_page.running_header
                    p_data["running_footer"] = t_page.running_footer or p_data.get("footer_text")

                    for elem, t_block in zip(p_data["elements"], t_page.blocks):
                        elem["semantic_type"] = t_block.component_type.value
                        elem["semantic_role"] = t_block.component_type.name

                        # 시맨틱 분류가 비표(캡션, 콜아웃, 제목, 주석, 본문)인데 요소 타입이 table인 경우 강제 변환함
                        if elem.get("type") == "table" and t_block.component_type in (
                            ComponentType.TABLE_CAPTION,
                            ComponentType.FIGURE_CAPTION,
                            ComponentType.TABLE_NOTE,
                            ComponentType.CALLOUT_BOX,
                            ComponentType.HEADING_L1,
                            ComponentType.HEADING_L2,
                            ComponentType.HEADING_L3,
                            ComponentType.LIST_BULLET,
                            ComponentType.LIST_NUMBERED,
                            ComponentType.PARAGRAPH,
                        ):
                            tbl_texts = []
                            for row in elem.get("rows", []):
                                tbl_texts.append(" ".join(c.get("text", "") for c in _row_cells(row)).strip())
                            combined_text = "\n".join(t for t in tbl_texts if t).strip()

                            if t_block.component_type == ComponentType.CALLOUT_BOX:
                                elem["type"] = "alert"
                            else:
                                elem["type"] = "paragraph"

                            elem["text"] = combined_text
                            elem.pop("rows", None)
        except Exception as exc:
            logging.getLogger(__name__).warning("문서 레이아웃 유형화 분석 중 예외 발생: %s", exc)

    # reading order key 작업을 수행함
    @staticmethod
    def _reading_order_key(item: Dict[str, Any]) -> Tuple[int, float, int]:
        bbox = item.get("bbox") or item.get("rect") or (item.get("x0", 0), item.get("y0", 0), 0, 0)
        y0 = float(item.get("y0", bbox[1]))
        x0 = float(item.get("x0", bbox[0]))
        source_order = int(item.get("source_order", 0))
        return (int(round(y0 / 3.0) * 3), x0, source_order)

    # valid 표(테이블) candidate 여부 및 유효성을 판별함
    def _is_valid_table_candidate(
        self,
        table: Any,
        raw_rows: List[List[Any]],
        strategy: Dict[str, Any],
    ) -> bool:
        """비표(본문 문단, 제목, 캡션, 마진 장식선, 콜아웃 상자)의 표 오인 분할을 엄격히 차단하고 실제 2D 데이터 표만 승인함."""
        if not raw_rows or len(raw_rows) < 2:
            return False

        col_counts = [len([c for c in row if c is not None and str(c).strip()]) for row in raw_rows]
        max_cols = max(col_counts, default=0)
        if max_cols < 2:
            return False

        bbox = getattr(table, "bbox", None) or (0, 0, 0, 0)
        tbl_w = bbox[2] - bbox[0]
        tbl_h = bbox[3] - bbox[1]

        # 세로 장식선 및 여백 마진 띠 배제함
        if tbl_w < 60:
            return False
        if tbl_h > 400 and tbl_w < 120:
            return False
        if tbl_h > 580 and len(raw_rows) < 12:
            return False

        all_non_empty = [str(c).strip() for row in raw_rows for c in row if c and str(c).strip()]
        if not all_non_empty:
            return False

        # 유효 행 수 및 비어있는 행 비율 검사함 (실제 데이터 표는 행의 65% 이상이 채워져 있어야 함)
        non_empty_rows = sum(1 for row in raw_rows if any(c and str(c).strip() for c in row))
        if non_empty_rows < 2:
            return False
        if (non_empty_rows / len(raw_rows)) < 0.65:
            return False

        # 첫 행이 완전히 빈 행인 경우 (박스 상단 패딩 또는 장식 박스 테두리임)
        if all(not c or not str(c).strip() for c in raw_rows[0]):
            return False

        # 단일 셀 콜아웃 / 안내 상자 차단함 (불릿 기호나 안내 문구로 시작하는 소형 상자)
        if len(all_non_empty) <= 4 and any(c.startswith(("*", "**", "※", "•", "′", "'", "- ", "안내", "참고")) for c in all_non_empty):
            return False

        # 표/그림 캡션 단독 상자 차단함
        if any(re.match(r"^\s*<[\s]*(?:표|그림|table|fig)", c, re.IGNORECASE) for c in all_non_empty[:3]):
            if len(raw_rows) <= 3 and len(all_non_empty) <= 6:
                return False

        v_strat = strategy.get("vertical_strategy")
        h_strat = strategy.get("horizontal_strategy")

        # 수평선 기반 반경계 표(text/lines): 장식용 밑줄 사이의 대형 비표 영역 오인 차단함
        if h_strat == "lines":
            num_rows = len(raw_rows)
            if num_rows > 0 and (tbl_h / num_rows) > 75:
                return False
            if num_rows < 2 and tbl_h > 60:
                return False
            if tbl_w < 100:
                return False

        # 목차(Table of Contents, 차례, 표목차, 그림목차) 오인 분할 원천 차단함
        toc_keywords = (
            "contents", "table of contents", "목차", "목 차", "차례", "차 례",
            "표 목차", "표목차", "그림 목차", "그림목차", "색인", "index"
        )
        outline_prefix = re.compile(
            r"^\s*("
            r"\d+[\.\)]"
            r"|[I|V|X|i|v|x]+[\.\)]"
            r"|[가-힣][\.\)]"
            r"|[A-Za-z][\.\)]"
            r"|<[^>]+>"
            r"|제\s*\d+\s*[장절편부관]"
            r")",
            re.IGNORECASE,
        )

        toc_keyword_found = False
        page_num_rows = 0
        outline_rows = 0
        valid_row_count = 0

        for row in raw_rows:
            row_non_empty = [str(c).strip() for c in row if c and str(c).strip()]
            if not row_non_empty:
                continue
            valid_row_count += 1
            row_text = " ".join(row_non_empty).lower()
            if any(kw in row_text for kw in toc_keywords):
                toc_keyword_found = True

            first_val = row_non_empty[0]
            last_val = row_non_empty[-1]
            if re.match(r"^\s*(?:p\.?|page)?\s*\d{1,4}\s*(?:p|쪽)?\s*$", last_val, re.IGNORECASE):
                page_num_rows += 1
            if outline_prefix.match(first_val):
                outline_rows += 1

        if valid_row_count >= 2:
            if toc_keyword_found and (page_num_rows / valid_row_count) >= 0.2:
                return False
            if (page_num_rows / valid_row_count) >= 0.5 and (outline_rows / valid_row_count) >= 0.35:
                return False

        # 서술형 본문 문단 오인 분할 차단함
        prose_endings = ("다.", "다,", "니다.", "나타남", "중심이나", "비교하면", "기록됨", "순으로", "보임.", "있음.", "였음.")
        prose_cells = sum(1 for c in all_non_empty if len(c) > 25 or any(c.endswith(pe) for pe in prose_endings))
        if len(all_non_empty) > 0 and (prose_cells / len(all_non_empty)) >= 0.28:
            return False

        # 형태소/문장 단절 및 제목/불릿 조합 검사함
        total_cells = 0
        non_empty_cells = 0
        sentence_breaks = 0
        has_heading = False
        has_bullet = False

        korean_particles = (
            "은", "는", "이", "가", "을", "를", "에", "에서", "로", "으로",
            "와", "과", "의", "며", "고", "도", "만", "인", "하며", "하여",
            "된", "별로", "까지", "부터", "에게"
        )

        for row in raw_rows:
            row_non_empty = [str(c).strip() for c in row if c and str(c).strip()]
            total_cells += len(row)
            non_empty_cells += len(row_non_empty)

            for c in row_non_empty:
                if re.match(r"^\s*\d+\.\s+[가-힣A-Za-z0-9]", c):
                    has_heading = True
                if re.match(r"^\s*\d+\)\s+[가-힣A-Za-z0-9]", c) or c.startswith("•") or c.startswith("※") or c.startswith("- "):
                    has_bullet = True

            for i in range(len(row_non_empty) - 1):
                left = row_non_empty[i]
                right = row_non_empty[i + 1]
                if any(right.startswith(p) for p in korean_particles):
                    sentence_breaks += 1
                if left.endswith("-") or left.endswith(":") or (len(left) >= 2 and left[-1] in "공포데플기분"):
                    sentence_breaks += 1

        fill_rate = non_empty_cells / max(1, total_cells)

        if has_heading and has_bullet and v_strat == "text":
            return False
        if sentence_breaks >= 2 and v_strat == "text":
            return False
        if fill_rate < 0.35 and v_strat == "text":
            return False

        return True

    # pdfplumber 표 목록 요소를 추출하여 반환함
    def _extract_pdfplumber_tables(self) -> Dict[int, List[Dict[str, Any]]]:
        tables_by_page = {}
        try:
            with pdfplumber.open(str(self.pdf_path)) as document:
                for page_index, page in enumerate(document.pages):
                    xs = sorted({float(e["x0"]) for e in page.edges if e.get("orientation") == "v"})
                    ys = sorted({float(e["top"]) for e in page.edges if e.get("orientation") == "h"})

                    settings = [
                        {"vertical_strategy": "lines", "horizontal_strategy": "lines"},
                    ]
                    # Tier 1.5: 명시적 벡터 선분 기반 표
                    if len(xs) > 1 and len(ys) > 1:
                        settings.append({
                            "vertical_strategy": "explicit",
                            "horizontal_strategy": "explicit",
                            "explicit_vertical_lines": xs,
                            "explicit_horizontal_lines": ys,
                            "snap_tolerance": 6,
                            "intersection_tolerance": 8,
                        })

                    # Tier 2: 수평선 기반 반경계 표 (상/하단 및 헤더 구분선)
                    settings.append({"vertical_strategy": "text", "horizontal_strategy": "lines"})

                    found = []
                    spans = [span for block in self.doc[page_index].get_text("dict")["blocks"]
                             for line in block.get("lines", []) for span in line["spans"]]
                    for strategy in settings:
                        try:
                            candidates = page.find_tables(table_settings=strategy)
                        except Exception as exc:
                            logging.getLogger(__name__).warning("PDF table strategy failed page=%s strategy=%s: %s", page_index + 1, strategy, exc)
                            continue
                        for table in candidates:
                            bbox = _bbox_tuple(table.bbox)
                            if any(max(_overlap_ratio(bbox, t["bbox"]), _overlap_ratio(t["bbox"], bbox)) > 0.6 for t in found):
                                continue
                            raw_extracted = table.extract() or []
                            if not self._is_valid_table_candidate(table, raw_extracted, strategy):
                                continue
                            rows = self._normalize_table_rows(raw_extracted)
                            if len(rows) < 2:
                                continue
                            if strategy['vertical_strategy'] == 'text' and max(r['col_count'] for r in rows) < 2:
                                continue
                            edges = sorted({x for cell in table.cells for x in (cell[0], cell[2])})
                            for row, geometry in zip(rows, table.rows):
                                for cell, rect in zip(row["cells"], geometry.cells):
                                    if rect is None:
                                        continue
                                    styled = [p for p in spans if rect[0] <= (p["bbox"][0]+p["bbox"][2])/2 <= rect[2]
                                              and rect[1] <= (p["bbox"][1]+p["bbox"][3])/2 <= rect[3]]
                                    dominant = max(styled, key=lambda p: len(p["text"]), default={})
                                    cell.update({"bbox": rect, "bold": bool(dominant.get("flags", 0) & 16),
                                                 "size": dominant.get("size", 9), "font": dominant.get("font", "맑은 고딕"),
                                                 "color": int_color_to_hex(dominant.get("color", 0)), "spans": styled})
                            found.append({"type": "table", "y0": bbox[1], "x0": bbox[0], "bbox": bbox,
                                          "rect": bbox, "rows": rows, "source": "pdfplumber",
                                          "col_edges": edges, "strategy": dict(strategy)})
                    if found:
                        tables_by_page[page_index] = sorted(found, key=self._reading_order_key)
        except Exception as exc:
            logging.getLogger(__name__).warning("PDF table extraction failed: %s", exc)
        return tables_by_page

    # 감지 표(테이블) 행 목록 from 텍스트 작업을 수행함
    def _detect_table_rows_from_text(
        self,
        structured_blocks: List[Dict[str, Any]],
        matched_blocks: set,
        page_width: float,
    ) -> List[Dict[str, Any]]:
        line_items: List[Dict[str, Any]] = []
        for b_idx, block in enumerate(structured_blocks):
            if b_idx in matched_blocks:
                continue
            for line in block["lines"]:
                spans = sorted(line, key=lambda s: s["bbox"][0])
                text = _join_cell_text([s["text"] for s in spans])
                if not text:
                    continue
                bbox = (
                    min(s["bbox"][0] for s in spans),
                    min(s["bbox"][1] for s in spans),
                    max(s["bbox"][2] for s in spans),
                    max(s["bbox"][3] for s in spans),
                )
                line_items.append({
                    "bbox": bbox,
                    "text": text,
                    "spans": spans,
                    "block_index": b_idx,
                })

        line_items.sort(key=lambda l: (round(l["bbox"][1] / 2.0) * 2, l["bbox"][0]))
        clusters: List[List[Dict[str, Any]]] = []
        for item in line_items:
            center_y = (item["bbox"][1] + item["bbox"][3]) / 2
            if clusters:
                last_center = sum((l["bbox"][1] + l["bbox"][3]) / 2 for l in clusters[-1]) / len(clusters[-1])
                if abs(center_y - last_center) <= 4:
                    clusters[-1].append(item)
                    continue
            clusters.append([item])

        candidates: List[Dict[str, Any]] = []
        for cluster in clusters:
            cluster.sort(key=lambda l: l["bbox"][0])
            cells = self._cluster_line_items_as_cells(cluster)
            if len(cells) < 2:
                continue
            row = self._cells_to_table_row(cells, page_width=page_width)
            if not row:
                continue
            bbox = (
                min(c["bbox"][0] for c in cells),
                min(c["bbox"][1] for c in cells),
                max(c["bbox"][2] for c in cells),
                max(c["bbox"][3] for c in cells),
            )
            candidates.append({
                "type": "table_row",
                "y0": bbox[1],
                "x0": bbox[0],
                "bbox": bbox,
                "cells": row,
                "blocks": [structured_blocks[i] for i in sorted({c["block_index"] for c in cells})],
                "block_indices": sorted({c["block_index"] for c in cells}),
                "source": "pymupdf_text_grid",
                "column_signature": tuple(round(c["bbox"][0] / 12) * 12 for c in cells),
            })

        if len(candidates) < 2:
            return []

        accepted: List[Dict[str, Any]] = []
        run: List[Dict[str, Any]] = []
        for row in candidates:
            if not run:
                run = [row]
                continue
            prev = run[-1]
            same_columns = self._similar_column_signature(prev["column_signature"], row["column_signature"])
            close_vertical = row["y0"] - prev["bbox"][3] <= max(28, (prev["bbox"][3] - prev["bbox"][1]) * 1.5)
            if same_columns and close_vertical:
                run.append(row)
            else:
                if len(run) >= 2 and not self._is_table_of_contents_run(run):
                    accepted.extend(run)
                run = [row]
        if len(run) >= 2 and not self._is_table_of_contents_run(run):
            accepted.extend(run)
        return accepted

    # 표(테이블) of contents run 여부 및 유효성을 판별함
    @staticmethod
    def _is_table_of_contents_run(run: List[Dict[str, Any]]) -> bool:
        """목차(TOC), 표목차, 그림목차 등 번호 개요와 우측 페이지 번호로 구성된 목록인지 판정함."""
        if not run or len(run) < 2:
            return False
        total = len(run)
        page_num_matches = 0
        outline_matches = 0

        toc_header_keywords = (
            "contents", "table of contents", "목차", "목 차", "차례", "차 례",
            "표 목차", "표목차", "그림 목차", "그림목차", "색인", "index"
        )
        outline_pattern = re.compile(
            r"^\s*("
            r"\d+[\.\)]"
            r"|[I|V|X|i|v|x]+[\.\)]"
            r"|[가-힣][\.\)]"
            r"|[A-Za-z][\.\)]"
            r"|<[^>]+>"
            r"|제\s*\d+\s*[장절편부관]"
            r"|부록|참고자료|contents|목차|차례|appendix"
            r")",
            re.IGNORECASE,
        )

        for row in run:
            cells = row.get("cells", {}).get("cells", [])
            if len(cells) < 2:
                continue
            first_text = str(cells[0].get("text", "")).strip()
            last_text = str(cells[-1].get("text", "")).strip()

            if re.match(r"^\s*(?:p\.?|page)?\s*\d{1,4}\s*(?:p|쪽)?\s*$", last_text, re.IGNORECASE):
                page_num_matches += 1
            if (
                outline_pattern.match(first_text)
                or "." * 3 in first_text
                or "…" in first_text
                or any(kw in first_text.lower() for kw in toc_header_keywords)
            ):
                outline_matches += 1

        if (page_num_matches / total) >= 0.6:
            if (outline_matches / total) >= 0.3 or page_num_matches == total:
                return True

        return False

    # similar 컬럼 signature 작업을 수행함
    @staticmethod
    def _similar_column_signature(left: Tuple[float, ...], right: Tuple[float, ...]) -> bool:
        if abs(len(left) - len(right)) > 1:
            return False
        shared = min(len(left), len(right))
        if shared < 2:
            return False
        return sum(abs(left[i] - right[i]) <= 24 for i in range(shared)) >= shared - 1

    # cluster line items as 셀 목록 작업을 수행함
    @staticmethod
    def _cluster_line_items_as_cells(items: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        cells: List[Dict[str, Any]] = []
        for item in items:
            if not cells:
                cells.append(dict(item))
                continue
            prev = cells[-1]
            gap = item["bbox"][0] - prev["bbox"][2]
            if gap <= 8:
                prev["text"] = _join_cell_text([prev["text"], item["text"]])
                prev["bbox"] = (
                    min(prev["bbox"][0], item["bbox"][0]),
                    min(prev["bbox"][1], item["bbox"][1]),
                    max(prev["bbox"][2], item["bbox"][2]),
                    max(prev["bbox"][3], item["bbox"][3]),
                )
            else:
                cells.append(dict(item))
        return cells

    # 표(테이블) 행 목록 데이터를 표준 형식으로 정규화함
    def _normalize_table_rows(self, raw_rows: List[List[Any]], *, preserve_newlines: bool = True) -> List[Dict[str, Any]]:
        rows = []
        for raw_row in raw_rows:
            cells = [dict(c) if isinstance(c, dict) else {"text": "" if c is None else str(c)} for c in raw_row]
            if not preserve_newlines:
                cells = [{**c, "text": c["text"].replace("\n", " ")} for c in cells]
            row = self._cells_to_table_row(cells)
            if row is not None:
                rows.append(row)
        return rows

    # 셀 목록 to 표(테이블) 행 작업을 수행함
    def _cells_to_table_row(self, cells: List[Any], page_width: Optional[float] = None) -> Optional[Dict[str, Any]]:
        if not cells:
            return None
        values = []
        for cell in cells:
            value = dict(cell) if isinstance(cell, dict) else {"text": "" if cell is None else str(cell)}
            value["text"] = str(value.get("text", ""))
            values.append(value)
        return {"type": f"{len(values)}col", "cells": values, "col_count": len(values),
                **{f"c{i}": c["text"] for i, c in enumerate(values)}}

    # 표(테이블) 행 스팬 목록 데이터를 분석하여 파싱함
    def _parse_table_row_spans(self, block: Dict[str, Any], page_width: Optional[float] = None,
                               *, col_edges: Optional[List[float]] = None) -> Optional[Dict[str, Any]]:
        # 여러 줄로 구성된 일반 문단 텍스트는 단일 table_row로 오인 분할되지 않도록 보호함
        if len(block.get("lines", [])) > 1 and not col_edges and not block.get("col_edges"):
            return None

        spans = [s for line in block["lines"] for s in line]
        if not spans:
            return None
        edges = col_edges or block.get("col_edges")
        if not edges:
            starts = sorted(float(s["bbox"][0]) for s in spans)
            edges = []
            for x in starts:
                if not edges or x - edges[-1] > 35:
                    edges.append(x)
        if len(edges) < 2:
            return None
        groups = [[] for _ in edges]
        for span in spans:
            index = max((i for i, x in enumerate(edges) if span["bbox"][0] >= x - 2), default=0)
            groups[index].append(span)
        cells = []
        for group in groups:
            ordered = sorted(group, key=lambda s: (round(s["bbox"][1] / 3), s["bbox"][0]))
            text, previous_y = [], None
            for span in ordered:
                y = span["bbox"][1]
                if previous_y is not None and abs(y - previous_y) > 3:
                    text.append("\n")
                text.append(span["text"])
                previous_y = y
            dominant = max(ordered, key=lambda s: len(s["text"]), default={})
            cells.append({"text": "".join(text), "bold": dominant.get("bold", False),
                          "size": dominant.get("size", 9), "font": dominant.get("font", "맑은 고딕"),
                          "color": dominant.get("color", "#000000"), "spans": ordered})
        # 단일 행 내 개요 번호와 우측 페이지 번호 구조는 표 행으로 분할하지 않음
        if len(cells) > 1 and not col_edges and not block.get("col_edges"):
            first_t = str(cells[0].get("text", "")).strip()
            last_t = str(cells[-1].get("text", "")).strip()
            if re.match(r"^\s*(?:p\.?|page)?\s*\d{1,4}\s*(?:p|쪽)?\s*$", last_t, re.IGNORECASE):
                if re.match(r"^\s*(\d+[\.\)]|[I|V|X|i|v|x]+[\.\)]|[가-힣][\.\)]|<[^>]+>|제\s*\d+\s*[장절편부관])", first_t):
                    return None
        return self._cells_to_table_row(cells) if len(cells) > 1 else None

    # HTML 웹 문서 형식으로 변환하여 반환함
    def to_html(self, title: Optional[str] = None) -> str:
        """
        @description 학습된 디자인 토큰 및 시맨틱 유형화(Typology)를 반영한 고충실도 반응형 HTML을 생성함
        @param title: 문서 제목 문자열임 (미지정 시 파일명 사용함)
        @return: 원본 서식과 시맨틱 구조가 보존된 HTML 문서 문자열을 반환함
        """
        doc_title = title or self.pdf_path.stem
        tokens = getattr(self, "tokens", None)
        primary_color = getattr(tokens, "primary_color", "#1a365d") if tokens else "#1a365d"
        dark_color = getattr(tokens, "dark_color", "#2d3748") if tokens else "#2d3748"
        muted_color = getattr(tokens, "muted_color", "#718096") if tokens else "#718096"
        base_font = getattr(tokens, "base_font_family", "Noto Sans KR") if tokens else "Noto Sans KR"
        body_size = f"{getattr(tokens, 'font_size_body', 9.5):.1f}pt" if tokens else "9.5pt"
        h1_size = f"{getattr(tokens, 'font_size_h1', 18.0):.1f}pt" if tokens else "1.35rem"
        h2_size = f"{getattr(tokens, 'font_size_h2', 14.0):.1f}pt" if tokens else "1.1rem"
        h3_size = f"{getattr(tokens, 'font_size_h3', 11.5):.1f}pt" if tokens else "1.0rem"
        caption_size = f"{getattr(tokens, 'font_size_caption', 8.5):.1f}pt" if tokens else "0.85rem"

        html = []
        html.append(f"""<!DOCTYPE html>
<html lang="ko">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{_escape_html(doc_title)}</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Noto+Sans+KR:wght@300;400;500;600;700;800&family=JetBrains+Mono:wght@400;600&display=swap" rel="stylesheet">
<style>
:root {{
  --pdf-bg: #f1f5f9;
  --paper-bg: #ffffff;
  --primary-navy: {primary_color};
  --primary-blue: {primary_color};
  --alert-bg: #edf2f7;
  --alert-border: {primary_color};
  --alert-text: {dark_color};
  --sec-bar-bg: #cad4df;
  --card-bg: #ffffff;
  --card-border: #e2e8ef;
  --table-header-bg: #edf2f6;
  --table-border: #cbd5e1;
  --text-main: {dark_color};
  --text-muted: {muted_color};
  --font-family-doc: '{base_font}', 'Noto Sans KR', sans-serif;
  --font-body-size: {body_size};
  --font-h1-size: {h1_size};
  --font-h2-size: {h2_size};
  --font-h3-size: {h3_size};
  --font-caption-size: {caption_size};
}}
* {{ box-sizing: border-box; margin: 0; padding: 0; }}
body {{
  background-color: var(--pdf-bg);
  font-family: var(--font-family-doc);
  color: var(--text-main);
  line-height: 1.6;
  padding: 2.5rem 1rem;
}}
.pdf-container {{
  max-width: 860px;
  margin: 0 auto;
  display: flex;
  flex-direction: column;
  gap: 2.5rem;
}}
.pdf-page-card {{
  background: var(--paper-bg);
  border: 1px solid #cbd5e1;
  border-radius: 8px;
  padding: 2.8rem 2.5rem;
  box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.08), 0 8px 10px -6px rgba(0, 0, 0, 0.04);
  position: relative;
  min-height: 1000px;
  display: flex;
  flex-direction: column;
}}
.pdf-page-cover {{
  justify-content: center;
  align-items: center;
  text-align: center;
  background: linear-gradient(180deg, #ffffff 0%, #f8fafc 100%);
}}
.pdf-page-divider {{
  justify-content: center;
  background: #f8fafc;
}}
.pdf-page-badge {{
  position: absolute;
  top: 1.25rem;
  right: 1.5rem;
  font-size: 0.75rem;
  font-weight: 700;
  color: var(--text-muted);
  background: #f1f5f9;
  border: 1px solid #e2e8f0;
  padding: 0.25rem 0.65rem;
  border-radius: 4px;
}}
.pdf-running-header {{
  display: flex;
  justify-content: space-between;
  align-items: center;
  font-size: 0.8rem;
  color: var(--text-muted);
  border-bottom: 1px solid #e2e8f0;
  padding-bottom: 0.45rem;
  margin-bottom: 1.25rem;
  font-weight: 500;
}}
.pdf-running-footer {{
  display: flex;
  justify-content: center;
  align-items: center;
  font-size: 0.85rem;
  color: var(--text-muted);
  border-top: 1px solid #e2e8f0;
  padding-top: 0.75rem;
  margin-top: auto;
}}
.pdf-heading-l1 {{
  font-size: var(--font-h1-size);
  color: var(--primary-navy);
  font-weight: 800;
  margin-top: 1.4rem;
  margin-bottom: 0.75rem;
  line-height: 1.35;
}}
.pdf-heading-l2 {{
  font-size: var(--font-h2-size);
  color: var(--primary-navy);
  font-weight: 700;
  margin-top: 1.2rem;
  margin-bottom: 0.5rem;
  border-bottom: 2px solid var(--primary-blue);
  padding-bottom: 0.35rem;
}}
.pdf-heading-l3 {{
  font-size: var(--font-h3-size);
  color: var(--text-main);
  font-weight: 600;
  margin-top: 1rem;
  margin-bottom: 0.4rem;
}}
.pdf-table-caption {{
  font-size: var(--font-caption-size);
  font-weight: 700;
  color: var(--primary-navy);
  margin-top: 1.1rem;
  margin-bottom: 0.4rem;
}}
.pdf-figure-caption {{
  font-size: var(--font-caption-size);
  font-weight: 600;
  color: var(--text-muted);
  text-align: center;
  margin-top: 0.5rem;
  margin-bottom: 0.75rem;
}}
.pdf-note {{
  font-size: 0.8rem;
  color: var(--text-muted);
  margin-top: 0.3rem;
  margin-bottom: 0.75rem;
  line-height: 1.45;
}}
.pdf-list-item {{
  position: relative;
  padding-left: 1.2rem;
  margin-bottom: 0.35rem;
  font-size: var(--font-body-size);
  line-height: 1.6;
}}
.pdf-list-item::before {{
  content: "•";
  position: absolute;
  left: 0.3rem;
  color: var(--primary-blue);
  font-weight: bold;
}}
.pdf-banner {{
  background-color: var(--primary-navy);
  color: #ffffff;
  padding: 1.4rem 1.6rem;
  border-radius: 6px;
  margin-bottom: 1.2rem;
  box-shadow: 0 4px 6px -1px rgba(26, 54, 93, 0.2);
}}
.pdf-banner h1 {{
  font-size: 1.35rem;
  font-weight: 800;
  color: #ffffff;
  margin-bottom: 0.45rem;
  letter-spacing: -0.02em;
}}
.pdf-banner p {{
  font-size: 0.875rem;
  color: #cbd5e1;
  line-height: 1.45;
}}
.pdf-alert {{
  background-color: var(--alert-bg);
  border-left: 4.5px solid var(--alert-border);
  border-radius: 0 6px 6px 0;
  padding: 1rem 1.25rem;
  margin-bottom: 1.25rem;
  display: flex;
  gap: 0.75rem;
  align-items: flex-start;
}}
.pdf-alert-badge {{
  background: var(--alert-border);
  color: #ffffff;
  font-size: 0.75rem;
  font-weight: 700;
  padding: 0.15rem 0.5rem;
  border-radius: 4px;
  white-space: nowrap;
  margin-top: 0.15rem;
}}
.pdf-alert-text {{
  font-size: 0.875rem;
  color: var(--alert-text);
  line-height: 1.55;
}}
.pdf-section-bar {{
  background-color: var(--sec-bar-bg);
  border-left: 4.5px solid var(--primary-navy);
  padding: 0.55rem 1rem;
  border-radius: 0 4px 4px 0;
  margin-top: 1.4rem;
  margin-bottom: 1rem;
  font-size: 1.05rem;
  font-weight: 700;
  color: var(--primary-navy);
}}
.pdf-card {{
  background: #ffffff;
  border: 1px solid var(--card-border);
  border-radius: 6px;
  padding: 1.25rem 1.4rem;
  margin-bottom: 1rem;
  box-shadow: 0 1px 3px rgba(0,0,0,0.03);
}}
.pdf-card-title {{
  font-size: 0.95rem;
  font-weight: 700;
  color: var(--primary-navy);
  margin-bottom: 0.35rem;
}}
.pdf-card-meta {{
  font-size: 0.8rem;
  color: var(--text-muted);
  border-bottom: 1px dashed #e2e8f0;
  padding-bottom: 0.5rem;
  margin-bottom: 0.75rem;
}}
.pdf-card-body {{
  font-size: 0.875rem;
  line-height: 1.7;
  color: var(--text-main);
}}
.highlight-id {{
  color: var(--primary-blue);
  font-weight: 600;
}}
.highlight-tag {{
  background: #edf2f7;
  padding: 0.1rem 0.4rem;
  border-radius: 3px;
  font-family: 'JetBrains Mono', monospace;
  font-size: 0.825rem;
  color: #c53030;
  border: 1px solid #e2e8f0;
}}
.pdf-table-container {{
  width: 100%;
  margin: 0.5rem 0 1.25rem 0;
  overflow-x: auto;
}}
.pdf-table {{
  width: 100%;
  border-collapse: collapse;
  font-size: 0.875rem;
  border: 1px solid var(--table-border);
}}
.pdf-table th, .pdf-table td {{
  border: 1px solid var(--table-border);
  padding: 0.65rem 0.85rem;
  vertical-align: middle;
}}
.pdf-table th {{
  background-color: var(--table-header-bg);
  color: var(--primary-navy);
  font-weight: 700;
  width: 22%;
  text-align: left;
}}
.pdf-table td {{
  background-color: #ffffff;
  color: var(--text-main);
}}
.pdf-footer {{
  margin-top: auto;
  text-align: center;
  font-size: 0.85rem;
  color: var(--text-muted);
  padding-top: 2rem;
}}
</style>
</head>
<body>
<div class="pdf-container">
""")

        for p in self.pages:
            pno = p["page_num"]
            tot = len(self.pages)
            archetype = p.get("archetype", "body")
            page_class = "pdf-page-card"
            if archetype == "cover":
                page_class += " pdf-page-cover"
            elif archetype == "chapter_divider":
                page_class += " pdf-page-divider"

            html.append(f'<div class="{page_class}" id="page-{pno}">')
            html.append(f'  <div class="pdf-page-badge">Page {pno} / {tot}</div>')

            r_header = p.get("running_header")
            if r_header and archetype not in ("cover", "front_matter"):
                html.append(f'  <div class="pdf-running-header"><span>{_escape_html(r_header)}</span><span>p. {pno}</span></div>')

            for elem in p["elements"]:
                etype = elem["type"]
                stype = elem.get("semantic_type", "")

                # 런닝 헤더/푸터는 상단/하단 메타 컴포넌트로 처리하므로 본문 중복 출력 방지함
                if stype in ("running_header", "running_footer"):
                    continue

                if stype == "heading_l1" or etype == "banner":
                    txt = "\n".join(b["text"] for b in elem.get("blocks", [])).strip() if "blocks" in elem else str(elem.get("text", "")).strip()
                    lines = [l.strip() for l in txt.split("\n") if l.strip()]
                    t1 = lines[0] if lines else "문서 제목"
                    t2 = " ".join(lines[1:]) if len(lines) > 1 else ""
                    html.append(f"""  <div class="pdf-banner">
    <h1>{_escape_html(t1)}</h1>
    {f'<p>{_escape_html(t2)}</p>' if t2 else ''}
  </div>""")

                elif stype == "heading_l2" or etype == "section_bar":
                    txt = " ".join(b["text"] for b in elem.get("blocks", [])).strip() if "blocks" in elem else str(elem.get("text", "")).strip()
                    html.append(f"""  <h2 class="pdf-heading-l2">{_escape_html(txt)}</h2>""")

                elif stype == "heading_l3":
                    txt = " ".join(b["text"] for b in elem.get("blocks", [])).strip() if "blocks" in elem else str(elem.get("text", "")).strip()
                    html.append(f"""  <h3 class="pdf-heading-l3">{_escape_html(txt)}</h3>""")

                elif stype == "table_caption":
                    txt = str(elem.get("text", "")).strip() or "".join(b["text"] for b in elem.get("blocks", [])).strip()
                    html.append(f"""  <div class="pdf-table-caption">{_escape_html(txt)}</div>""")

                elif stype == "figure_caption":
                    txt = str(elem.get("text", "")).strip() or "".join(b["text"] for b in elem.get("blocks", [])).strip()
                    html.append(f"""  <div class="pdf-figure-caption">{_escape_html(txt)}</div>""")

                elif stype in ("table_note", "figure_note"):
                    txt = str(elem.get("text", "")).strip() or "".join(b["text"] for b in elem.get("blocks", [])).strip()
                    html.append(f"""  <div class="pdf-note">{_escape_html(txt)}</div>""")

                elif etype == "alert" or stype == "callout_box":
                    txt = "\n".join(b["text"] for b in elem.get("blocks", [])).strip() if "blocks" in elem else str(elem.get("text", "")).strip()
                    clean_alert = re.sub(r"^\s*안내\s*:\s*", "", txt)
                    html.append(f"""  <div class="pdf-alert">
    <span class="pdf-alert-badge">안내</span>
    <div class="pdf-alert-text">{_escape_html(clean_alert).replace(chr(10), '<br/>')}</div>
  </div>""")

                elif etype == "card":
                    txt = "\n".join(b["text"] for b in elem["blocks"]).strip()
                    lines = txt.split("\n")
                    c_title = lines[0] if lines else "상세 내역"
                    c_meta = ""
                    c_body_lines = lines[1:]
                    if len(lines) > 1 and any(k in lines[1] for k in ("접수일시", "일시", "담당자", "상담원:")) and len(lines[1]) < 60:
                        c_meta = lines[1]
                        c_body_lines = lines[2:]

                    body_str = "<br/>".join(_escape_html(l) for l in c_body_lines)
                    body_str = re.sub(r"(\[RRN Omitted[^\]]*\])", r'<span class="highlight-tag">\1</span>', body_str)
                    body_str = re.sub(r"(010-\d{4}-\d{4})", r'<span class="highlight-id">\1</span>', body_str)
                    body_str = re.sub(r"([a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+)", r'<span class="highlight-id">\1</span>', body_str)

                    html.append(f"""  <div class="pdf-card">
    <div class="pdf-card-title">{_escape_html(c_title)}</div>
    {f'<div class="pdf-card-meta">{_escape_html(c_meta)}</div>' if c_meta else ''}
    <div class="pdf-card-body">{body_str}</div>
  </div>""")

                elif etype == "image":
                    b64 = base64.b64encode(elem["image_bytes"]).decode("ascii")
                    mime = "image/jpeg" if elem.get("format", "").lower() in ("jpg", "jpeg") else "image/png"
                    html.append(f'''  <div class="pdf-image-container" style="text-align:center; margin:16px 0;">
    <img src="data:{mime};base64,{b64}" style="max-width:100%; height:auto; border-radius:8px; box-shadow:0 2px 4px rgba(0,0,0,0.06);" alt="문서 이미지" />
  </div>''')

                elif etype == "table":
                    html.append('  <div class="pdf-table-container">\n    <table class="pdf-table">')
                    html.append('<colgroup>' + ''.join(f'<col style="width:{width:.4f}%"/>' for width in _column_widths(elem, 100)) + '</colgroup>')
                    for row in elem["rows"]:
                        html.append("      <tr>")
                        for index, cell in enumerate(_row_cells(row)):
                            tag = "th" if index == 0 and len(cell["text"]) < 40 else "td"
                            color = cell.get("color", "#000000")
                            if isinstance(color, int):
                                color = int_color_to_hex(color)
                            style = f'white-space:pre-wrap;font-size:{cell.get("size", 9)}pt;color:{color};'
                            if cell.get('font'):
                                style += 'font-family:' + str(cell['font']).replace(';', '').replace('"', '').replace("'", '') + ';'
                            if cell.get("bold"):
                                style += "font-weight:bold;"
                            content = _escape_html(cell['text']).replace(chr(10), '<br/>')
                            content += ''.join(_inline_image_html(image) for image in cell.get('inline_images', []))
                            html.append(f'<{tag} style="{_escape_html(style)}">{content}</{tag}>')
                        html.append("      </tr>")
                    html.append("    </table>\n  </div>")

                elif etype == "paragraph":
                    txt = "\n".join(b["text"] for b in elem["blocks"]).strip()
                    if txt:
                        if stype == "list_bullet":
                            html.append(f'  <div class="pdf-list-item">{_escape_html(txt)}</div>')
                        else:
                            html.append(f'  <p style="margin-bottom:0.75rem; font-size:var(--font-body-size);">{_escape_html(txt).replace(chr(10), "<br/>")}</p>')

                elif etype == "footer":
                    html.append(f"""  <div class="pdf-footer">{_escape_html(elem.get('text', f'- {pno} -'))}</div>""")

            r_footer = p.get("running_footer") or p.get("footer_text")
            if r_footer and archetype != "cover":
                html.append(f'  <div class="pdf-running-footer">{_escape_html(r_footer)}</div>')

            html.append("</div>\n")

        html.append("""</div>
<script>
(function() {
  // 부모 창으로부터 특정 페이지 이동 메시지 수신 시 부드러운 스크롤 이동 실행함
  window.addEventListener('message', function(e) {
    if (e.data && e.data.type === 'scrollToPage') {
      var pno = e.data.page;
      var el = document.getElementById('page-' + pno);
      if (el) {
        el.scrollIntoView({ behavior: 'smooth', block: 'start' });
      }
    }
  });

  // 사용자가 HTML 내부를 스크롤할 때 현재 보이는 페이지 번호를 부모 창에 실시간 전송함
  var pageCards = document.querySelectorAll('.pdf-page-card');
  var lastReported = 1;
  var scrollTimer = null;
  window.addEventListener('scroll', function() {
    clearTimeout(scrollTimer);
    scrollTimer = setTimeout(function() {
      var current = 1;
      for (var i = 0; i < pageCards.length; i++) {
        var rect = pageCards[i].getBoundingClientRect();
        if (rect.top <= 220) {
          var m = pageCards[i].id.match(/page-(\\d+)/);
          if (m) current = parseInt(m[1], 10);
        } else {
          break;
        }
      }
      if (current !== lastReported) {
        lastReported = current;
        if (window.parent && window.parent !== window) {
          window.parent.postMessage({ type: 'pageScrolled', page: current }, '*');
        }
      }
    }, 60);
  }, { passive: true });
})();
</script>
</body>
</html>""")
        return "\n".join(html)

    # 워드(DOCX) 형식으로 변환하여 반환함
    def to_docx(self, output_path: Path) -> None:
        """
        @description 학습된 타이포그래피 토큰과 시맨틱 유형(Typology)을 적용하여 고충실도 Word(.docx) 문서를 생성함
        @param output_path: 저장할 docx 파일 경로임
        @return: 없음 (지정 경로로 파일 저장함)
        """
        doc = docx.Document()

        for section in doc.sections:
            section.top_margin = Inches(0.7)
            section.bottom_margin = Inches(0.7)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        tokens = getattr(self, "tokens", None)
        h1_size = getattr(tokens, "font_size_h1", 16.0) if tokens else 16.0
        h2_size = getattr(tokens, "font_size_h2", 13.0) if tokens else 13.0
        h3_size = getattr(tokens, "font_size_h3", 11.0) if tokens else 11.0
        body_size = getattr(tokens, "font_size_body", 9.5) if tokens else 9.5
        caption_size = getattr(tokens, "font_size_caption", 8.5) if tokens else 8.5
        p_color = getattr(tokens, "primary_color", "#1A365D") if tokens else "#1A365D"
        try:
            primary_rgb = RGBColor.from_string(p_color.lstrip("#"))
        except Exception:
            primary_rgb = RGBColor(26, 54, 93)

        for p in self.pages:
            pno = p["page_num"]
            if pno > 1:
                doc.add_page_break()

            r_header = p.get("running_header")
            if r_header and p.get("archetype") not in ("cover", "front_matter"):
                p_head = doc.add_paragraph()
                r_head = p_head.add_run(f"{r_header}   |   Page {pno}")
                r_head.font.name = "맑은 고딕"
                r_head.font.size = Pt(8.0)
                r_head.font.color.rgb = RGBColor(113, 128, 150)
                p_head.paragraph_format.space_after = Pt(8)

            for elem in p["elements"]:
                etype = elem["type"]
                stype = elem.get("semantic_type", "")

                # 런닝 헤더/푸터는 상단/하단 메타 컴포넌트로 처리하므로 본문 중복 출력 방지함
                if stype in ("running_header", "running_footer"):
                    continue

                if stype == "heading_l1":
                    txt = "\n".join(b["text"] for b in elem.get("blocks", [])).strip() if "blocks" in elem else str(elem.get("text", "")).strip()
                    p_h1 = doc.add_paragraph()
                    r_h1 = p_h1.add_run(txt)
                    r_h1.font.name = "맑은 고딕"
                    r_h1.font.bold = True
                    r_h1.font.size = Pt(max(13.0, h1_size))
                    r_h1.font.color.rgb = primary_rgb
                    p_h1.paragraph_format.space_before = Pt(12)
                    p_h1.paragraph_format.space_after = Pt(6)

                elif stype == "heading_l2":
                    txt = " ".join(b["text"] for b in elem.get("blocks", [])).strip() if "blocks" in elem else str(elem.get("text", "")).strip()
                    p_h2 = doc.add_paragraph()
                    r_h2 = p_h2.add_run(txt)
                    r_h2.font.name = "맑은 고딕"
                    r_h2.font.bold = True
                    r_h2.font.size = Pt(max(11.0, h2_size))
                    r_h2.font.color.rgb = primary_rgb
                    p_h2.paragraph_format.space_before = Pt(10)
                    p_h2.paragraph_format.space_after = Pt(4)

                elif stype == "heading_l3":
                    txt = " ".join(b["text"] for b in elem.get("blocks", [])).strip() if "blocks" in elem else str(elem.get("text", "")).strip()
                    p_h3 = doc.add_paragraph()
                    r_h3 = p_h3.add_run(txt)
                    r_h3.font.name = "맑은 고딕"
                    r_h3.font.bold = True
                    r_h3.font.size = Pt(max(10.0, h3_size))
                    p_h3.paragraph_format.space_before = Pt(6)
                    p_h3.paragraph_format.space_after = Pt(2)

                elif stype == "table_caption":
                    txt = str(elem.get("text", "")).strip() or "".join(b["text"] for b in elem.get("blocks", [])).strip()
                    p_cap = doc.add_paragraph()
                    r_cap = p_cap.add_run(txt)
                    r_cap.font.name = "맑은 고딕"
                    r_cap.font.bold = True
                    r_cap.font.size = Pt(caption_size)
                    r_cap.font.color.rgb = primary_rgb
                    p_cap.paragraph_format.space_before = Pt(8)
                    p_cap.paragraph_format.space_after = Pt(2)

                elif stype == "figure_caption":
                    txt = str(elem.get("text", "")).strip() or "".join(b["text"] for b in elem.get("blocks", [])).strip()
                    p_fcap = doc.add_paragraph()
                    p_fcap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r_fcap = p_fcap.add_run(txt)
                    r_fcap.font.name = "맑은 고딕"
                    r_fcap.font.bold = True
                    r_fcap.font.size = Pt(caption_size)
                    p_fcap.paragraph_format.space_before = Pt(4)
                    p_fcap.paragraph_format.space_after = Pt(6)

                elif stype in ("table_note", "figure_note"):
                    txt = str(elem.get("text", "")).strip() or "".join(b["text"] for b in elem.get("blocks", [])).strip()
                    p_note = doc.add_paragraph()
                    r_note = p_note.add_run(txt)
                    r_note.font.name = "맑은 고딕"
                    r_note.font.size = Pt(8.0)
                    r_note.font.italic = True
                    r_note.font.color.rgb = RGBColor(113, 128, 150)
                    p_note.paragraph_format.space_after = Pt(4)

                elif etype == "banner":
                    txt = "\n".join(b["text"] for b in elem["blocks"]).strip()
                    lines = [l.strip() for l in txt.split("\n") if l.strip()]
                    t1 = lines[0] if lines else "문서 제목"
                    t2 = " ".join(lines[1:]) if len(lines) > 1 else ""

                    tbl = doc.add_table(rows=1, cols=1)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    tbl.autofit = False
                    tbl.columns[0].width = Inches(6.8)

                    cell = tbl.cell(0, 0)
                    _docx_set_cell_background(cell, "1A365D")
                    _docx_set_cell_margins(cell, top=180, bottom=180, left=240, right=240)
                    _docx_set_cell_border(cell)

                    p_title = cell.paragraphs[0]
                    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run1 = p_title.add_run(t1)
                    run1.font.name = "맑은 고딕"
                    run1.font.size = Pt(14)
                    run1.font.bold = True
                    run1.font.color.rgb = RGBColor(255, 255, 255)

                    if t2:
                        p_sub = cell.add_paragraph()
                        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run2 = p_sub.add_run(t2)
                        run2.font.name = "맑은 고딕"
                        run2.font.size = Pt(9.5)
                        run2.font.color.rgb = RGBColor(203, 213, 225)

                    doc.add_paragraph().paragraph_format.space_after = Pt(4)

                elif etype == "alert":
                    txt = "\n".join(b["text"] for b in elem["blocks"]).strip()
                    tbl = doc.add_table(rows=1, cols=1)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    tbl.autofit = False
                    tbl.columns[0].width = Inches(6.8)

                    cell = tbl.cell(0, 0)
                    _docx_set_cell_background(cell, "EDF2F7")
                    _docx_set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
                    _docx_set_cell_border(cell, left={"sz": 24, "val": "single", "color": "3182CE"},
                                                top={"sz": 4, "val": "single", "color": "CBD5E1"},
                                                bottom={"sz": 4, "val": "single", "color": "CBD5E1"},
                                                right={"sz": 4, "val": "single", "color": "CBD5E1"})

                    p_alert = cell.paragraphs[0]
                    run_badge = p_alert.add_run("[안내] ")
                    run_badge.font.name = "맑은 고딕"
                    run_badge.font.size = Pt(9)
                    run_badge.font.bold = True
                    run_badge.font.color.rgb = RGBColor(49, 130, 206)

                    run_text = p_alert.add_run(txt)
                    run_text.font.name = "맑은 고딕"
                    run_text.font.size = Pt(9)
                    run_text.font.color.rgb = RGBColor(26, 54, 93)

                    doc.add_paragraph().paragraph_format.space_after = Pt(4)

                elif etype == "section_bar":
                    txt = " ".join(b["text"] for b in elem["blocks"]).strip()
                    tbl = doc.add_table(rows=1, cols=1)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    tbl.autofit = False
                    tbl.columns[0].width = Inches(6.8)

                    cell = tbl.cell(0, 0)
                    _docx_set_cell_background(cell, "CAD4DF")
                    _docx_set_cell_margins(cell, top=80, bottom=80, left=150, right=150)
                    _docx_set_cell_border(cell, left={"sz": 24, "val": "single", "color": "1A365D"})

                    p_sec = cell.paragraphs[0]
                    run_sec = p_sec.add_run(txt)
                    run_sec.font.name = "맑은 고딕"
                    run_sec.font.size = Pt(10.5)
                    run_sec.font.bold = True
                    run_sec.font.color.rgb = RGBColor(26, 54, 93)

                    doc.add_paragraph().paragraph_format.space_after = Pt(4)

                elif etype == "card":
                    txt = "\n".join(b["text"] for b in elem["blocks"]).strip()
                    lines = txt.split("\n")
                    c_title = lines[0] if lines else "상세 내역"
                    c_meta = ""
                    c_body_lines = lines[1:]

                    if len(lines) > 1 and any(k in lines[1] for k in ("접수일시", "일시", "담당자", "상담원:")) and len(lines[1]) < 60:
                        c_meta = lines[1]
                        c_body_lines = lines[2:]

                    tbl = doc.add_table(rows=1, cols=1)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    tbl.autofit = False
                    tbl.columns[0].width = Inches(6.8)

                    cell = tbl.cell(0, 0)
                    _docx_set_cell_background(cell, "FCFCFD")
                    _docx_set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
                    _docx_set_cell_border(cell, top={"sz": 4, "val": "single", "color": "E2E8EF"},
                                                bottom={"sz": 4, "val": "single", "color": "E2E8EF"},
                                                left={"sz": 4, "val": "single", "color": "E2E8EF"},
                                                right={"sz": 4, "val": "single", "color": "E2E8EF"})

                    p_title = cell.paragraphs[0]
                    run_title = p_title.add_run(c_title)
                    run_title.font.name = "맑은 고딕"
                    run_title.font.size = Pt(10)
                    run_title.font.bold = True
                    run_title.font.color.rgb = RGBColor(26, 54, 93)

                    if c_meta:
                        p_meta = cell.add_paragraph()
                        run_meta = p_meta.add_run(c_meta)
                        run_meta.font.name = "맑은 고딕"
                        run_meta.font.size = Pt(8.5)
                        run_meta.font.color.rgb = RGBColor(113, 128, 150)
                        p_meta.paragraph_format.space_after = Pt(4)

                    p_body = cell.add_paragraph()
                    p_body.paragraph_format.line_spacing = 1.2
                    body_text = "\n".join(c_body_lines)
                    run_body = p_body.add_run(body_text)
                    run_body.font.name = "맑은 고딕"
                    run_body.font.size = Pt(9)
                    run_body.font.color.rgb = RGBColor(45, 55, 72)

                    doc.add_paragraph().paragraph_format.space_after = Pt(4)

                elif etype == "image":
                    try:
                        from PIL import Image as PILImage
                        img_stream = io.BytesIO(elem["image_bytes"])
                        with PILImage.open(img_stream) as pil_img:
                            converted_stream = io.BytesIO()
                            if pil_img.mode in ("RGBA", "LA") or (pil_img.mode == "P" and "transparency" in pil_img.info):
                                pil_img.save(converted_stream, format="PNG")
                            else:
                                pil_img.convert("RGB").save(converted_stream, format="PNG")
                            converted_stream.seek(0)
                            doc.add_picture(converted_stream, width=Pt(min(450, max(1, elem['bbox'][2] - elem['bbox'][0]))))
                            p_img = doc.paragraphs[-1]
                            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            p_img.paragraph_format.space_after = Pt(8)
                    except Exception as exc:
                        logging.getLogger(__name__).warning("DOCX 이미지 추가 건너뜀 (page %s): %s", pno, exc)

                elif etype == "table":
                    rows_count = len(elem["rows"])
                    widths = _column_widths(elem, 450)
                    tbl = doc.add_table(rows=rows_count, cols=len(widths))
                    tbl.autofit = False
                    for r_i, row in enumerate(elem["rows"]):
                        for c_i, cell in enumerate(_row_cells(row)):
                            target = tbl.cell(r_i, c_i)
                            target.width = Pt(widths[c_i])
                            target.text = cell["text"]
                            _docx_set_cell_margins(target)
                            _docx_set_cell_border(target, **{edge: {"sz": 4, "val": "single", "color": "CBD5E1"} for edge in ("top", "bottom", "left", "right")})
                            if c_i == 0 and len(cell["text"]) < 40:
                                _docx_set_cell_background(target, "EDF2F6")
                            for paragraph in target.paragraphs:
                                for run in paragraph.runs:
                                    run.font.name = cell.get("font") or "맑은 고딕"
                                    run.font.size = Pt(cell.get("size") or 9)
                                    run.bold = bool(cell.get("bold", c_i == 0))
                                    color = cell.get("color", "#000000")
                                    if isinstance(color, int):
                                        color = int_color_to_hex(color)
                                    run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
                            for image in cell.get('inline_images', []):
                                try:
                                    from PIL import Image as PILImage
                                    width = min(widths[c_i], image['bbox'][2] - image['bbox'][0])
                                    with PILImage.open(io.BytesIO(image['image_bytes'])) as pil_img:
                                        converted_stream = io.BytesIO()
                                        if pil_img.mode in ("RGBA", "LA") or (pil_img.mode == "P" and "transparency" in pil_img.info):
                                            pil_img.save(converted_stream, format="PNG")
                                        else:
                                            pil_img.convert("RGB").save(converted_stream, format="PNG")
                                        converted_stream.seek(0)
                                        target.paragraphs[-1].add_run().add_picture(converted_stream, width=Pt(max(1, width)))
                                except Exception as exc:
                                    logging.getLogger(__name__).warning("DOCX 셀 인라인 이미지 추가 건너뜀 (page %s): %s", pno, exc)
                    doc.add_paragraph().paragraph_format.space_after = Pt(4)

                elif etype == "paragraph":
                    txt = "\n".join(b["text"] for b in elem["blocks"]).strip()
                    if txt:
                        p_para = doc.add_paragraph()
                        r_para = p_para.add_run(txt)
                        r_para.font.name = "맑은 고딕"
                        r_para.font.size = Pt(9.5)
                        p_para.paragraph_format.space_after = Pt(4)

                elif etype == "footer":
                    p_foot = doc.add_paragraph()
                    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r_foot = p_foot.add_run(elem.get("text", f"- {pno} -"))
                    r_foot.font.name = "맑은 고딕"
                    r_foot.font.size = Pt(9)
                    r_foot.font.color.rgb = RGBColor(113, 128, 150)

        doc.save(output_path)

    # excel 형식으로 변환하여 반환함
    def to_excel(self, output_path: Path) -> None:
        """Render 3-sheet Excel spreadsheet."""
        wb = openpyxl.Workbook()
        default_sheet = wb.active

        f_banner_title = Font(name="Malgun Gothic", size=13, bold=True, color="FFFFFF")
        f_banner_sub = Font(name="Malgun Gothic", size=9, bold=False, color="CBD5E1")
        f_alert = Font(name="Malgun Gothic", size=9, bold=False, color="1A365D")
        f_sec = Font(name="Malgun Gothic", size=10, bold=True, color="1A365D")
        f_card_title = Font(name="Malgun Gothic", size=10, bold=True, color="1A365D")
        f_card_meta = Font(name="Malgun Gothic", size=9, bold=False, color="718096")
        f_card_body = Font(name="Malgun Gothic", size=9, bold=False, color="2D3748")
        f_tbl_hdr = Font(name="Malgun Gothic", size=9, bold=True, color="1A365D")
        f_tbl_cell = Font(name="Malgun Gothic", size=9, bold=False, color="2D3748")
        f_footer = Font(name="Malgun Gothic", size=9, bold=False, color="718096")

        fill_banner = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
        fill_alert = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
        fill_sec = PatternFill(start_color="CAD4DF", end_color="CAD4DF", fill_type="solid")
        fill_card = PatternFill(start_color="FCFCFD", end_color="FCFCFD", fill_type="solid")
        fill_tbl_hdr = PatternFill(start_color="EDF2F6", end_color="EDF2F6", fill_type="solid")

        side_thin = Side(border_style="thin", color="CBD5E1")
        border_box = Border(left=side_thin, right=side_thin, top=side_thin, bottom=side_thin)
        border_alert = Border(left=Side(border_style="medium", color="3182CE"), right=side_thin, top=side_thin, bottom=side_thin)

        doc_max_cols = max(
            [4] + [len(_row_cells(r)) for p in self.pages for elem in p["elements"] if elem["type"] == "table" for r in elem.get("rows", [])]
        )

        # Sheet 1: 전체_문서_서식 (Visual Replica)
        ws1 = wb.create_sheet(title="전체_문서_서식")
        ws1.views.sheetView[0].showGridLines = True
        for c_idx in range(1, doc_max_cols + 1):
            ws1.column_dimensions[get_column_letter(c_idx)].width = 20 if c_idx % 2 == 1 else 38

        r1 = 1
        for p in self.pages:
            pno = p["page_num"]
            if pno > 1:
                r1 += 1
                ws1.cell(row=r1, column=1, value=f"--- [Page {pno}] ---").font = f_footer
                r1 += 1

            for elem in p["elements"]:
                etype = elem["type"]

                if etype == "banner":
                    txt = "\n".join(b["text"] for b in elem["blocks"]).strip()
                    lines = [l.strip() for l in txt.split("\n") if l.strip()]
                    t1 = lines[0] if lines else "문서 제목"
                    t2 = " ".join(lines[1:]) if len(lines) > 1 else ""

                    ws1.merge_cells(start_row=r1, start_column=1, end_row=r1, end_column=doc_max_cols)
                    c = ws1.cell(row=r1, column=1, value=t1)
                    c.font = f_banner_title; c.fill = fill_banner
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    ws1.row_dimensions[r1].height = 28
                    r1 += 1

                    if t2:
                        ws1.merge_cells(start_row=r1, start_column=1, end_row=r1, end_column=doc_max_cols)
                        c = ws1.cell(row=r1, column=1, value=t2)
                        c.font = f_banner_sub; c.fill = fill_banner
                        c.alignment = Alignment(horizontal="center", vertical="center")
                        ws1.row_dimensions[r1].height = 20
                        r1 += 1
                    r1 += 1

                elif etype == "alert":
                    txt = "\n".join(b["text"] for b in elem["blocks"]).strip()
                    ws1.merge_cells(start_row=r1, start_column=1, end_row=r1, end_column=doc_max_cols)
                    c = ws1.cell(row=r1, column=1, value=f"[안내] {txt}")
                    c.font = f_alert; c.fill = fill_alert; c.border = border_alert
                    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    ws1.row_dimensions[r1].height = 44
                    r1 += 2

                elif etype == "section_bar":
                    txt = " ".join(b["text"] for b in elem["blocks"]).strip()
                    ws1.merge_cells(start_row=r1, start_column=1, end_row=r1, end_column=doc_max_cols)
                    c = ws1.cell(row=r1, column=1, value=txt)
                    c.font = f_sec; c.fill = fill_sec
                    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
                    ws1.row_dimensions[r1].height = 24
                    r1 += 1

                elif etype == "card":
                    txt = "\n".join(b["text"] for b in elem["blocks"]).strip()
                    lines = txt.split("\n")
                    c_title = lines[0] if lines else "상세 내역"
                    c_body = "\n".join(lines[1:])

                    ws1.merge_cells(start_row=r1, start_column=1, end_row=r1, end_column=doc_max_cols)
                    c_t = ws1.cell(row=r1, column=1, value=c_title)
                    c_t.font = f_card_title; c_t.fill = fill_card; c_t.border = border_box
                    ws1.row_dimensions[r1].height = 22
                    r1 += 1

                    if c_body:
                        ws1.merge_cells(start_row=r1, start_column=1, end_row=r1, end_column=doc_max_cols)
                        c_b = ws1.cell(row=r1, column=1, value=c_body)
                        c_b.font = f_card_body; c_b.fill = fill_card; c_b.border = border_box
                        c_b.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                        ws1.row_dimensions[r1].height = max(40, min(120, len(lines) * 18))
                        r1 += 1
                    r1 += 1

                elif etype == "table":
                    for row in elem["rows"]:
                        for index, cell in enumerate(_row_cells(row), 1):
                            target = ws1.cell(r1, index, cell["text"])
                            _excel_inline_images(ws1, r1, index, cell.get('inline_images', []))
                            target.data_type = "s"
                            target.font = Font(name=cell.get("font") or "Malgun Gothic", size=cell.get("size") or 9,
                                               color=str(cell.get('color') or '#000000').lstrip('#'),
                                               bold=bool(cell.get("bold", index == 1)))
                            target.alignment = Alignment(wrap_text=True, vertical="top")
                            target.border = border_box
                            if index == 1:
                                target.fill = fill_tbl_hdr
                        r1 += 1
                    r1 += 1

                elif etype == "paragraph":
                    txt = "\n".join(b["text"] for b in elem["blocks"]).strip()
                    if txt:
                        ws1.merge_cells(start_row=r1, start_column=1, end_row=r1, end_column=doc_max_cols)
                        c = ws1.cell(row=r1, column=1, value=txt)
                        c.font = f_card_body
                        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                        r1 += 1

                elif etype == "footer":
                    ws1.merge_cells(start_row=r1, start_column=1, end_row=r1, end_column=doc_max_cols)
                    c = ws1.cell(row=r1, column=1, value=elem.get("text", f"- {pno} -"))
                    c.font = f_footer
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    r1 += 1

        # Sheet 2: 정형_데이터_테이블
        ws2 = wb.create_sheet(title="정형_데이터_테이블")
        ws2.views.sheetView[0].showGridLines = True
        for c_idx in range(1, doc_max_cols + 1):
            ws2.column_dimensions[get_column_letter(c_idx)].width = 22 if c_idx % 2 == 1 else 38

        ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=doc_max_cols)
        h_cell = ws2.cell(row=1, column=1, value="문서 추출 정형 데이터 테이블")
        h_cell.font = f_sec; h_cell.fill = fill_sec
        ws2.row_dimensions[1].height = 26

        r2 = 3
        for p in self.pages:
            for elem in p["elements"]:
                if elem["type"] == "table":
                    for row in elem["rows"]:
                        for index, cell in enumerate(_row_cells(row), 1):
                            target = ws2.cell(r2, index, cell["text"])
                            _excel_inline_images(ws2, r2, index, cell.get('inline_images', []))
                            target.data_type = "s"
                            target.font = Font(name=cell.get("font") or "Malgun Gothic", size=cell.get("size") or 9,
                                               color=str(cell.get('color') or '#000000').lstrip('#'),
                                               bold=bool(cell.get("bold", index == 1)))
                            target.alignment = Alignment(wrap_text=True, vertical="top")
                            target.border = border_box
                        r2 += 1

        # Sheet 3: 상담_비정형_데이터
        ws3 = wb.create_sheet(title="상담_비정형_데이터")
        ws3.views.sheetView[0].showGridLines = True
        ws3.column_dimensions["A"].width = 8
        ws3.column_dimensions["B"].width = 25
        ws3.column_dimensions["C"].width = 22
        ws3.column_dimensions["D"].width = 18
        ws3.column_dimensions["E"].width = 65

        rec_headers = ["번호", "상담/신청 구분", "접수일시", "담당자/신청인", "상세 상담 및 처리 내용"]
        for c_i, h in enumerate(rec_headers):
            c = ws3.cell(row=1, column=c_i+1, value=h)
            c.font = f_banner_title; c.fill = fill_banner
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws3.row_dimensions[1].height = 26

        r3 = 2
        idx = 1
        for p in self.pages:
            for elem in p["elements"]:
                if elem["type"] == "card":
                    txt = "\n".join(b["text"] for b in elem["blocks"]).strip()
                    lines = txt.split("\n")
                    title_line = lines[0] if lines else ""

                    date_val = ""
                    agent_val = ""
                    body_lines = lines[1:]

                    if len(lines) > 1 and ("일시" in lines[1] or "담당자" in lines[1]):
                        meta_line = lines[1]
                        body_lines = lines[2:]
                        m_date = re.search(r"(\d{4}-\d{2}-\d{2}\s+\d{2}:\d{2}:\d{2})", meta_line)
                        if m_date:
                            date_val = m_date.group(1)
                        m_agent = re.search(r"담당자:\s*([^|\n]+)", meta_line)
                        if m_agent:
                            agent_val = m_agent.group(1).strip()

                    content_val = "\n".join(body_lines)

                    ws3.cell(row=r3, column=1, value=idx).alignment = Alignment(horizontal="center")
                    ws3.cell(row=r3, column=2, value=title_line).font = f_card_title
                    ws3.cell(row=r3, column=3, value=date_val).alignment = Alignment(horizontal="center")
                    ws3.cell(row=r3, column=4, value=agent_val).alignment = Alignment(horizontal="center")
                    c_body = ws3.cell(row=r3, column=5, value=content_val)
                    c_body.alignment = Alignment(wrap_text=True)

                    for c_i in range(1, 6):
                        ws3.cell(row=r3, column=c_i).border = border_box

                    ws3.row_dimensions[r3].height = max(35, min(100, len(body_lines) * 16))
                    r3 += 1
                    idx += 1

        if default_sheet in wb.worksheets:
            wb.remove(default_sheet)

        wb.save(output_path)

    # 한글 표준(HWPX) 형식으로 변환하여 반환함
    def to_hwpx(self, output_path: Path) -> None:
        """
        @description 공식 python-hwpx 엔진을 활용하여 표준 HWPX 문서로 변환 저장함
        @param output_path: 저장할 hwpx 파일 경로임
        @return: 없음 (지정 경로로 파일 저장함)
        """
        from hwpx.document import HwpxDocument

        doc = HwpxDocument.new()
        PAGE_WIDTH = 42520  # standard printable width

        for p in self.pages:
            pno = p["page_num"]

            r_header = p.get("running_header")
            if r_header and p.get("archetype") not in ("cover", "front_matter"):
                doc.add_paragraph(f"{r_header}   |   p. {pno}")

            for elem in p["elements"]:
                etype = elem["type"]
                stype = elem.get("semantic_type", "")

                # 런닝 헤더/푸터는 상단/하단 메타 컴포넌트로 처리하므로 본문 중복 출력 방지함
                if stype in ("running_header", "running_footer"):
                    continue

                if stype == "heading_l1":
                    txt = "\n".join(b["text"] for b in elem.get("blocks", [])).strip() if "blocks" in elem else str(elem.get("text", "")).strip()
                    doc.add_paragraph(f"■ {txt}")

                elif stype == "heading_l2":
                    txt = " ".join(b["text"] for b in elem.get("blocks", [])).strip() if "blocks" in elem else str(elem.get("text", "")).strip()
                    doc.add_paragraph(f"▶ {txt}")

                elif stype == "heading_l3":
                    txt = " ".join(b["text"] for b in elem.get("blocks", [])).strip() if "blocks" in elem else str(elem.get("text", "")).strip()
                    doc.add_paragraph(f"● {txt}")

                elif stype == "table_caption":
                    txt = str(elem.get("text", "")).strip() or "".join(b["text"] for b in elem.get("blocks", [])).strip()
                    doc.add_paragraph(txt)

                elif stype == "figure_caption":
                    txt = str(elem.get("text", "")).strip() or "".join(b["text"] for b in elem.get("blocks", [])).strip()
                    doc.add_paragraph(txt)

                elif stype in ("table_note", "figure_note"):
                    txt = str(elem.get("text", "")).strip() or "".join(b["text"] for b in elem.get("blocks", [])).strip()
                    doc.add_paragraph(f"* {txt}")

                elif etype == "banner":
                    txt = "\n".join(b["text"] for b in elem["blocks"]).strip()
                    lines = [l.strip() for l in txt.split("\n") if l.strip()]
                    t1 = lines[0] if lines else "문서 제목"
                    t2 = " ".join(lines[1:]) if len(lines) > 1 else ""

                    tbl = doc.add_table(rows=1, cols=1)
                    tbl.set_column_widths([PAGE_WIDTH])
                    tbl.set_cell_shading(0, 0, "#1A365D")
                    banner_txt = f"■ {t1}\n{t2}" if t2 else f"■ {t1}"
                    tbl.set_cell_text(0, 0, banner_txt)

                elif etype == "alert":
                    txt = "\n".join(b["text"] for b in elem["blocks"]).strip()
                    tbl = doc.add_table(rows=1, cols=1)
                    tbl.set_column_widths([PAGE_WIDTH])
                    tbl.set_cell_shading(0, 0, "#EDF2F7")
                    tbl.set_cell_text(0, 0, f"[안내사항]\n{txt}")

                elif etype == "section_bar":
                    txt = " ".join(b["text"] for b in elem["blocks"]).strip()
                    tbl = doc.add_table(rows=1, cols=1)
                    tbl.set_column_widths([PAGE_WIDTH])
                    tbl.set_cell_shading(0, 0, "#CAD4DF")
                    tbl.set_cell_text(0, 0, f"▶ {txt}")

                elif etype == "card":
                    txt = "\n".join(b["text"] for b in elem["blocks"]).strip()
                    lines = txt.split("\n")
                    c_title = lines[0] if lines else "상세 기록"
                    c_meta = ""
                    c_body_lines = lines[1:]

                    if len(lines) > 1 and any(k in lines[1] for k in ("접수일시", "일시", "담당자", "상담원:")) and len(lines[1]) < 60:
                        c_meta = lines[1]
                        c_body_lines = lines[2:]

                    card_content = f"【{c_title}】"
                    if c_meta:
                        card_content += f"\n접수 정보: {c_meta}"
                    if c_body_lines:
                        card_content += "\n\n" + "\n".join(c_body_lines)

                    tbl = doc.add_table(rows=1, cols=1)
                    tbl.set_column_widths([PAGE_WIDTH])
                    tbl.set_cell_shading(0, 0, "#FCFCFD")
                    tbl.set_cell_text(0, 0, card_content)

                elif etype == "table":
                    rows_data = elem["rows"]
                    widths = _column_widths(elem, PAGE_WIDTH)
                    tbl = doc.add_table(rows=len(rows_data), cols=len(widths))
                    widths = [round(w) for w in widths]
                    widths[-1] += PAGE_WIDTH - sum(widths)
                    tbl.set_column_widths(widths)
                    for r_idx, row in enumerate(rows_data):
                        for c_idx, cell in enumerate(_row_cells(row)):
                            target = tbl.cell(r_idx, c_idx)
                            paragraph = target.paragraphs[0] if target.paragraphs else target.add_paragraph("")
                            paragraph.add_run(cell["text"], bold=bool(cell.get("bold", c_idx == 0)),
                                              font=cell.get("font"), size=cell.get("size"), color=cell.get("color"),
                                              expand_special_characters=True)
                            for image in cell.get('inline_images', []):
                                binary = doc.media.add_image(image['image_bytes'], image.get('format', 'png'))
                                width = max(1, min(widths[c_idx], round((image['bbox'][2] - image['bbox'][0]) * 100)))
                                height = max(1, round(width * (image['bbox'][3] - image['bbox'][1]) / max(1, image['bbox'][2] - image['bbox'][0])))
                                paragraph.add_picture(binary.item_id, width=width, height=height)
                            if c_idx == 0:
                                tbl.set_cell_shading(r_idx, c_idx, "#EDF2F6")

                elif etype == "paragraph":
                    txt = "\n".join(b["text"] for b in elem["blocks"]).strip()
                    if txt:
                        doc.add_paragraph(txt)

                elif etype == "image":
                    width_mm = min(PAGE_WIDTH / (7200 / 25.4), (elem['bbox'][2] - elem['bbox'][0]) * 25.4 / 72)
                    doc.add_picture(elem['image_bytes'], elem.get('format', 'png'), width_mm=max(0.1, width_mm))

                elif etype == "footer":
                    doc.add_paragraph(elem.get('text', f"- {pno} -"))

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save_to_path(output_path)

    # 마크다운 형식으로 변환하여 반환함
    def to_markdown(self) -> str:
        """Render clean, structured Markdown with typology awareness."""
        md_parts = []
        for p in self.pages:
            pno = p["page_num"]
            md_parts.append(f"## Page {pno}\n")

            r_header = p.get("running_header")
            if r_header and p.get("archetype") not in ("cover", "front_matter"):
                md_parts.append(f"*{r_header}*\n\n")

            for elem in p["elements"]:
                etype = elem["type"]
                stype = elem.get("semantic_type", "")

                # 런닝 헤더/푸터는 페이지 상단/하단 메타로 처리되므로 본문 중복 출력 방지함
                if stype in ("running_header", "running_footer"):
                    continue

                if stype == "heading_l1" or etype == "banner":
                    txt = "\n".join(b["text"] for b in elem.get("blocks", [])).strip() if "blocks" in elem else str(elem.get("text", "")).strip()
                    md_parts.append(f"# {txt}\n\n")
                elif stype == "heading_l2" or etype == "section_bar":
                    txt = " ".join(b["text"] for b in elem.get("blocks", [])).strip() if "blocks" in elem else str(elem.get("text", "")).strip()
                    md_parts.append(f"## {txt}\n\n")
                elif stype == "heading_l3":
                    txt = " ".join(b["text"] for b in elem.get("blocks", [])).strip() if "blocks" in elem else str(elem.get("text", "")).strip()
                    md_parts.append(f"### {txt}\n\n")
                elif stype == "table_caption":
                    txt = str(elem.get("text", "")).strip() or "".join(b["text"] for b in elem.get("blocks", [])).strip()
                    md_parts.append(f"**{txt}**\n\n")
                elif stype == "figure_caption":
                    txt = str(elem.get("text", "")).strip() or "".join(b["text"] for b in elem.get("blocks", [])).strip()
                    md_parts.append(f"*{txt}*\n\n")
                elif stype in ("table_note", "figure_note"):
                    txt = str(elem.get("text", "")).strip() or "".join(b["text"] for b in elem.get("blocks", [])).strip()
                    md_parts.append(f"> *{txt}*\n\n")
                elif etype == "alert" or stype == "callout_box":
                    txt = "\n".join(b["text"] for b in elem.get("blocks", [])).strip() if "blocks" in elem else str(elem.get("text", "")).strip()
                    md_parts.append(f"> **[안내사항]** {txt}\n\n")
                elif etype == "card":
                    txt = "\n".join(b["text"] for b in elem.get("blocks", [])).strip() if "blocks" in elem else str(elem.get("text", "")).strip()
                    md_parts.append(f"```text\n{txt}\n```\n\n")
                elif etype == "table":
                    rows_data = elem["rows"]
                    count = max(len(_row_cells(row)) for row in rows_data)
                    first_row_cells = [_escape_md_cell(c["text"]) + ''.join(_inline_image_html(image) for image in c.get('inline_images', [])) for c in _row_cells(rows_data[0])]
                    if len(rows_data) >= 2 and any(c.strip() for c in first_row_cells):
                        t_rows = [
                            "| " + " | ".join(first_row_cells + [""] * (count - len(first_row_cells))) + " |",
                            "| " + " | ".join("---" for _ in range(count)) + " |",
                        ]
                        for row in rows_data[1:]:
                            values = [_escape_md_cell(c["text"]) + ''.join(_inline_image_html(image) for image in c.get('inline_images', [])) for c in _row_cells(row)]
                            t_rows.append("| " + " | ".join(values + [""] * (count - len(values))) + " |")
                    else:
                        t_rows = ["| " + " | ".join(f"열 {i+1}" for i in range(count)) + " |",
                                  "| " + " | ".join("---" for _ in range(count)) + " |"]
                        for row in rows_data:
                            values = [_escape_md_cell(c["text"]) + ''.join(_inline_image_html(image) for image in c.get('inline_images', [])) for c in _row_cells(row)]
                            t_rows.append("| " + " | ".join(values + [""] * (count - len(values))) + " |")
                    md_parts.append("\n".join(t_rows) + "\n\n")
                elif etype == "paragraph":
                    txt = "\n".join(b["text"] for b in elem.get("blocks", [])).strip() if "blocks" in elem else str(elem.get("text", "")).strip()
                    if txt:
                        if stype == "list_bullet" and not txt.startswith("- ") and not txt.startswith("* "):
                            md_parts.append(f"- {txt}\n\n")
                        else:
                            md_parts.append(f"{txt}\n\n")

                elif etype == 'image':
                    encoded = base64.b64encode(elem['image_bytes']).decode('ascii')
                    mime = 'image/jpeg' if elem.get('format') in ('jpg', 'jpeg') else 'image/png'
                    md_parts.append(f'![Document image](data:{mime};base64,{encoded})\n\n')

            footer_str = p.get("running_footer") or p.get("footer_text") or f"- {pno} -"
            md_parts.append(f"*{footer_str}*\n\n---\n")

        return "\n".join(md_parts).strip()


# PDF 문서 to high 충실도 HTML 웹 문서 데이터를 대상 포맷으로 변환함
def convert_pdf_to_high_fidelity_html(pdf_path: Path, title: Optional[str] = None) -> str:
    """Convert PDF to high-fidelity responsive HTML (95%+ visual match)."""
    from synthetic_engine.exporters.ocr_table_reconstructor import (
        is_scanned_pdf, extract_scanned_pdf_pages, convert_ocr_result_to_html
    )
    if is_scanned_pdf(pdf_path):
        pages = extract_scanned_pdf_pages(pdf_path)
        return convert_ocr_result_to_html(pages)

    doc = HighFidelityPdfDoc(pdf_path)
    try:
        return doc.to_html(title=title)
    finally:
        doc.close()


# PDF 문서 to high 충실도 워드(DOCX) 데이터를 대상 포맷으로 변환함
def convert_pdf_to_high_fidelity_docx(pdf_path: Path, output_path: Path) -> None:
    """Convert PDF to high-fidelity Word (.docx) document."""
    from synthetic_engine.exporters.ocr_table_reconstructor import (
        is_scanned_pdf, extract_scanned_pdf_pages, convert_ocr_result_to_docx
    )
    if is_scanned_pdf(pdf_path):
        pages = extract_scanned_pdf_pages(pdf_path)
        convert_ocr_result_to_docx(pages, output_path)
        return

    doc = HighFidelityPdfDoc(pdf_path)
    try:
        doc.to_docx(output_path)
    finally:
        doc.close()


# PDF 문서 to high 충실도 excel 데이터를 대상 포맷으로 변환함
def convert_pdf_to_high_fidelity_excel(pdf_path: Path, output_path: Path) -> None:
    """Convert PDF to styled Excel (.xlsx) spreadsheet with 3 distinct sheets."""
    doc = HighFidelityPdfDoc(pdf_path)
    try:
        doc.to_excel(output_path)
    finally:
        doc.close()


# PDF 문서 to high 충실도 한글 표준(HWPX) 데이터를 대상 포맷으로 변환함
def convert_pdf_to_high_fidelity_hwpx(pdf_path: Path, output_path: Path) -> None:
    """Convert PDF to standard HWPX document."""
    from synthetic_engine.exporters.ocr_table_reconstructor import (
        is_scanned_pdf, extract_scanned_pdf_pages, convert_ocr_result_to_hwpx
    )
    if is_scanned_pdf(pdf_path):
        pages = extract_scanned_pdf_pages(pdf_path)
        convert_ocr_result_to_hwpx(pages, output_path)
        return

    doc = HighFidelityPdfDoc(pdf_path)
    try:
        doc.to_hwpx(output_path)
    finally:
        doc.close()


# PDF 문서 to high 충실도 한글(HWP) 데이터를 대상 포맷으로 변환함
def convert_pdf_to_high_fidelity_hwp(pdf_path: Path, output_path: Path) -> None:
    """Convert PDF to binary HWP through an intermediate HWPX package and Hancom COM."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as tmpdir:
        intermediate_hwpx = Path(tmpdir) / f"{Path(pdf_path).stem}.hwpx"
        convert_pdf_to_high_fidelity_hwpx(pdf_path, intermediate_hwpx)

        from synthetic_engine.common.com_session import win32_com_session
        try:
            with win32_com_session("HWPFrame.HwpObject") as hwp:
                opened = hwp.Open(str(intermediate_hwpx.resolve()), "HWPX", "versionwarning:False;forcedopen:True")
                if not opened:
                    opened = hwp.Open(str(intermediate_hwpx.resolve()))
                if not opened:
                    raise RuntimeError("한컴오피스에서 중간 HWPX 문서를 열지 못했습니다.")
                saved = hwp.SaveAs(str(output_path.resolve()), "HWP", "")
                if not saved or not output_path.exists() or output_path.stat().st_size == 0:
                    raise RuntimeError("한컴오피스 HWP 저장에 실패했습니다.")
        except ImportError as exc:
            raise RuntimeError("HWP 저장에는 Windows 한컴오피스 COM 환경이 필요합니다. HWPX 출력을 사용하세요.") from exc


# PDF 문서 to high 충실도 마크다운 데이터를 대상 포맷으로 변환함
def convert_pdf_to_high_fidelity_markdown(pdf_path: Path) -> str:
    """Convert PDF to structured Markdown."""
    from synthetic_engine.exporters.ocr_table_reconstructor import (
        is_scanned_pdf, extract_scanned_pdf_pages, convert_ocr_result_to_markdown
    )
    if is_scanned_pdf(pdf_path):
        pages = extract_scanned_pdf_pages(pdf_path)
        return convert_ocr_result_to_markdown(pages)

    doc = HighFidelityPdfDoc(pdf_path)
    try:
        return doc.to_markdown()
    finally:
        doc.close()

