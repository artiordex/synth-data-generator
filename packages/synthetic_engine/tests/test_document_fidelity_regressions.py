# -*- coding: utf-8 -*-
# =============================================================================
# 파일명: test_document_fidelity_regressions.py
# 경로: packages/synthetic_engine/tests/test_document_fidelity_regressions.py
# 목적: 문서 포맷 변환 충실도 회귀 방지 테스트를 수행함.
# 작성자: AI Agent
# 작성일: 2026-09-13
# 수정일: 2026-09-13
# =============================================================================
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

import docx
import openpyxl
import pytest
from bs4 import BeautifulSoup

from synthetic_engine.exporters.pdf_high_fidelity_converter import HighFidelityPdfDoc
from synthetic_engine.exporters.hwp_high_fidelity_hwpx_converter import (
    _build_hwpx_from_html, _extract_col_widths_hwpunit, _html_table_grid,
    markdown_text_to_hwpx,
)
from synthetic_engine.exporters.hwp_high_fidelity_docx_converter import _resolve_style, HwpHtmlToDocxBuilder


# document six 컬럼 목록 all renderers 기능의 정상 동작 및 제약조건을 테스트함
def test_document_six_columns_all_renderers(tmp_path):
    model = HighFidelityPdfDoc.__new__(HighFidelityPdfDoc)
    model.pdf_path = Path('six.pdf')
    values = ['label', 'line1\nline2', 'third', 'fourth', 'fifth', 'sixth']
    row = model._cells_to_table_row([{'text': text, 'bold': True, 'font': 'Arial', 'size': 11, 'color': '#123456'} for text in values])
    assert row['col_count'] == 6
    model.pages = [{'page_num': 1, 'elements': [{'type': 'table', 'rows': [row], 'col_edges': [0, 40, 120, 160, 200, 240, 300]}]}]
    html = BeautifulSoup(model.to_html(), 'html.parser')
    assert len(html.find('table').find_all(['td', 'th'])) == 6
    assert 'line1<br/>line2' in str(html)
    assert 'sixth' in model.to_markdown()
    model.to_docx(tmp_path / 'six.docx')
    table = docx.Document(tmp_path / 'six.docx').tables[0]
    assert [c.text for c in table.rows[0].cells] == values
    model.to_excel(tmp_path / 'six.xlsx')
    wb = openpyxl.load_workbook(tmp_path / 'six.xlsx')
    assert any(tuple(values) == tuple(row[:6]) for ws in wb for row in ws.iter_rows(values_only=True))
    model.to_hwpx(tmp_path / 'six.hwpx')
    with zipfile.ZipFile(tmp_path / 'six.hwpx') as archive:
        xml = archive.read('Contents/section0.xml')
        root = ET.fromstring(xml)
        assert any(e.attrib.get('colCnt') == '6' for e in root.iter())
        assert 'sixth' in xml.decode()


# document nested 표(테이블) 너비 목록 and 병합 기능의 정상 동작 및 제약조건을 테스트함
def test_document_nested_table_widths_and_merge(tmp_path):
    source = '<table><col style="width:20%"/><col style="width:80%"/><tr><td rowspan="2">outer<table><tr><td>inner</td><td>value</td></tr></table>end</td><td>A</td></tr><tr><td>B</td></tr></table>'
    table = BeautifulSoup(source, 'html.parser').table
    assert _extract_col_widths_hwpunit(table, 48000) == [9600, 38400]
    path = _build_hwpx_from_html(source, tmp_path / 'nested.hwpx')
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read('Contents/section0.xml'))
        assert len([e for e in root.iter() if e.tag.endswith('}tbl')]) == 2
        for value in ['outer', 'inner', 'value', 'end', 'A', 'B']:
            assert value in ''.join(root.itertext())


# document overlapping 스팬 목록 rejected 기능의 정상 동작 및 제약조건을 테스트함
def test_document_overlapping_spans_rejected():
    table = BeautifulSoup('<table><tr><td>A</td><td rowspan="2">B</td></tr><tr><td colspan="2">C</td></tr></table>', 'html.parser').table
    with pytest.raises(ValueError, match='Overlapping'):
        _html_table_grid(table)


# document 마크다운 runs 기능의 정상 동작 및 제약조건을 테스트함
def test_document_markdown_runs(tmp_path):
    path = markdown_text_to_hwpx('plain **bold** *italic* `code`', tmp_path / 'inline.hwpx')
    with zipfile.ZipFile(path) as archive:
        text = ''.join(ET.fromstring(archive.read('Contents/section0.xml')).itertext())
        assert 'plain bold italic code' in text
        assert b'bold' in archive.read('Contents/header.xml')


# document css cascade 기능의 정상 동작 및 제약조건을 테스트함
def test_document_css_cascade():
    element = BeautifulSoup('<div style="color:red"><p class="a" style="font-size:12pt">text</p></div>', 'html.parser').p
    style = _resolve_style(element, {'p': {'font-size': '8pt'}, '.a': {'font-size': '10pt', 'font-weight': 'bold'}})
    assert style['font-size'] == '12pt'
    assert style['color'] == 'red'
    assert style['font-weight'] == 'bold'


# document 워드(DOCX) mixed 텍스트 and nested 표(테이블) 기능의 정상 동작 및 제약조건을 테스트함
def test_document_docx_mixed_text_and_nested_table(tmp_path):
    source = tmp_path / 'mixed.html'
    source.write_text('<html><body><p>before <b>bold <i>inside</i></b> after<br/>next</p><table><tr><td>outer<table><tr><td>inner</td></tr></table>end</td></tr></table></body></html>', encoding='utf-8')
    target = tmp_path / 'mixed.docx'
    HwpHtmlToDocxBuilder(source, None, None).build(target)
    document = docx.Document(target)
    assert document.paragraphs[0].text == 'before bold inside after\nnext'
    cell = document.tables[0].cell(0, 0)
    assert cell.tables[0].cell(0, 0).text == 'inner'
    assert 'outer' in cell.text and 'end' in cell.text


# document PDF 문서 repeated 이미지 목록 한글 표준(HWPX) 기능의 정상 동작 및 제약조건을 테스트함
def test_document_pdf_repeated_images_hwpx(tmp_path):
    import pymupdf as fitz
    source = tmp_path / 'images.pdf'
    with fitz.open() as pdf:
        page = pdf.new_page()
        pixels = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 10, 10), False)
        pixels.clear_with(90)
        image = pixels.tobytes('png')
        page.insert_image(fitz.Rect(20, 20, 40, 40), stream=image)
        page.insert_image(fitz.Rect(60, 20, 80, 40), stream=image)
        pdf.save(source)
    model = HighFidelityPdfDoc(source)
    try:
        assert len([e for e in model.pages[0]['elements'] if e['type'] == 'image']) == 2
        model.to_hwpx(tmp_path / 'images.hwpx')
        with zipfile.ZipFile(tmp_path / 'images.hwpx') as archive:
            xml = ET.fromstring(archive.read('Contents/section0.xml'))
            assert len([e for e in xml.iter() if e.tag.endswith('}pic')]) == 2
    finally:
        model.close()
