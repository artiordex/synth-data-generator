# -*- coding: utf-8 -*-
# =============================================================================
# 파일명: test_pipeline.py
# 경로: packages/synthetic_engine/tests/document_conversion/test_pipeline.py
# 목적: 전체 문서 변환 파이프라인 실행 흐름을 테스트함.
# 작성자: AI Agent
# 작성일: 2026-09-13
# 수정일: 2026-09-13
# =============================================================================
import json
from zipfile import ZipFile
import pytest
from synthetic_engine.document_conversion.pipeline import convert_document
from synthetic_engine.document_conversion.exceptions import UnsupportedFeatureError
from synthetic_engine.document_conversion.parsers.hwpx_parser import HP, HH, HS


# sample 작업을 수행함
def sample(path):
    with ZipFile(path, 'w') as archive:
        archive.writestr('Contents/header.xml', f'<head xmlns="{HH}"/>')
        archive.writestr('Contents/section0.xml', f'<sec xmlns="{HS}" xmlns:p="{HP}"><p:p><p:run><p:t> A\u00a0\u3000①食藥處<p:tab/>B</p:t></p:run></p:p></sec>')
    return path


# 한글 표준(HWPX) ir render and 분석 리포트 기능의 정상 동작 및 제약조건을 테스트함
@pytest.mark.parametrize('target', ['html', 'docx'])
def test_hwpx_ir_render_and_report(tmp_path, target):
    source = sample(tmp_path/'sample.hwpx')
    result = convert_document(source, target)
    assert result.exists()
    report = json.loads(result.with_suffix('.conversion-report.json').read_text(encoding='utf-8'))
    assert report['unsupported_features']
    if target == 'html':
        assert report['text_fidelity'] == 1.0
        assert '食藥處' in result.read_text(encoding='utf-8')
    else:
        from docx import Document
        assert Document(result).paragraphs[0].text == ' A\u00a0\u3000①食藥處\tB'


# strict failure does not publish 기능의 정상 동작 및 제약조건을 테스트함
def test_strict_failure_does_not_publish(tmp_path):
    source = sample(tmp_path/'sample.hwpx')
    with pytest.raises(UnsupportedFeatureError):
        convert_document(source, 'html', strict=True)
    assert not list(tmp_path.glob('*.html'))


# empty renderer output does not publish 기능의 정상 동작 및 제약조건을 테스트함
def test_empty_renderer_output_does_not_publish(tmp_path, monkeypatch):
    source = sample(tmp_path/'sample.hwpx')

    class EmptyRenderer:
        capabilities = type('Capabilities', (), {})()

        # render 작업을 수행함
        def render(self, document, output_path):
            output_path.write_bytes(b'')

    monkeypatch.setattr('synthetic_engine.document_conversion.pipeline.renderer_for', lambda _: EmptyRenderer())
    with pytest.raises(ValueError, match='empty output'):
        convert_document(source, 'html')
    assert not list(tmp_path.glob('*.html'))


# unparseable renderer output does not publish 기능의 정상 동작 및 제약조건을 테스트함
def test_unparseable_renderer_output_does_not_publish(tmp_path, monkeypatch):
    source = sample(tmp_path/'sample.hwpx')

    class InvalidRenderer:
        capabilities = type('Capabilities', (), {})()

        # render 작업을 수행함
        def render(self, document, output_path):
            output_path.write_text('not an html document', encoding='utf-8')

    monkeypatch.setattr('synthetic_engine.document_conversion.pipeline.renderer_for', lambda _: InvalidRenderer())
    with pytest.raises(Exception):
        convert_document(source, 'html')
    assert not list(tmp_path.glob('*.html'))
    assert not list(tmp_path.glob('*.conversion-report.json'))


# standalone public api uses workspace engine 기능의 정상 동작 및 제약조건을 테스트함
def test_standalone_public_api_uses_workspace_engine(tmp_path):
    from docengine import convert_document as public_convert
    output = public_convert(sample(tmp_path / 'public.hwpx'), 'html')
    assert output.is_file()
    assert output.with_suffix('.conversion-report.json').is_file()


# real 한글 표준(HWPX) package to 워드(DOCX) 표(테이블) roundtrip 기능의 정상 동작 및 제약조건을 테스트함
def test_real_hwpx_package_to_docx_table_roundtrip(tmp_path):
    from hwpx.document import HwpxDocument
    from synthetic_engine.document_conversion.parsers.docx_parser import DocxParser
    from synthetic_engine.document_conversion.core.ir import TableIR
    source = tmp_path / 'actual.hwpx'
    doc = HwpxDocument.new()
    doc.add_paragraph('한글\u00a0① 원문')
    table = doc.add_table(rows=2, cols=2)
    table.set_cell_text(0, 0, '가')
    table.set_cell_text(0, 1, '나')
    table.set_cell_text(1, 0, '다')
    table.set_cell_text(1, 1, '라')
    doc.save_to_path(source)
    result = convert_document(source, 'docx')
    parsed = DocxParser().parse(result)
    tables = [b for b in parsed.iter_blocks() if isinstance(b, TableIR)]
    assert len(tables) == 1
    assert sum(len(row) for row in tables[0].rows) == 4
