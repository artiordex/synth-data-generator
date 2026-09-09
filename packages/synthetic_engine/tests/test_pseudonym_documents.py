# -*- coding: utf-8 -*-
from __future__ import annotations

import tempfile
from pathlib import Path
import pandas as pd

from synthetic_engine import (
    read_table,
    scan_pii_columns,
    ColumnPlan,
    apply_pii,
    export_pseudonymized_document
)
from synthetic_engine.privacy.masker import SmartMasker


def test_smart_masker_full_text():
    sample_text = (
        "작성자 홍길동(주민번호 900101-1234567, 전화번호 010-1234-5678, 이메일 hong@test.com)은 "
        "서울특별시 강남구 테헤란로 152 23층에 거주하고 있습니다."
    )
    masked = SmartMasker.mask_full_text(sample_text)
    
    assert "900101-1******" in masked
    assert "010-****-5678" in masked
    assert "hon*****@test.com" in masked
    assert "서울특별시 강남구 ********" in masked


def test_markdown_document_pseudonymization_and_export():
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)
        md_file = tmp_path / "sample_doc.md"
        md_file.write_text(
            "# 개인정보 서류\n\n"
            "본 문서는 홍길동 (주민번호: 950215-1098765, 연락처: 010-9876-5432) 님의 신청서입니다.\n\n"
            "이메일 문의: gildong@company.or.kr",
            encoding="utf-8"
        )

        df = read_table(md_file)
        assert not df.empty
        assert "문서_내용" in df.columns or "컬럼_1" in df.columns

        pii_detected = scan_pii_columns(df)
        assert len(pii_detected) > 0

        plan = ColumnPlan(
            categorical=[],
            numerical=[],
            ignored=[],
            pii={col: {"action": "mask", "faker": "unstructured_text", "pii_type": "unstructured_text"} for col in df.columns},
            rules={}
        )
        pseudo_df, summary = apply_pii(df, plan)

        all_text = " ".join(pseudo_df.astype(str).values.flatten())
        assert "950215-1******" in all_text or "010-****-5432" in all_text

        for fmt in ["pdf", "docx", "hwpx", "md", "txt", "csv", "xlsx"]:
            out_file = tmp_path / f"out_pseudo.{fmt}"
            exported = export_pseudonymized_document(pseudo_df, target_fmt=fmt, output_path=out_file)
            assert exported.exists()
            assert exported.stat().st_size > 0


def test_document_dataframe_export_formats():
    df = pd.DataFrame({
        "성명": ["홍길동", "김철수"],
        "주민번호": ["900101-1******", "850312-2******"],
        "전화번호": ["010-****-1234", "010-****-5678"],
        "소속": ["기술혁신팀", "데이터분석팀"]
    })

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)
        for fmt in ["pdf", "docx", "hwpx", "hwp", "md", "txt", "csv", "xlsx", "json", "parquet"]:
            out_file = tmp_path / f"test_export.{fmt}"
            exported = export_pseudonymized_document(df, target_fmt=fmt, output_path=out_file, original_filename="고객목록.xlsx")
            assert exported.exists()
            assert exported.stat().st_size > 0


def test_in_place_docx_replacement():
    import docx
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)
        docx_file = tmp_path / "original.docx"

        doc = docx.Document()
        doc.add_heading("원본 개인정보 문서", level=1)
        doc.add_paragraph("담당자 홍길동 님의 연락처는 010-1234-5678 입니다.")
        doc.save(docx_file)

        out_file = tmp_path / "pseudonymized.docx"
        df = pd.DataFrame({"문서_내용": ["담당자 홍*동 님의 연락처는 010-****-5678 입니다."]})
        replacements = [
            ("홍길동", "홍*동"),
            ("010-1234-5678", "010-****-5678")
        ]

        exported = export_pseudonymized_document(
            df,
            target_fmt="docx",
            output_path=out_file,
            original_filepath=docx_file,
            replacements=replacements
        )
        assert exported.exists()
        
        # Verify in-place replaced text
        res_doc = docx.Document(exported)
        full_text = "\n".join(p.text for p in res_doc.paragraphs)
        assert "홍*동" in full_text
        assert "010-****-5678" in full_text
        assert "홍길동" not in full_text
