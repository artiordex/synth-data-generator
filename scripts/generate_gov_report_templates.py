# -*- coding: utf-8 -*-
# =============================================================================
# 파일명: generate_gov_report_templates.py
# 경로: scripts/generate_gov_report_templates.py
# 목적: 충북청주시 기술검토(오픈API) 및 데이터설명서(파일데이터) 실무 양식을 반영하여
#       공공기관/공무원 맞춤형 표 중심의 DOCX 및 HWPX 보고서 템플릿을 생성하는 통합 엔진
# 작성자: 개발팀
# 작성일: 2026-09-16
# =============================================================================
"""Unified Generator for Government-Standard API and File Dataset Reports (DOCX & HWPX)."""
from __future__ import annotations

import argparse
from datetime import datetime
import io
import os
from pathlib import Path
import sys
from typing import Any, Dict, List, Optional, Tuple

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "apps" / "api" / "src"))
sys.path.insert(0, str(ROOT / "packages" / "synthetic_engine"))

import docx
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

# ---------------------------------------------------------------------------
# DOCX 공문서 스타일 유틸리티
# ---------------------------------------------------------------------------
COLOR_PRIMARY = RGBColor(30, 58, 138)     # Deep Navy (#1E3A8A)
COLOR_TEXT = RGBColor(31, 41, 55)         # Dark Gray (#1F2937)
COLOR_MUTED = RGBColor(107, 114, 128)     # Gray (#6B7280)
COLOR_HEADER_BG = "EBF2FA"                # Soft Blue-Gray Header
COLOR_ALT_ROW = "F9FAFB"                  # Soft Zebra Gray
COLOR_BORDER = "CBD5E1"                   # Slate Border (#CBD5E1)


def set_cell_background(cell, fill_hex: str):
    """셀 배경색 설정."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """셀 내부 여백 (dxa 단위: 20 dxa = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """테이블 전체 괘선 설정."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def format_table(table, col_widths: List[float], header_bg: str = COLOR_HEADER_BG):
    """공문서 표 스타일 일괄 적용."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, COLOR_BORDER)
    for row_idx, row in enumerate(table.rows):
        # 헤더 반복 속성
        if row_idx == 0:
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for col_idx, cell in enumerate(row.cells):
            # 열 너비 설정
            if col_idx < len(col_widths):
                cell.width = Inches(col_widths[col_idx])

            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)

            # 배경색 설정
            if row_idx == 0:
                set_cell_background(cell, header_bg)
            elif col_idx == 0 and len(row.cells) > 2 and row_idx % 2 == 1:
                set_cell_background(cell, "F8FAFC")

            # 셀 내 텍스트 폰트 설정
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    run.font.name = "맑은 고딕"
                    run._r.get_or_add_rPr().append(parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="맑은 고딕"/>'))
                    if row_idx == 0:
                        run.font.bold = True
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = COLOR_TEXT
                    else:
                        run.font.size = Pt(9)
                        run.font.color.rgb = COLOR_TEXT


def add_heading_1(doc, text: str):
    """대제목 (예: 1. 서비스 명세, Ⅰ. 데이터셋 개요)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "맑은 고딕"
    run._r.get_or_add_rPr().append(parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="맑은 고딕"/>'))
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p


def add_heading_2(doc, text: str):
    """중제목 (예: 가. API 서비스 개요, 1.1 데이터 일반 현황)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "맑은 고딕"
    run._r.get_or_add_rPr().append(parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="맑은 고딕"/>'))
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_TEXT
    return p


def add_heading_3(doc, text: str):
    """소제목 (예: 1) [이상/일반 상황 목록 서비스] 상세기능명세, □ 품질 진단)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "맑은 고딕"
    run._r.get_or_add_rPr().append(parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="맑은 고딕"/>'))
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = COLOR_TEXT
    return p


def add_notice_p(doc, text: str):
    """공무원 안내 문구 (예: ※ 항목구분 : 필수(1), 옵션(0))."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "맑은 고딕"
    run._r.get_or_add_rPr().append(parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="맑은 고딕"/>'))
    run.font.size = Pt(8.5)
    run.font.color.rgb = COLOR_MUTED
    return p


def add_code_box(doc, text: str):
    """코드/URL/XML/JSON 블록 박스."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    set_table_borders(table, "CBD5E1", sz="4")
    cell.width = Inches(6.5)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._r.get_or_add_rPr().append(parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="맑은 고딕"/>'))
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(15, 23, 42)


# ===========================================================================
# 1. API 데이터 DOCX 템플릿 생성기 (충북청주시 기술검토 표준 양식 기반)
# ===========================================================================
def build_api_docx_template(output_path: Path):
    """충북청주시 기술검토 보고서 양식을 정밀 반영한 공공 오픈API AI 가이드 DOCX 생성."""
    doc = Document()

    # 페이지 마진: 20mm
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 표지 / 헤더 타이틀
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(18)
    run_t = title_p.add_run("공공데이터 오픈API 활용가이드 및 AI 기술검토서")
    run_t.font.name = "맑은 고딕"
    run_t.font.size = Pt(20)
    run_t.font.bold = True
    run_t.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(24)
    run_sub = sub_p.add_run("AI 친화 공공데이터 표준 명세 및 오픈API 연계 계약 가이드라인")
    run_sub.font.name = "맑은 고딕"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = COLOR_MUTED

    # -----------------------------------------------------------------------
    # 1. 서비스 명세
    # -----------------------------------------------------------------------
    add_heading_1(doc, "1. 서비스 명세")
    add_heading_2(doc, "1.1 AI 친화 공공데이터 오픈API 서비스 개요")

    # 가. API 서비스 개요 표 (Table 1 형태)
    add_heading_3(doc, "가. API 서비스 개요")
    t1 = doc.add_table(rows=4, cols=3)
    t1_data = [
        ["API 서비스 정보", "API명(영문)", "{{dataset.identifier_or_api_name}}"],
        ["API 서비스 정보", "API명(국문)", "{{dataset.title}}"],
        ["API 서비스 정보", "API 설명", "{{dataset.description}}"],
        ["API 서비스\n보안적용기술 수준", "서비스 인증/권한", "[O] serviceKey    [ ] 인증서 (GPKI/NPKI)\n[ ] OAuth 2.0    [ ] 없음 (완전 공개)"]
    ]
    for r_idx, row in enumerate(t1.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = t1_data[r_idx][c_idx]
    # 병합
    t1.cell(0, 0).merge(t1.cell(2, 0))
    format_table(t1, [1.5, 1.4, 3.6])

    # 나. 상세기능 목록 표 (Table 2 형태)
    add_heading_3(doc, "나. 상세기능 목록")
    t2 = doc.add_table(rows=3, cols=4)
    t2.rows[0].cells[0].text = "번호"
    t2.rows[0].cells[1].text = "API명(국문)"
    t2.rows[0].cells[2].text = "상세기능명(영문)"
    t2.rows[0].cells[3].text = "상세기능명(국문)"

    t2.rows[1].cells[0].text = "1"
    t2.rows[1].cells[1].text = "{{dataset.title}}"
    t2.rows[1].cells[2].text = "getAiDatasetList"
    t2.rows[1].cells[3].text = "AI 데이터 목록 조회 서비스"

    t2.rows[2].cells[0].text = "2"
    t2.rows[2].cells[1].text = "{{dataset.title}}"
    t2.rows[2].cells[2].text = "getAiDatasetDetail"
    t2.rows[2].cells[3].text = "AI 데이터 상세 조회 서비스"
    t2.cell(1, 1).merge(t2.cell(2, 1))
    format_table(t2, [0.6, 2.3, 1.8, 1.8])

    # 다. 상세기능내역 (Table 3, 4, 5, 6 형태)
    add_heading_3(doc, "다. 상세기능내역")
    add_heading_3(doc, "1) [AI 데이터 목록 조회 서비스] 상세기능명세")

    t3 = doc.add_table(rows=4, cols=4)
    t3.rows[0].cells[0].text = "상세기능 번호"; t3.rows[0].cells[1].text = "1"
    t3.rows[0].cells[2].text = "상세기능 유형"; t3.rows[0].cells[3].text = "조회 (목록)"
    t3.rows[1].cells[0].text = "상세기능명(국문)"; t3.rows[1].cells[1].text = "AI 데이터 목록 조회 서비스"
    t3.cell(1, 1).merge(t3.cell(1, 3))
    t3.rows[2].cells[0].text = "상세기능 설명"; t3.rows[2].cells[1].text = "{{dataset.description}}에 대한 페이징 및 조건별 목록 검색을 제공합니다."
    t3.cell(2, 1).merge(t3.cell(2, 3))
    t3.rows[3].cells[0].text = "Call Back URL"; t3.rows[3].cells[1].text = "{{structure.api_endpoint_url}}/getAiDatasetList"
    t3.cell(3, 1).merge(t3.cell(3, 3))
    format_table(t3, [1.5, 1.8, 1.4, 1.8])

    # 요청 메시지 명세
    add_notice_p(doc, "■ 요청 메시지 명세 (Request Parameters)")
    add_notice_p(doc, "※ 항목구분 : 필수(1), 옵션(0)")
    t4 = doc.add_table(rows=5, cols=6)
    t4_headers = ["항목명(영문)", "항목명(국문)", "항목크기", "항목구분", "샘플데이터", "항목설명"]
    for c_idx, h in enumerate(t4_headers):
        t4.rows[0].cells[c_idx].text = h
    t4_rows = [
        ["serviceKey", "공공데이터 인증키", "100", "1", "인증키(URL Encoded)", "공공데이터포털 발급 API 인증키"],
        ["pageNo", "페이지 번호", "4", "0", "1", "조회할 페이지 번호 (기본 1)"],
        ["numOfRows", "페이지당 결과 수", "4", "0", "10", "한 페이지당 표출 데이터 수 (기본 10)"],
        ["returnType", "반환 응답 포맷", "4", "0", "json", "응답 포맷 (json / xml)"]
    ]
    for r_idx, row in enumerate(t4_rows, start=1):
        for c_idx, val in enumerate(row):
            t4.rows[r_idx].cells[c_idx].text = val
    format_table(t4, [1.4, 1.3, 0.7, 0.7, 1.1, 1.3])

    # 응답 메시지 명세
    add_notice_p(doc, "■ 응답 메시지 명세 (Response Parameters)")
    t5 = doc.add_table(rows=6, cols=6)
    for c_idx, h in enumerate(t4_headers):
        t5.rows[0].cells[c_idx].text = h
    t5_rows = [
        ["resultCode", "결과코드", "2", "1", "00", "정상 처리 결과 코드 (00: 정상)"],
        ["resultMsg", "결과메시지", "50", "1", "NORMAL_SERVICE", "결과 처리 상태 메시지"],
        ["totalCount", "총 데이터 수", "10", "1", "{{dataset.record_count}}", "전체 누적 레코드 건수"],
        ["pageNo", "페이지 번호", "4", "1", "1", "현재 반환된 페이지 번호"],
        ["items", "데이터 목록", "-", "1", "[배열 객체]", "상세 관측 데이터 필드 배열"]
    ]
    for r_idx, row in enumerate(t5_rows, start=1):
        for c_idx, val in enumerate(row):
            t5.rows[r_idx].cells[c_idx].text = val
    format_table(t5, [1.4, 1.3, 0.7, 0.7, 1.1, 1.3])

    # 요청 및 응답 예시
    add_notice_p(doc, "■ 요청 URL 예시")
    add_code_box(doc, "{{structure.api_endpoint_url}}/getAiDatasetList?serviceKey={{인증키}}&pageNo=1&numOfRows=10&returnType=json")

    add_notice_p(doc, "■ 응답 JSON 메시지 예시")
    add_code_box(doc, '{\n  "response": {\n    "header": {\n      "resultCode": "00",\n      "resultMsg": "NORMAL_SERVICE"\n    },\n    "body": {\n      "pageNo": 1,\n      "totalCount": 50000,\n      "numOfRows": 10,\n      "items": [\n        {"id": 1, "name": "샘플데이터1", "status": "Y"}\n      ]\n    }\n  }\n}')

    # -----------------------------------------------------------------------
    # 2. OpenAPI 에러 코드 정리
    # -----------------------------------------------------------------------
    add_heading_1(doc, "2. OpenAPI 에러 코드 정리")
    add_heading_2(doc, "2.1 공공데이터포털 표준 에러코드 (Gateway Error)")
    t11 = doc.add_table(rows=5, cols=3)
    t11_headers = ["에러코드", "에러메시지", "설명 및 조치방법"]
    for c_idx, h in enumerate(t11_headers):
        t11.rows[0].cells[c_idx].text = h
    t11_rows = [
        ["01", "APPLICATION_ERROR", "제공기관 어플리케이션 서버 응답 오류"],
        ["04", "HTTP_ERROR", "HTTP 통신 실패 및 연결 시간 초과"],
        ["10", "INVALID_REQUEST_PARAMETER_ERROR", "잘못된 요청 파라미터 전달"],
        ["12", "NO_OPENAPI_SERVICE_ERROR", "해당 오픈API 서비스가 없거나 폐기됨"]
    ]
    for r_idx, row in enumerate(t11_rows, start=1):
        for c_idx, val in enumerate(row):
            t11.rows[r_idx].cells[c_idx].text = val
    format_table(t11, [1.2, 2.3, 3.0])

    add_heading_2(doc, "2.2 제공기관 자체 에러코드 (Institutional Error)")
    t12 = doc.add_table(rows=4, cols=3)
    for c_idx, h in enumerate(t11_headers):
        t12.rows[0].cells[c_idx].text = h
    t12_rows = [
        ["90", "DATABASE_TIMEOUT_ERROR", "기관 내부 DB 쿼리 수행 시간 초과"],
        ["91", "RATE_LIMIT_EXCEEDED", "일일 허용 트래픽 호출 한도 초과"],
        ["99", "UNKNOWN_ERROR", "알 수 없는 시스템 내부 예외"]
    ]
    for r_idx, row in enumerate(t12_rows, start=1):
        for c_idx, val in enumerate(row):
            t12.rows[r_idx].cells[c_idx].text = val
    format_table(t12, [1.2, 2.3, 3.0])

    # -----------------------------------------------------------------------
    # 3. AI 친화도 및 모델 연계 가이드
    # -----------------------------------------------------------------------
    add_heading_1(doc, "3. AI 친화도 및 모델 연계 가이드")
    t_ai = doc.add_table(rows=5, cols=2)
    t_ai_data = [
        ["구분", "AI 연계 지침 및 분석 내용"],
        ["추천 AI Task", "{{ai.tasks}} (실시간 추론, 이상 탐지, 시계열 예측 등)"],
        ["RAG 및 LLM 파이프라인 연계", "JSON 응답 봉투(Envelope) 제거 후 순수 데이터 노드 임베딩 권고"],
        ["학습/검증 데이터 분할", "Train : Validation : Test = 70 : 15 : 15 (시간 흐름 보존 분할)"],
        ["대용량 API 호출 최적화", "대량 적재 시 벌크 API 활용 및 Apache Parquet 전환 권고"]
    ]
    for r_idx, row in enumerate(t_ai.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = t_ai_data[r_idx][c_idx]
    format_table(t_ai, [1.8, 4.7])

    # 저장
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"[성공] API용 DOCX 템플릿 생성 완료: {output_path.resolve()} ({output_path.stat().st_size:,} bytes)")


# ===========================================================================
# 2. 파일데이터 DOCX 템플릿 생성기 (제조공정 로봇티칭 및 데이터설명서 표준 양식 기반)
# ===========================================================================
def build_file_docx_template(output_path: Path):
    """제조공정 로봇티칭 HWP 데이터설명서 양식을 정밀 반영한 파일데이터 AI 가이드 DOCX 생성."""
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(18)
    run_t = title_p.add_run("AI 친화·고가치 공공데이터셋 설명서")
    run_t.font.name = "맑은 고딕"
    run_t.font.size = Pt(20)
    run_t.font.bold = True
    run_t.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(24)
    run_sub = sub_p.add_run("파일데이터(CSV / XLSX) 품질 진단 및 AI 학습데이터 표준 명세서")
    run_sub.font.name = "맑은 고딕"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = COLOR_MUTED

    # 1. 데이터셋 개요
    add_heading_1(doc, "Ⅰ. 데이터셋 일반 개요")
    t1 = doc.add_table(rows=6, cols=4)
    t1_data = [
        ["데이터명", "{{dataset.title}}", "제공기관", "{{dataset.publisher}}"],
        ["구축/생성목적", "{{dataset.purpose}}", "운영부서", "{{governance.managing_department}}"],
        ["데이터 규모", "{{dataset.record_count}} 건 / {{dataset.byte_size}} bytes", "데이터 형식", "{{dataset.formats}}"],
        ["시간적 범위", "{{dataset.temporal}}", "공간적 범위", "{{dataset.spatial}}"],
        ["갱신주기", "{{dataset.update_frequency}}", "배포 버전", "{{dataset.version}}"],
        ["관련 근거법령", "{{governance.legal_basis}}", "문서 상태", "{{document.status}}"]
    ]
    for r_idx, row in enumerate(t1.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = t1_data[r_idx][c_idx]
    format_table(t1, [1.4, 1.85, 1.4, 1.85])

    # 2. 데이터 사전 (Data Dictionary)
    add_heading_1(doc, "Ⅱ. 전체 필드 데이터 사전 (Data Dictionary)")
    add_notice_p(doc, "※ 데이터셋에 포함된 모든 관측 컬럼에 대한 물리명, 논리명, 데이터타입 및 제약조건 명세")
    t2 = doc.add_table(rows=5, cols=6)
    t2_headers = ["컬럼 물리명", "컬럼 논리명", "자료형", "필수여부", "단위/코드", "항목 설명 및 제약조건"]
    for c_idx, h in enumerate(t2_headers):
        t2.rows[0].cells[c_idx].text = h
    t2_rows = [
        ["SEQ_NO", "일련번호", "INTEGER", "필수(1)", "무차원", "레코드 고유 식별자 (Candidate Key)"],
        ["ORG_CD", "기관코드", "VARCHAR(7)", "필수(1)", "행정코드", "행정표준코드관리시스템 공공기관 고유코드"],
        ["CRTR_YMD", "기준일자", "DATE", "필수(1)", "YYYY-MM-DD", "공시 및 측정 기준일시"],
        ["MESR_VAL", "측정수치", "NUMERIC(10,2)", "옵션(0)", "ppm", "센서 정밀 계측 수치 (이상치 필터링 적용)"]
    ]
    for r_idx, row in enumerate(t2_rows, start=1):
        for c_idx, val in enumerate(row):
            t2.rows[r_idx].cells[c_idx].text = val
    format_table(t2, [1.1, 1.1, 0.9, 0.7, 0.9, 1.8])

    # 3. 공통 6대 품질 진단 결과
    add_heading_1(doc, "Ⅲ. AI 친화도 및 6대 품질 진단 결과 ([REF-02])")
    add_notice_p(doc, "※ 정부 표준 AI 데이터 품질관리 가이드 v4.0에 근거한 결정적 품질 프로파일링 점수")
    t3 = doc.add_table(rows=7, cols=4)
    t3_headers = ["품질 평가 지표", "점수 / 적합율", "판정 상태", "측정 근거 및 세부 내역"]
    for c_idx, h in enumerate(t3_headers):
        t3.rows[0].cells[c_idx].text = h
    t3_rows = [
        ["1. 완전성 (Completeness)", "{{quality.completeness.score}}점", "{{quality.completeness.status}}", "결측치 관측률 0.01% 미만, 필수 컬럼 100% 충족"],
        ["2. 유효성 (Validity)", "{{quality.validity.score}}점", "{{quality.validity.status}}", "표준 데이터 타입 및 날짜 포맷 규칙 100% 준수"],
        ["3. 일관성 (Consistency)", "{{quality.consistency.score}}점", "{{quality.consistency.status}}", "상하위 도메인 및 코드값 간 상호 논리 모순 없음"],
        ["4. 정확성 (Accuracy)", "{{quality.accuracy.score}}점", "{{quality.accuracy.status}}", "3-Sigma 이상치 범위 내 정밀도 확보"],
        ["5. 유일성 (Uniqueness)", "{{quality.uniqueness.score}}점", "{{quality.uniqueness.status}}", "중복 레코드 0건, 고유 후보키 식별 완료"],
        ["6. 적시성 (Timeliness)", "{{quality.timeliness.score}}점", "{{quality.timeliness.status}}", "갱신 주기(반기) 준수 및 최신 공시 기준일 확인"]
    ]
    for r_idx, row in enumerate(t3_rows, start=1):
        for c_idx, val in enumerate(row):
            t3.rows[r_idx].cells[c_idx].text = val
    format_table(t3, [1.8, 1.0, 1.0, 2.7])

    # 4. 데이터 구축 및 가공 공정
    add_heading_1(doc, "Ⅳ. 데이터 구축·정제·가공 공정 및 계보")
    t4 = doc.add_table(rows=5, cols=4)
    t4_headers = ["구축 공정 단계", "주요 처리 규칙", "사용 도구/알고리즘", "품질 검증 기준"]
    for c_idx, h in enumerate(t4_headers):
        t4.rows[0].cells[c_idx].text = h
    t4_rows = [
        ["1. 수집/획득", "원천 행정 시스템 연계 추출", "ETL 배치 파이프라인", "원천 레코드 수 전수 대조"],
        ["2. 정제/가공", "인코딩 변환(UTF-8) 및 공백 제거", "데이터 정제 스크립트", "비표준 문자 및 결측 제거"],
        ["3. 비식별화", "개인정보(주민번호/연락처) 마스킹", "개인정보 비식별화 엔진", "재식별 위험성 평가 통과"],
        ["4. 검수/승인", "6대 품질지표 정량 측정 검증", "품질관리 검수 시스템", "품질 합격 기준 충족"]
    ]
    for r_idx, row in enumerate(t4_rows, start=1):
        for c_idx, val in enumerate(row):
            t4.rows[r_idx].cells[c_idx].text = val
    format_table(t4, [1.4, 2.0, 1.6, 1.5])

    # 5. AI 모델 연계 및 대용량 최적화
    add_heading_1(doc, "Ⅴ. AI 모델 연계 및 대용량 파이프라인 가이드")
    t5 = doc.add_table(rows=4, cols=2)
    t5_data = [
        ["구분", "세부 권고 내용"],
        ["추천 AI 알고리즘", "{{ai.recommended_models}} (예: Gradient Boosting, Transformer 정형 분류)"],
        ["데이터셋 분할 권고", "학습용(Train) 70% : 검증용(Val) 15% : 평가용(Test) 15%"],
        ["대용량 열지향 전환", "대규모 배치 학습 시 I/O 성능 향상을 위해 Apache Parquet 포맷 전환 필수 권고"]
    ]
    for r_idx, row in enumerate(t5.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = t5_data[r_idx][c_idx]
    format_table(t5, [1.8, 4.7])

    # 6. 이용 조건 및 라이선스
    add_heading_1(doc, "Ⅵ. 데이터 이용·배포 및 거버넌스")
    t6 = doc.add_table(rows=4, cols=4)
    t6_data = [
        ["공식 라이선스", "{{usage.license}}", "이용 허용범위", "상업적·비상업적 이용 및 변형 가능"],
        ["개인정보 포함여부", "미포함 (비식별화 완료)", "보안 수준", "공개 (전국민 개방)"],
        ["담당 부서", "{{governance.managing_department}}", "담당자 연락처", "{{governance.contact.phone}}"],
        ["문서 발간 상태", "{{document.status}}", "기관 확인 필요", "{{reviewRequired.count}} 건 미결"]
    ]
    for r_idx, row in enumerate(t6.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = t6_data[r_idx][c_idx]
    format_table(t6, [1.4, 1.85, 1.4, 1.85])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"[성공] 파일데이터용 DOCX 템플릿 생성 완료: {output_path.resolve()} ({output_path.stat().st_size:,} bytes)")


# ===========================================================================
# 3. HWPX 확장 렌더링 함수 (구조적 결손 보완)
# ===========================================================================
def build_expanded_hwpx_templates(adr_dir: Path):
    """기존 ai_guide_template을 활용하여 완벽한 표 기반의 API용 및 파일데이터용 HWPX 생성."""
    from synthetic_api.application.services.ai_guide_template import (
        render_template, TemplateGuideRequest, FieldAnnotation
    )

    # 1. 파일데이터용 HWPX
    file_c_meta = {
        "schema_version": "2.0",
        "data_category": "file",
        "format": "csv",
        "root_type": "table",
        "sha256": "94c1064a7898f0e175f62d0fae82935b9a8f1e317b381087f0567e009854d60f",
        "byte_size": 2457600,
        "title": "공공 행정 표준 파일데이터셋 가이드",
        "tables": [{"name": "TB_PUBLIC_STD", "row_count": 100000, "col_count": 10}],
        "record_sets": [{"name": "전체 표준 공시 레코드", "count": 100000}],
        "fields": [
            {"path": "일련번호", "types": ["integer"], "occurrences": 100000, "null_count": 0, "empty_count": 0, "examples": [1, 2, 3]},
            {"path": "기관코드", "types": ["string"], "occurrences": 100000, "null_count": 0, "empty_count": 0, "examples": ["1470000"]},
            {"path": "데이터셋명", "types": ["string"], "occurrences": 100000, "null_count": 0, "empty_count": 0, "examples": ["표준 안전 데이터"]},
            {"path": "기준일자", "types": ["string"], "occurrences": 100000, "null_count": 0, "empty_count": 0, "examples": ["2026-06-30"]},
            {"path": "분류구분", "types": ["string"], "occurrences": 100000, "null_count": 15, "empty_count": 0, "examples": ["보건의료"]},
            {"path": "측정수치", "types": ["number"], "occurrences": 100000, "null_count": 50, "empty_count": 0, "examples": [102.5]},
            {"path": "품질적합여부", "types": ["string"], "occurrences": 100000, "null_count": 0, "empty_count": 0, "examples": ["Y"]},
            {"path": "처리상태코드", "types": ["string"], "occurrences": 100000, "null_count": 0, "empty_count": 0, "examples": ["01"]}
        ],
        "quality_metrics": [
            {"category": "COMPLETENESS", "status": "MEASURED", "scope": "전체 10만 행 전수 검사", "observed": 799935, "missing": 65, "score": 99},
            {"category": "VALIDITY", "status": "MEASURED", "scope": "표준 데이터 타입 검증", "score": 100},
            {"category": "CONSISTENCY", "status": "MEASURED", "scope": "기관코드 논리 정합성", "score": 100},
            {"category": "ACCURACY", "status": "REVIEW_REQUIRED", "scope": "도메인 수치 정밀 검증", "score": None},
            {"category": "UNIQUENESS", "status": "MEASURED", "scope": "복합키 유일성 검사", "score": 100},
            {"category": "TIMELINESS", "status": "REVIEW_REQUIRED", "scope": "공시 갱신주기 검토", "score": None}
        ],
        "warnings": ["대용량 권고: 10만 건 이상의 데이터셋은 Parquet 압축 포맷 전환이 권장됩니다."]
    }

    file_i_meta = {
        "publisher": "대한민국 식품의약품안전처",
        "description": "공공 행정 표준 파일데이터셋 설명서 및 AI 학습용 가이드라인 보고서입니다.",
        "department": "디지털안전정보과",
        "legal_basis": "공공데이터의 제공 및 이용 활성화에 관한 법률",
        "landing_page": "https://www.data.go.kr",
        "contact": "데이터 담당관 (043-719-0000)",
        "license": "공공누리 제1유형: 출처표시 (상업적 이용 및 변형 가능)",
        "rights": "식품의약품안전처 공공저작물",
        "update_frequency": "반기",
        "version": "v1.0.0",
        "issued": "2026-01-01",
        "modified": "2026-06-30",
        "temporal": "2026-01-01 ~ 2026-06-30",
        "spatial": "대한민국 전역",
        "source_datasets": "식약처 내부 통합 업무 DB",
        "transformation": "식별자 매핑 및 결측치 플래그 처리",
        "imputation": "선형 보간법",
        "training_split": "지도학습 분류 (Train:Val:Test = 70:15:15)",
        "limitations": "특정 공시 데이터로서 도메인 외 적용 시 주의 필요"
    }

    file_ann = {
        "일련번호": FieldAnnotation(label="연번", description="레코드 고유 식별자", unit="무차원", codes="정수"),
        "기관코드": FieldAnnotation(label="기관코드", description="공공기관 표준 식별코드", unit="코드", codes="7자리"),
        "데이터셋명": FieldAnnotation(label="공시명", description="품목허가 공시 명칭", unit="문자열", codes="해당없음"),
        "기준일자": FieldAnnotation(label="기준일", description="측정 기준일자 (ISO 8601)", unit="날짜", codes="YYYY-MM-DD"),
        "분류구분": FieldAnnotation(label="분류", description="업무 대분류", unit="범주", codes="식품/의약품/의료기기"),
        "측정수치": FieldAnnotation(label="계측값", description="실험실 계측 수치", unit="ppm", codes="수치"),
        "품질적합여부": FieldAnnotation(label="적합여부", description="공식 규격 적합 판정", unit="플래그", codes="Y/N"),
        "처리상태코드": FieldAnnotation(label="상태코드", description="행정 처리 상태", unit="코드", codes="01, 02, 03")
    }

    req_file = TemplateGuideRequest(canonical_metadata=file_c_meta, metadata=file_i_meta, field_annotations=file_ann)
    file_hwpx_bytes = render_template(req_file)
    p_file = adr_dir / "AI친화_고가치_데이터셋_파일데이터용_템플릿.hwpx"
    p_file.write_bytes(file_hwpx_bytes)
    print(f"[성공] 파일데이터용 HWPX 템플릿 생성 완료: {p_file.resolve()} ({len(file_hwpx_bytes):,} bytes)")

    # 2. API용 HWPX
    api_c_meta = {
        "schema_version": "2.0",
        "data_category": "api",
        "format": "json",
        "root_type": "object",
        "sha256": "a1b2c3d4e5f67890123456789abcdef0123456789abcdef0123456789abcdef0",
        "byte_size": 524288,
        "title": "공공데이터 오픈API 조회 서비스 가이드",
        "tables": [{"name": "API_RESPONSE_ITEMS", "row_count": 50000, "col_count": 8}],
        "record_sets": [{"name": "응답 items 배열", "count": 50000}],
        "fields": [
            {"path": "response/header/resultCode", "types": ["string"], "occurrences": 50000, "null_count": 0, "empty_count": 0, "examples": ["00"]},
            {"path": "response/header/resultMsg", "types": ["string"], "occurrences": 50000, "null_count": 0, "empty_count": 0, "examples": ["NORMAL_SERVICE"]},
            {"path": "response/body/pageNo", "types": ["integer"], "occurrences": 50000, "null_count": 0, "empty_count": 0, "examples": [1]},
            {"path": "response/body/totalCount", "types": ["integer"], "occurrences": 50000, "null_count": 0, "empty_count": 0, "examples": [50000]},
            {"path": "response/body/items/*/id", "types": ["integer"], "occurrences": 50000, "null_count": 0, "empty_count": 0, "examples": [1001]},
            {"path": "response/body/items/*/name", "types": ["string"], "occurrences": 50000, "null_count": 0, "empty_count": 0, "examples": ["의약품 표준항목"]},
            {"path": "response/body/items/*/status", "types": ["string"], "occurrences": 50000, "null_count": 0, "empty_count": 0, "examples": ["Y"]}
        ],
        "quality_metrics": [
            {"category": "COMPLETENESS", "status": "MEASURED", "scope": "API 응답 페이로드", "observed": 350000, "missing": 0, "score": 100},
            {"category": "VALIDITY", "status": "MEASURED", "scope": "JSON 구문 및 스키마", "score": 100}
        ],
        "warnings": ["API 호출 시 serviceKey URL 인코딩 여부를 확인하세요."]
    }

    api_i_meta = {
        "publisher": "대한민국 식품의약품안전처",
        "description": "공공데이터 오픈API 조회 서비스 활용 가이드라인 및 기술검토 명세서입니다.",
        "endpoint": "http://apis.data.go.kr/1471000/FoodDrugInfoService",
        "http_method": "GET",
        "authentication": "공공데이터포털 발급 serviceKey (URL Encoded)",
        "request_parameters": "serviceKey(필수), pageNo(옵션), numOfRows(옵션), returnType(옵션)",
        "pagination": "pageNo, numOfRows 지원 (최대 100건/페이지)",
        "error_codes": "00: 정상, 01: 어플리케이션 에러, 10: 잘못된 요청 파라미터",
        "response_path": "response/body/items",
        "license": "공공누리 제1유형: 출처표시",
        "training_split": "실시간 API 추론 및 RAG 컨텍스트 주입용",
        "limitations": "일일 호출 한도 10,000건 초과 시 추가 승인 필요"
    }

    api_ann = {
        "response/header/resultCode": FieldAnnotation(label="결과코드", description="응답 상태 코드", unit="코드", codes="00: 정상"),
        "response/header/resultMsg": FieldAnnotation(label="결과메시지", description="응답 상태 메시지", unit="문자열", codes="NORMAL_SERVICE"),
        "response/body/pageNo": FieldAnnotation(label="페이지번호", description="현재 페이지 번호", unit="정수", codes="1 이상"),
        "response/body/totalCount": FieldAnnotation(label="총건수", description="전체 데이터 건수", unit="건", codes="정수"),
        "response/body/items/*/id": FieldAnnotation(label="식별ID", description="항목 고유 식별자", unit="무차원", codes="정수"),
        "response/body/items/*/name": FieldAnnotation(label="항목명", description="공시 항목 명칭", unit="문자열", codes="해당없음"),
        "response/body/items/*/status": FieldAnnotation(label="상태", description="공시 상태", unit="플래그", codes="Y/N")
    }

    req_api = TemplateGuideRequest(canonical_metadata=api_c_meta, metadata=api_i_meta, field_annotations=api_ann)
    api_hwpx_bytes = render_template(req_api)
    p_api = adr_dir / "AI친화_고가치_데이터셋_API용_템플릿.hwpx"
    p_api.write_bytes(api_hwpx_bytes)
    print(f"[성공] API용 HWPX 템플릿 생성 완료: {p_api.resolve()} ({len(api_hwpx_bytes):,} bytes)")


# ===========================================================================
# 4. 메인 실행 함수
# ===========================================================================
def main():
    parser = argparse.ArgumentParser(description="공공기관 표준 API 및 파일데이터 DOCX/HWPX 템플릿 생성기")
    parser.add_argument("--format", "-f", choices=["docx", "hwpx", "all"], default="all", help="생성할 포맷 (기본: all)")
    parser.add_argument("--target-dir", "-d", default=str(ROOT / "docs" / "adr"), help="템플릿 생성 대상 디렉토리")

    args = parser.parse_args()
    target_dir = Path(args.target_dir)
    target_dir.mkdir(parents=True, exist_ok=True)

    print("==================================================================")
    print(" 공공기관 표준 보고서 템플릿 생성 시작")
    print(" - 대상 디렉토리:", target_dir.resolve())
    print(" - 생성 포맷:", args.format)
    print("==================================================================")

    if args.format in ["docx", "all"]:
        build_api_docx_template(target_dir / "AI친화_고가치_데이터셋_API용_템플릿.docx")
        build_file_docx_template(target_dir / "AI친화_고가치_데이터셋_파일데이터용_템플릿.docx")

    if args.format in ["hwpx", "all"]:
        build_expanded_hwpx_templates(target_dir)

    print("==================================================================")
    print(" 모든 공공기관용 DOCX 및 HWPX 템플릿 생성이 성공적으로 완료되었습니다.")
    print("==================================================================")


if __name__ == "__main__":
    main()
